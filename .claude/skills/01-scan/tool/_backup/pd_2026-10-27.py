# -*- coding: utf-8 -*-
"""pd.py — công cụ cố định của skill 01-scan. Model chỉ ghi JSON; file PROJECT_DATA do script này dựng.

Đường dẫn: <goc> = thư mục dự án (vd D:\\COSTING EXCEL\\_ACTIVE\\NORTON), <dot> = tên thư mục trong 00_input.
  đầu vào : <goc>\\00_input\\<dot>
  dữ liệu : <goc>\\01_scan\\_data\\<dot>\\<SHEET>.<buoc>.json
  kết quả : <goc>\\01_scan\\PROJECT_DATA_<dot>.xlsx

Lệnh (chạy: py -I pd.py <lệnh> ...):
  kiemke <goc> <dot> --ds                    chỉ liệt kê gọn theo thư mục (không đọc nội dung) để chọn file cần quét
  kiemke <goc> <dot>                      kiểm kê file + trang PDF -> FILES/PAGES .kiemke.json, cache text
  tim    <goc> <dot> <regex> [--file chuoi]  tìm trang có chữ khớp regex (từ cache text)
  ungvien <goc> <dot> [--top 12]             sau bước boq: mỗi Mã -> trang PDF ứng viên (mã / từ khóa Hạng mục) -> _ungvien.txt
  chu    <pdf> <trang> [--clip x0,y0,x1,y1]  in lớp text của trang (clip theo tỉ lệ 0..1 của trang)
  so     <pdf> <trang> [--clip ...]          in các số trong lớp text kèm toạ độ (đọc kích thước)
  anh    <pdf> <trang> <png> [--dpi 150] [--clip ...]  xuất ảnh trang/vùng để xem bằng Read
  xem    <xlsx> [sheet] [A1:Z60]             in giá trị ô (đọc BoQ, Material List)
  gemini <a.png> [b.png …] [--claude "900,450,450"] [--quansat "Tủ lavabo (V01)"] [--hoi "câu hỏi"]
                                             Gemini (agy) đọc độc lập: số kích thước / bộ trường quan sát / câu hỏi riêng
  mail   <file.msg>                          in tiêu đề, người gửi/nhận, đính kèm, nội dung mail Outlook
  build  <goc> <dot>                         ghép JSON theo form -> PROJECT_DATA_<dot>.xlsx (giữ sửa tay của user)
  check  <goc> <dot> [--den buoc]            kiểm; exit 1 nếu có LỖI. --den: chỉ đòi cột bắt buộc của các bước tới buoc
  cot    [--buoc boq]                        in các cột bước đó được ghi (từ form) + ghi chú
"""
import sys, os, re, json, glob, shutil, datetime
sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
FORM = os.path.join(HERE, "..", "form", "PROJECT_DATA_FORM.xlsx")
BUOC = ["kiemke", "boq", "spec", "banve", "doichieu", "user"]
SKIP_EXT = {".dwl", ".dwl2", ".bak", ".tmp", ".log", ".plt", ".ctb", ".stb"}
# L0: file nhiễu theo tên / thư mục (khung tên CAD, xref, bản lưu) — bỏ khỏi kiểm kê, in số lượng khi chạy
SKIP_NAME = re.compile(r"^(frame|xref)[-_ ]|^thumbs\.db$", re.I)
SKIP_DIR = re.compile(r"(^|[\\/])(archive|backup|old|_old|cũ)([\\/]|$)", re.I)


# ---------------------------------------------------------------- đường dẫn
def paths(goc, dot):
    goc = os.path.abspath(goc)
    return dict(goc=goc, inp=os.path.join(goc, "00_input", dot), data=os.path.join(goc, "01_scan", "_data", dot),
                out=os.path.join(goc, "01_scan", f"PROJECT_DATA_{dot}.xlsx"), bak=os.path.join(goc, "01_scan", "_backup"))


def parse_clip(args, page):
    if "--clip" not in args: return None
    import pymupdf
    x0, y0, x1, y1 = [float(v) for v in args[args.index("--clip") + 1].split(",")]
    r = page.rect
    return pymupdf.Rect(r.x0 + x0 * r.width, r.y0 + y0 * r.height, r.x0 + x1 * r.width, r.y0 + y1 * r.height)


def opt(args, name, default=None):
    return args[args.index(name) + 1] if name in args else default


# ---------------------------------------------------------------- form
def load_form():
    import openpyxl
    wb = openpyxl.load_workbook(FORM)
    sheets = {}
    for r in wb["_SHEETS"].iter_rows(min_row=2, values_only=True):
        if r[0]: sheets[r[0]] = dict(kieu=r[1], khoa=[k.strip() for k in (r[2] or "").split(",") if k.strip()])
    cols = {}
    for r in wb["_COT"].iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        sh, c, buoc, bb, kieu, ds, rong, nhom, ct, note = r[:10]
        cols.setdefault(sh, {})[c] = dict(buoc=[b.strip() for b in (buoc or "").split(",") if b.strip()], batbuoc=bb == "x",
                                          kieu=kieu, ds=[x.strip() for x in (ds or "").split("|") if x.strip()], ct=ct, note=note)
    for sh in sheets:
        if sh in wb.sheetnames and sheets[sh]["kieu"] != "AUTO":
            hdr = [c.value for c in wb[sh][1] if c.value is not None]
            sheets[sh]["hdr"] = hdr
            for h in hdr:
                if h not in cols.get(sh, {}): raise SystemExit(f"Form lỗi: {sh}!{h} không có trong _COT")
    return wb, sheets, cols


# ---------------------------------------------------------------- dữ liệu JSON
def load_json(data):
    """-> {sheet: [(buoc, row_dict), ...]} theo thứ tự BUOC."""
    out, errs = {}, []
    files = glob.glob(os.path.join(data, "*.json"))
    def order(f):
        b = os.path.basename(f).split(".")[1].split("-")[0] if os.path.basename(f).count(".") >= 2 else "?"
        return BUOC.index(b) if b in BUOC else 99
    for f in sorted(files, key=order):
        name = os.path.basename(f)
        if name.startswith("_"): continue
        parts = name.split(".")
        if len(parts) != 3 or parts[1].split("-")[0] not in BUOC:
            errs.append(f"Tên file JSON sai '{name}' (đúng: <SHEET>.<buoc>.json hoặc <SHEET>.<buoc>-<nhóm>.json, buoc ∈ {BUOC})"); continue
        try:
            obj = json.load(open(f, encoding="utf-8"))
        except Exception as e:
            errs.append(f"{name}: JSON hỏng: {e}"); continue
        rows = obj["rows"] if isinstance(obj, dict) else obj
        for r in rows: out.setdefault(parts[0], []).append((parts[1], r))
    return out, errs


def norm_links(v, files_by_id, inp):
    """Giá trị ô LINK -> (text, [(nhan, abs_file, trang, o_dich)]). Chấp nhận chuỗi, list link, hoặc {text, links}."""
    if v is None or v == "": return "", []
    if isinstance(v, (str, int, float)): return str(v), []
    text, links = "", v
    if isinstance(v, dict): text, links = v.get("text", ""), v.get("links", [])
    res = []
    for l in links:
        if isinstance(l, (list, tuple)):
            l = dict(zip(["nhan", "file", "trang", "o"], l))
        f = str(l.get("file", ""))
        if re.fullmatch(r"F\d{3,4}", f): f = files_by_id.get(f, f)
        elif f and not os.path.isabs(f): f = os.path.join(inp, f)
        res.append((str(l.get("nhan") or os.path.basename(f)), os.path.normpath(f) if f else "", l.get("trang") or "", l.get("o") or ""))
    if not text and res:
        text = res[0][0] if len(res) == 1 else f"{res[0][0]}  ▾{len(res)}"
    return text, res


def merge(sheets, cols, raw):
    """Ghép dữ liệu các bước. KHOA: theo khóa, mỗi ô 1 bước; NOI: nối dòng. user ghi đè sau cùng."""
    errs, merged, owner = [], {}, {}
    for sh, items in raw.items():
        items = [(b, r) for b, r in items]
        if sh not in sheets or sheets[sh]["kieu"] == "AUTO":
            errs.append(f"Sheet '{sh}' không có trong form (hoặc là AUTO)"); continue
        meta, keys = cols[sh], sheets[sh]["khoa"]
        if sheets[sh]["kieu"] == "NOI":
            rows = []
            for bf, r in items:
                b = bf.split("-")[0]
                for c in r:
                    if c not in meta: errs.append(f"{sh}: cột lạ '{c}' (bước {bf})")
                    elif b != "user" and meta[c]["buoc"] and b not in meta[c]["buoc"]: errs.append(f"{sh}!{c}: bước {b} không được ghi cột này")
                rr = dict(r); rr.setdefault("Bước", bf) if "Bước" in meta else None
                rows.append(rr)
            merged[sh] = rows; continue
        rows, own = {}, {}
        for bf, r in items:
            b = bf.split("-")[0]
            miss = [k for k in keys if r.get(k) in (None, "")]
            if miss: errs.append(f"{sh}: dòng thiếu khóa {miss} (bước {b}): {str(r)[:80]}"); continue
            key = tuple(str(r[k]).strip() for k in keys)
            row = rows.setdefault(key, {k: r[k] for k in keys})
            for c, v in r.items():
                if c in keys: continue
                if c not in meta: errs.append(f"{sh}: cột lạ '{c}' (bước {b})"); continue
                if b != "user" and meta[c]["buoc"] and b not in meta[c]["buoc"]:
                    errs.append(f"{sh}!{c}: bước {b} không được ghi cột này (chỉ {meta[c]['buoc']})"); continue
                if v in (None, ""): continue
                prev = own.get((key, c))
                if bf == "user" and prev and row.get(c) != v:
                    owner.setdefault("user_de", []).append(f"{sh} {'/'.join(key)} · {c}: user '{str(v)[:40]}' thay model '{str(row.get(c))[:40]}'")
                elif prev and prev != bf and prev != "kiemke" and row.get(c) != v:  # giá trị máy đoán (kiemke) được sửa đè
                    errs.append(f"{sh} {'/'.join(key)} · {c}: 2 bước ghi khác nhau ({prev} vs {bf})"); continue
                row[c] = v; own[(key, c)] = bf
        merged[sh] = list(rows.values())
    return merged, errs, owner.get("user_de", [])


# ---------------------------------------------------------------- kiểm kê
def cmd_ds(goc, dot):
    """kiemke --ds: liệt kê gọn theo thư mục (không đọc nội dung, không ghi file) để model xem phạm vi rồi quyết định bỏ gì vào _bo_qua.txt."""
    import pymupdf
    P = paths(goc, dot)
    if not os.path.isdir(P["inp"]): raise SystemExit(f"Không có thư mục {P['inp']}")
    dirs = {}
    for dp, dn, fn in os.walk(P["inp"]):
        for f in sorted(fn):
            if f.startswith("~$") or os.path.splitext(f)[1].lower() in SKIP_EXT: continue
            dirs.setdefault(os.path.relpath(dp, P["inp"]), []).append(f)
    for d in sorted(dirs):
        fs, ext, mb, pg = dirs[d], {}, 0.0, 0
        for f in fs:
            e = os.path.splitext(f)[1].lower(); ext[e] = ext.get(e, 0) + 1
            full = os.path.join(P["inp"], d, f); mb += os.path.getsize(full) / 1048576
            if e == ".pdf":
                try: pg += pymupdf.open(full).page_count
                except Exception: pass
        print(f"[{d}]  {len(fs)} file · {mb:.1f} MB · " + " ".join(f"{k}×{v}" for k, v in sorted(ext.items())) + (f" · PDF {pg} trang" if pg else ""))
        if len(fs) <= 8: print("    " + " | ".join(fs))
        else: print("    " + " | ".join(fs[:3]) + f" … {fs[-1]}")


def cmd_kiemke(goc, dot):
    import pymupdf
    P = paths(goc, dot)
    if not os.path.isdir(P["inp"]): raise SystemExit(f"Không có thư mục {P['inp']}")
    os.makedirs(os.path.join(P["data"], "_text"), exist_ok=True)
    allf, skipped = [], []
    # L0 riêng dự án: <goc>\01_scan\_bo_qua.txt — mỗi dòng 1 regex (so với đường dẫn tương đối, không phân biệt hoa/thường); dòng bắt đầu # là ghi chú
    rf = os.path.join(goc, "01_scan", "_bo_qua.txt"); user_rx = None
    if os.path.isfile(rf):
        pats = [l.strip() for l in open(rf, encoding="utf-8") if l.strip() and not l.lstrip().startswith("#")]
        if pats: user_rx = re.compile("|".join(f"(?:{p})" for p in pats), re.I); print(f"L0 luật riêng dự án: {len(pats)} dòng trong {rf}")
    for dp, dn, fn in os.walk(P["inp"]):
        dn.sort()
        for f in sorted(fn):
            rel = os.path.relpath(os.path.join(dp, f), P["inp"])
            if f.startswith("~$") or os.path.splitext(f)[1].lower() in SKIP_EXT: continue
            if SKIP_NAME.search(f) or SKIP_DIR.search(rel) or (user_rx and user_rx.search(rel)): skipped.append(rel); continue
            allf.append(rel)
    if skipped: print(f"L0 bỏ {len(skipped)} file nhiễu (frame/xref/Archive/log…): " + "; ".join(os.path.basename(s) for s in skipped[:6]) + (" …" if len(skipped) > 6 else ""))
    dup_report(P["inp"], allf)
    FILES, PAGES, npg = [], [], 0
    tok = re.compile(r"\b[A-Z0-9]{1,6}(?:[.\-_][A-Z0-9]{1,6}){2,}\b")
    for i, rel in enumerate(allf, 1):
        fid, ext = f"F{i:03d}", os.path.splitext(rel)[1].lower()
        full = os.path.join(P["inp"], rel)
        row = {"ID": fid, "Loại": loai_file(rel), "Tên file": os.path.basename(rel), "Đường dẫn": rel, "Ext": ext,
               "MB": round(os.path.getsize(full) / 1048576, 1), "Mở": [[os.path.basename(rel), rel]]}
        if ext == ".pdf":
            try:
                d = pymupdf.open(full)
            except Exception as e:
                row["Số trang"] = f"LỖI MỞ: {e}"[:60]; FILES.append(row); continue
            row["Số trang"] = d.page_count; texts = []
            for p in d:
                t = p.get_text(); texts.append(t); n = len(t.strip())
                W, H = p.rect.width, p.rect.height
                so, ten = title_guess(p, tok)
                PAGES.append({"File ID": fid, "Trang": p.number + 1, "Số BV": so, "Tên BV": ten,
                              "Lớp text": "CÓ" if n > 200 else ("ÍT" if n else "KHÔNG"),
                              "Mở": [[f"{fid} p{p.number + 1}", fid, p.number + 1]]})
            json.dump(texts, open(os.path.join(P["data"], "_text", fid + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
            npg += d.page_count
        FILES.append(row)
    dump(P["data"], "FILES.kiemke.json", FILES); dump(P["data"], "PAGES.kiemke.json", PAGES)
    by = {}
    for f in FILES: by[f["Loại"]] = by.get(f["Loại"], 0) + 1
    notext = sum(1 for p in PAGES if p["Lớp text"] == "KHÔNG")
    print(f"{len(FILES)} file, {npg} trang PDF ({notext} trang không có lớp text). Theo loại: {by}")
    print("Loại 'Khác' cần model xem tên/đường dẫn rồi ghi lại vào FILES.kiemke.json nếu đoán sai.")


def dup_report(inp, rels):
    """L1: in nhóm file trùng nội dung (cùng hash) và nhóm cùng tên khác ngày/REV (chỉ báo, không xóa) — model quyết giữ bản nào."""
    import hashlib
    by_size = {}
    for r in rels: by_size.setdefault(os.path.getsize(os.path.join(inp, r)), []).append(r)
    for sz, g in by_size.items():
        if len(g) < 2 or sz == 0: continue
        h = {}
        for r in g: h.setdefault(hashlib.md5(open(os.path.join(inp, r), "rb").read()).hexdigest(), []).append(r)
        for v in h.values():
            if len(v) > 1: print("L1 TRÙNG NỘI DUNG: " + " = ".join(v))
    stem = {}
    for r in rels:
        b, e = os.path.splitext(os.path.basename(r))
        k = re.sub(r"[\s_\-]*rev[\s._-]*\d+\w*|^\d{6}(?!\d)[\s_\-]*|[\s_\-]+\d{6}$", "", b.lower()).strip(" _-")
        stem.setdefault((os.path.dirname(r), k, e.lower()), []).append(r)
    for v in stem.values():
        if len(v) > 1: print("L1 CÙNG TÊN KHÁC NGÀY/REV (giữ bản mới nhất, ghi NOTES): " + " | ".join(os.path.basename(x) for x in v))


RAC = re.compile(r"PERMISSION|COPYRIGHT|WWW\.|@|TƯ VẤN|CONSULTANT|ARCHITECT|DRAWN|CHECKED|APPROVED|SCALE|TỶ LỆ|TỈ LỆ|"
                 r"HẠNG MỤC|ITEM:|BY:|DIRECTOR|DESIGNER|NGÀY|DATE|REV|CHỦ ĐẦU TƯ|CLIENT|PROJECT|DỰ ÁN|SỐ BẢN VẼ|DRAWING NO|TÊN BẢN VẼ|DRAWING TITLE|SHEET|ISSUE", re.I)


def title_guess(p, tok):
    """Đoán số tờ + tên tờ trong khung tên (góc phải-dưới theo chiều nhìn): số = mã gần nhãn DRAWING NO (+ mảnh tiền tố);
    tên = khối chữ to nhất gần nhãn DRAWING NAME. Chỉ là đoán — bước bản vẽ sửa khi sai."""
    import pymupdf
    W, H, M = p.rect.width, p.rect.height, p.rotation_matrix
    bl = []
    for b in p.get_text("dict")["blocks"]:
        if b.get("type") != 0: continue
        r = pymupdf.Rect(b["bbox"]) * M  # toạ độ theo chiều nhìn (trang xoay 90/270)
        if r.x0 < W * 0.6 or r.y0 < H * 0.55: continue
        sp = [s for l in b["lines"] for s in l["spans"]]
        t = "\n".join("".join(s["text"] for s in l["spans"]) for l in b["lines"])
        bl.append((r.x0, r.y0, r.x1, r.y1, t, max((s["size"] for s in sp), default=0)))
    tx = lambda b: " ".join(b[4].split())
    nos = lambda b: b[4].replace(" ", "").replace("\n", "").upper()
    ctr = lambda b: ((b[0] + b[2]) / 2, (b[1] + b[3]) / 2)
    dist = lambda a, b: ((ctr(a)[0] - ctr(b)[0]) ** 2 + (ctr(a)[1] - ctr(b)[1]) ** 2) ** 0.5
    frag = re.compile(r"^[A-Z0-9][A-Z0-9.\-_]*[-.]$")
    lab_no = [b for b in bl if re.search(r"KÝHIỆUBẢNVẼ|DRAWINGNO|SỐBẢNVẼ|SHEETNO", nos(b))]
    lab_ten = [b for b in bl if re.search(r"TÊNBẢNVẼ|DRAWINGNAME|DRAWINGTITLE", nos(b))]
    num = [(b, m) for b in bl for m in tok.findall(b[4]) if not RAC.search(tx(b)) and not frag.match(tx(b))]
    if not num: return "", ""
    nb, so = min(num, key=lambda x: dist(x[0], lab_no[0])) if lab_no else num[-1]
    fr = [b for b in bl if frag.match(tx(b))]
    if fr:
        f = min(fr, key=lambda b: dist(b, nb))
        if dist(f, nb) < W * 0.08: so = tx(f) + so
    anchor = lab_ten[0] if lab_ten else nb
    ok = [b for b in bl if b is not nb and b is not anchor and 6 <= len(tx(b)) <= 120 and not RAC.search(tx(b))
          and not tok.fullmatch(tx(b)) and not frag.match(tx(b)) and re.search(r"[A-ZÀ-Ỹ]{3}", tx(b))
          and not re.search(r"CO\.,? ?LTD|STREET|WARD|TEL[: ]", tx(b), re.I)]
    ok = [b for b in ok if dist(b, anchor) < W * 0.15] or ok
    ok.sort(key=lambda b: (-round(b[5]), dist(b, anchor)))  # tên tờ = chữ to nhất gần nhãn
    return so, " / ".join(tx(b) for b in ok[:2])[:120]


def loai_file(rel):
    s, ext = rel.lower(), os.path.splitext(rel)[1].lower()
    if ext in (".msg", ".eml"): return "Mail"
    if ext in (".dwg", ".dxf", ".dwt"): return "CAD"
    if re.search(r"boq|bill|khoi luong|khối lượng", s) or (ext in (".xlsx", ".xls", ".xlsm") and re.search(r"(^|\\)\d+_", s)): return "BoQ"
    if re.search(r"material|vat lieu|vật liệu|\bml\b", s): return "Material List"
    if re.search(r"spec|chỉ dẫn|chi dan|tieu chuan|tiêu chuẩn|\d\d \d\d \d\d", s): return "Spec"
    if re.search(r"drawing|ban ve|bản vẽ|banve|\\bv\\|mass units|plan", s) or ext == ".pdf": return "Bản vẽ"
    if ext in (".xlsx", ".xls", ".xlsm", ".csv"): return "Excel khác"
    return "Khác"


def dump(d, name, rows):
    os.makedirs(d, exist_ok=True)
    json.dump({"rows": rows}, open(os.path.join(d, name), "w", encoding="utf-8"), ensure_ascii=False, indent=0)


# ---------------------------------------------------------------- công cụ đọc
def cmd_tim(goc, dot, rx, args):
    P = paths(goc, dot)
    files = {r["ID"]: r for _, r in load_json(P["data"])[0].get("FILES", [])}
    fl, pat, n = opt(args, "--file"), re.compile(rx, re.I), 0
    for f in sorted(glob.glob(os.path.join(P["data"], "_text", "F*.json"))):
        fid = os.path.basename(f)[:-5]; name = files.get(fid, {}).get("Tên file", "")
        if fl and fl.lower() not in name.lower(): continue
        for i, t in enumerate(json.load(open(f, encoding="utf-8")), 1):
            hits = sorted(set(m.group(0) for m in pat.finditer(t)))
            if hits:
                n += 1; print(f"{fid} p{i}\t{name[:50]}\t{', '.join(hits[:12])}")
    print(f"-- {n} trang")


def _fold(s):
    import unicodedata
    s = unicodedata.normalize("NFD", str(s or "").lower().replace("đ", "d"))
    return re.sub(r"\s+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn"))


STOP = set("va voi cho cua the loai dang kieu bang theo tren duoi trong ngoai khong co mot cac cai bo and the with for of in on".split())


def ngoai_pham_vi(qs_rows):
    """Mã bước boq đã phân loại ngoài scope nhà máy (5 xưởng BOHO): QUAN_SAT Trường 'Phạm vi nhà máy', Giá trị 'KHÔNG THUỘC SCOPE NHÀ MÁY…'."""
    return {str(r.get("Mã")) for r in qs_rows
            if r.get("Trường") == "Phạm vi nhà máy" and str(r.get("Giá trị") or "").strip().upper().startswith("KHÔNG THUỘC SCOPE")}


def cmd_ungvien(goc, dot, args):
    """Sau bước boq: với mỗi Mã trong BOQ_TONGHOP, liệt kê trang PDF ứng viên (mã xuất hiện, hoặc ≥2 từ khóa của Hạng mục) từ cache text.
    Ghi _data\\<dot>\\_ungvien.txt để subagent bước banve chỉ mở các trang này (mô phỏng cách user: đọc BoQ -> tìm tên tờ -> chi tiết)."""
    P = paths(goc, dot); sheets = load_json(P["data"])[0]
    boq = [r for _, r in sheets.get("BOQ_TONGHOP", [])]
    if not boq: raise SystemExit("Chưa có BOQ_TONGHOP (chạy bước boq trước)")
    ngoai = ngoai_pham_vi([r for _, r in sheets.get("QUAN_SAT", [])])
    boq = [r for r in boq if str(r.get("Mã")) not in ngoai]
    files = {r["ID"]: r for _, r in sheets.get("FILES", [])}
    titles = {(r["File ID"], int(r["Trang"])): r.get("Tên BV") or r.get("Số BV") or "" for _, r in sheets.get("PAGES", []) if r.get("Trang")}
    top = int(opt(args, "--top", 12)); texts = {}
    for f in sorted(glob.glob(os.path.join(P["data"], "_text", "F*.json"))):
        fid = os.path.basename(f)[:-5]
        texts[fid] = [_fold(t) for t in json.load(open(f, encoding="utf-8"))]
    out = []
    for r in boq:
        ma, hm = str(r.get("Mã") or "").strip(), r.get("Hạng mục") or ""
        kws = sorted({w for w in re.findall(r"[a-z0-9]{3,}", _fold(hm)) if w not in STOP})
        mrx = re.compile(r"(?<![a-z0-9])" + re.escape(_fold(ma)) + r"(?![a-z0-9])") if ma else None
        res = []
        for fid, pages in texts.items():
            for i, t in enumerate(pages, 1):
                mh = len(mrx.findall(t)) if mrx else 0
                kh = [w for w in kws if w in t]
                if mh or len(kh) >= min(2, len(kws)) and kh:
                    res.append((mh * 5 + len(kh), fid, i, mh, kh))
        res.sort(key=lambda x: -x[0])
        out.append(f"## {ma} · {hm}   (từ khóa: {', '.join(kws)})")
        if not res: out.append("   KHÔNG CÓ ỨNG VIÊN — tìm theo tên tờ / Drawing List, hoặc thật sự chưa có tờ")
        for sc, fid, i, mh, kh in res[:top]:
            out.append(f"   {fid} p{i}  điểm {sc}  mã×{mh}  [{', '.join(kh)}]  {files.get(fid, {}).get('Tên file', '')[:40]} | {titles.get((fid, i), '')[:50]}")
        if len(res) > top: out.append(f"   … còn {len(res) - top} trang điểm thấp hơn (dùng `tim` nếu cần)")
    txt = "\n".join(out)
    open(os.path.join(P["data"], "_ungvien.txt"), "w", encoding="utf-8").write(txt + "\n")
    print(txt); print(f"-- {len(boq)} mã -> {os.path.join(P['data'], '_ungvien.txt')}; bỏ {len(ngoai)} mã ngoài scope nhà máy")


def cmd_chu(pdf, pg, args):
    import pymupdf
    p = pymupdf.open(pdf)[int(pg) - 1]
    print(p.get_text(clip=parse_clip(args, p), sort=True))


def cmd_so(pdf, pg, args):
    import pymupdf
    p = pymupdf.open(pdf)[int(pg) - 1]
    W, H = p.rect.width, p.rect.height
    for w in p.get_text("words", clip=parse_clip(args, p), sort=True):
        if re.fullmatch(r"[\d.,x×/+\-]+", w[4]) and re.search(r"\d", w[4]):
            print(f"{w[4]}\t x={w[0] / W:.3f} y={w[1] / H:.3f}")


def cmd_anh(pdf, pg, png, args):
    import pymupdf
    p = pymupdf.open(pdf)[int(pg) - 1]
    pix = p.get_pixmap(dpi=int(opt(args, "--dpi", 150)), clip=parse_clip(args, p))
    os.makedirs(os.path.dirname(os.path.abspath(png)), exist_ok=True)
    pix.save(png); print(png, pix.width, "x", pix.height)


GEM_MODEL = "gemini-3.8-flash-low"  # benchmark 24/10: low đủ chính xác (~13 s/crop); medium/high quá chậm
GEM_SO = ("It is a crop of an interior/joinery drawing (dimensions in millimetres). List EVERY dimension number printed along "
          "dimension lines in the crop, including vertical dimension lines, one entry per printed number, in reading order. "
          "Do NOT include detail numbers inside circles, drawing numbers, scale text, level markers (FFL/FCL/AFFL) or note text. "
          "If a number is unreadable write \"?\". Never guess. Output ONLY a JSON array of strings, nothing else.")


GEM_TRUONG = ["Lắp đặt (gắn tường/đứng rời/âm tường)", "Kích thước BV (R×S×C mm, chỉ số in trên bản vẽ)", "Có cánh (số cánh, mở/lùa)",
              "Ngăn kéo (số)", "Thanh treo", "Kệ/đợt (số)", "Hốc/ô tivi", "Mặt đá", "Gương", "Kính", "Kim loại",
              "Giấy/vải", "Mã hoàn thiện trên BV (nhãn WD/PT/LAM/SS… chỉ vào món này)", "Khung xương", "Cong/bo",
              "Phần ngoài phạm vi (chậu, vòi, thiết bị… thấy trên hình)", "Số chi tiết/hình đang xem"]
GEM_QS = ("It is a page or crop of an interior/joinery drawing. Inspect ONLY this item: {vat}. "
          "For each field below, report what is VISIBLE in the image. Output ONLY a JSON object: "
          "{{\"<field>\": {{\"value\": \"…\", \"status\": \"YES|NO|NOT VISIBLE\", \"where\": \"view/detail number + position in the image\"}}}}. "
          "YES = present, NO = clearly absent, NOT VISIBLE = cannot tell from this image. Count only what you can see; never guess, "
          "never infer from typical designs. Fields (keep these exact Vietnamese keys): {fields}")


def agy_exe():
    import shutil as sh
    p = sh.which("agy")
    if p: return p
    g = glob.glob(os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "WinGet", "Packages", "Google.AntigravityCLI*", "**", "agy.exe"), recursive=True)
    if not g: raise SystemExit("Không thấy agy.exe (Antigravity CLI). Bỏ qua kiểm chéo Gemini, ghi vào kết quả trả về.")
    return g[0]


def cmd_gemini(args):
    """Gemini đọc độc lập ảnh crop (người đọc thứ 2). Mặc định: liệt kê số kích thước; --hoi: câu hỏi riêng.
    --quansat "Tủ lavabo (V01)": trả JSON đủ bộ trường quan sát (GEM_TRUONG) để so từng trường với Claude.
    --claude "900,450,450": so với số Claude đọc -> KHỚP / chỉ Claude / chỉ Gemini. Kết quả lưu cạnh ảnh: <ảnh>.gemini.txt"""
    import subprocess, concurrent.futures as cf
    from collections import Counter
    pngs = [a for a in args if a.lower().endswith(".png")]
    hoi, mine, par = opt(args, "--hoi"), opt(args, "--claude"), int(opt(args, "--par", 4))
    vat = opt(args, "--quansat")
    if vat: hoi = GEM_QS.format(vat=vat, fields=json.dumps(GEM_TRUONG, ensure_ascii=False))
    exe = agy_exe()
    def run(png):
        png = os.path.abspath(png)
        out = png + ".gemini.txt"
        if os.path.exists(out) and os.path.getsize(out) > 2 and not hoi: return png, open(out, encoding="utf-8").read()
        prompt = (f"Do NOT run shell commands; use your file-viewing tool directly on this absolute path: {png} . "
                  + ((hoi if vat else hoi + " Answer only from what is visible in the image; if not visible say NOT VISIBLE. Never guess.") if hoi else GEM_SO))
        try:
            r = subprocess.run([exe, "--model", GEM_MODEL, "--print", prompt, "--print-timeout", "280s"], cwd=os.path.dirname(png),
                               capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=320)
            txt = r.stdout.strip() or f"LỖI agy: {r.stderr.strip()[:200]}"
        except subprocess.TimeoutExpired:
            txt = "LỖI agy: quá 320 s"
        if not txt.startswith("LỖI") and not hoi: open(out, "w", encoding="utf-8").write(txt)
        return png, txt
    with cf.ThreadPoolExecutor(par) as ex:
        res = list(ex.map(run, pngs))
    allg = []
    for png, txt in res:
        print(f"== {os.path.basename(png)}\n{txt}")
        if not hoi:
            m = re.search(r"\[.*\]", txt, re.S)
            try: allg += [str(x).strip() for x in json.loads(m.group(0))] if m else []
            except ValueError: print("  (không đọc được JSON của Gemini)")
    if mine and not hoi:
        a = Counter(x.strip() for x in mine.split(",") if x.strip()); g = Counter(x for x in allg if x != "?")
        print(f"-- SO SÁNH: khớp {sum((a & g).values())} · chỉ Claude {sorted((a - g).elements())} · chỉ Gemini {sorted((g - a).elements())}")


def cmd_mail(msg):
    """In tiêu đề, người gửi, người nhận, tệp đính kèm, nội dung của file Outlook .msg."""
    import olefile
    o = olefile.OleFileIO(msg)
    def get(tag, base=""):
        for enc, suf in (("utf-16-le", "001F"), ("cp1252", "001E")):
            p = f"{base}__substg1.0_{tag}{suf}"
            if o.exists(p): return o.openstream(p).read().decode(enc, "replace").strip("\x00")
        return ""
    print("Tiêu đề :", get("0037")); print("Từ      :", get("0C1A"), get("0C1F"))
    print("Tới     :", get("0E04")); print("Cc      :", get("0E03"))
    att = sorted({e[0] for e in o.listdir() if e[0].startswith("__attach")})
    print("Đính kèm:", [get("3707", a + "/") or get("3704", a + "/") for a in att])
    print("-" * 60); print(get("1000"))


def cmd_xem(xlsx, args):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if not args:
        for ws in wb.worksheets: print(f"[{ws.title}] {ws.max_row}x{ws.max_column}")
        return
    ws = wb[args[0]]
    rng = args[1] if len(args) > 1 else f"A1:{openpyxl.utils.get_column_letter(min(ws.max_column or 26, 30))}{min(ws.max_row or 60, 80)}"
    from openpyxl.utils import range_boundaries
    c0, r0, c1, r1 = range_boundaries(rng)
    for i, row in enumerate(ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True), r0):
        v = [f"{openpyxl.utils.get_column_letter(c0 + j)}={str(x)[:60]}" for j, x in enumerate(row) if x not in (None, "")]
        if v: print(i, " | ".join(v))


# ---------------------------------------------------------------- build
def snapshot_rows(ws, hdr, cols):
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v in (None, "") for v in r): continue
        rows.append({h: r[i] for i, h in enumerate(hdr) if i < len(r) and cols[h]["kieu"] not in ("CT", "AUTO", "LINK")})
    return rows


def pull_user(P, sheets, cols):
    """So file xlsx hiện có với ảnh chụp lần build trước -> sửa tay của user ghi vào <SHEET>.user.json."""
    snap_f = os.path.join(P["data"], "_snapshot.json")
    if not (os.path.exists(P["out"]) and os.path.exists(snap_f)): return []
    import openpyxl
    snap, wb, notes = json.load(open(snap_f, encoding="utf-8")), openpyxl.load_workbook(P["out"]), []
    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO" or sh not in wb.sheetnames or sh not in snap: continue
        ws = wb[sh]; hdr = [c.value for c in ws[1]]
        if [h for h in hdr if h] != meta["hdr"]: notes.append(f"{sh}: user đổi cột trong file — không lấy sửa tay sheet này"); continue
        cur, old = snapshot_rows(ws, meta["hdr"], cols[sh]), snap[sh]
        uf = os.path.join(P["data"], f"{sh}.user.json")
        user = json.load(open(uf, encoding="utf-8"))["rows"] if os.path.exists(uf) else []
        if meta["kieu"] == "KHOA":
            k = meta["khoa"]; kk = lambda r: tuple(str(r.get(x)).strip() for x in k)
            om = {kk(r): r for r in old}; um = {kk(r): r for r in user}
            for r in cur:
                o = om.get(kk(r))
                ch = {c: v for c, v in r.items() if c not in k and (o is None or norm(v) != norm(o.get(c))) and not (o is None and v in (None, ""))}
                if ch:
                    um.setdefault(kk(r), {x: r[x] for x in k}).update(ch)
                    notes.append(f"{sh} {'/'.join(kk(r))}: user sửa {list(ch)}")
            user = list(um.values())
        else:
            olds = [json.dumps({c: norm(v) for c, v in r.items()}, ensure_ascii=False, sort_keys=True) for r in old]
            seen = [json.dumps({c: norm(v) for c, v in r.items()}, ensure_ascii=False, sort_keys=True) for r in user]
            for r in cur:
                s = json.dumps({c: norm(v) for c, v in r.items()}, ensure_ascii=False, sort_keys=True)
                if s in olds: olds.remove(s)
                elif s not in seen:
                    user.append({c: v for c, v in r.items() if v not in (None, "")}); seen.append(s)
                    notes.append(f"{sh}: user thêm dòng {str(list(r.values())[:3])[:60]}")
        if user: dump(P["data"], f"{sh}.user.json", user)
    return notes


def norm(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()


def cmd_build(goc, dot):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.hyperlink import Hyperlink
    P = paths(goc, dot)
    lock = os.path.join(os.path.dirname(P["out"]), "~$" + os.path.basename(P["out"]))
    if os.path.exists(lock): raise SystemExit(f"DỪNG: {os.path.basename(P['out'])} đang mở trong Excel. Đóng file rồi chạy lại.")
    wb, sheets, cols = load_form()
    notes = pull_user(P, sheets, cols)
    raw, errs = load_json(P["data"])
    merged, e2, user_de = merge(sheets, cols, raw)
    errs += e2
    files_by_id = {r["ID"]: os.path.join(P["inp"], r["Đường dẫn"]) for r in merged.get("FILES", [])}
    LK = Font(color="0563C1", underline="single")
    REG, snap = [], {}
    pos = {sh: {h: i + 1 for i, h in enumerate(m.get("hdr", []))} for sh, m in sheets.items()}

    def fx(tpl, sh, r):
        def rep(m):
            s, c = (m.group(1).split("!") if "!" in m.group(1) else (sh, m.group(1)))
            col = L(pos[s][c])
            return f"'{s}'!${col}:${col}" if "!" in m.group(1) else f"{col}{r}"
        return "=" + re.sub(r"\{([^{}]+)\}", rep, tpl)

    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO": continue
        ws, hdr, rows = wb[sh], meta["hdr"], merged.get(sh, [])
        if sh == "FILES": auto_file_status(rows, merged, files_by_id)
        for i, row in enumerate(rows, 2):
            for j, h in enumerate(hdr, 1):
                m, v, cell = cols[sh][h], row.get(h), ws.cell(i, j)
                if m["kieu"] == "CT":
                    cell.value = fx(m["ct"], sh, i)
                elif h == "STT":
                    cell.value = i - 1
                elif m["kieu"] == "LINK":
                    text, links = norm_links(v, files_by_id, P["inp"])
                    cell.value = text or None
                    if links:
                        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sh}'!{cell.coordinate}", display=text); cell.font = LK
                        for l in links: REG.append((sh, cell.coordinate) + l)
                else:
                    if isinstance(v, (list, dict)): v = json.dumps(v, ensure_ascii=False)
                    cell.value = v
                    if isinstance(v, str) and v.startswith("="): cell.data_type = "s"
                if m["kieu"] in ("TEXT",) and isinstance(cell.value, str) and len(cell.value) > 40:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        if rows: ws.auto_filter.ref = f"A1:{L(len(hdr))}{len(rows) + 1}"
        snap[sh] = [{h: r.get(h) for h in hdr if cols[sh][h]["kieu"] not in ("CT", "AUTO", "LINK")} for r in rows]
    # _LINKS
    ws, outdir = wb["_LINKS"], os.path.dirname(P["out"])
    for i, (s, c, lab, f, pg, sub) in enumerate(REG, 2):
        rel = os.path.relpath(f, outdir) if f else ""
        ws.append([f"{s}!{c}", s, c, lab, rel, pg, sub, "Mở", 1.0])
        if rel:
            ws.cell(i, 8).hyperlink = rel + (("#" + sub) if sub else ""); ws.cell(i, 8).font = LK
    res = run_checks(P, sheets, cols, merged, REG, errs, None)
    dashboard(wb, sheets, merged, res, dot, notes + [f"Ghi đè: {u}" for u in user_de])
    for s in ("_SHEETS", "_COT", "_HUONG_DAN"):
        if s in wb.sheetnames: wb[s].sheet_state = "hidden"
    if os.path.exists(P["out"]):
        os.makedirs(P["bak"], exist_ok=True)
        shutil.copy2(P["out"], os.path.join(P["bak"], f"PROJECT_DATA_{dot}_{datetime.datetime.now():%y%m%d-%H%M%S}.xlsx"))
    wb.save(P["out"])
    json.dump(snap, open(os.path.join(P["data"], "_snapshot.json"), "w", encoding="utf-8"), ensure_ascii=False, default=str)
    print(f"Đã ghi {P['out']}  ({sum(len(v) for v in merged.values())} dòng, {len(REG)} link)")
    for n in notes + user_de: print("  user:", n)
    report(res)
    return res


def auto_file_status(rows, merged, files_by_id):
    used = {}
    for sh, rs in merged.items():
        if sh in ("FILES", "PAGES"): continue
        for r in rs:
            for v in r.values():
                if isinstance(v, (list, dict)):
                    for l in (v.get("links", []) if isinstance(v, dict) else v):
                        f = l[1] if isinstance(l, (list, tuple)) and len(l) > 1 else (l.get("file") if isinstance(l, dict) else None)
                        if f: used[f] = used.get(f, 0) + 1
    read = {}
    for p in merged.get("PAGES", []):
        if p.get("Đã đọc bằng"): read[p["File ID"]] = read.get(p["File ID"], 0) + 1
    for r in rows:
        n = used.get(r["ID"], 0) + used.get(r["Đường dẫn"], 0)
        pg = r.get("Số trang")
        s = f"đọc {read.get(r['ID'], 0)}/{pg} trang" if isinstance(pg, int) else ""
        r["Trạng thái scan"] = "; ".join(x for x in [s, f"{n} link dẫn tới" if n else "chưa có link dẫn tới"] if x)


# ---------------------------------------------------------------- check
def run_checks(P, sheets, cols, merged, REG, errs, den):
    import pymupdf
    E, W = list(errs), []
    upto = BUOC[:BUOC.index(den) + 1] if den else BUOC
    ngoai = ngoai_pham_vi(merged.get("QUAN_SAT", []))
    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO": continue
        for i, r in enumerate(merged.get(sh, []), 2):
            for h, m in cols[sh].items():
                v = r.get(h)
                empty = v in (None, "") or (isinstance(v, (list, dict)) and not v)
                if str(r.get("Mã")) in ngoai and m["buoc"] and "banve" in m["buoc"]: continue  # mã ngoài scope: không đọc bản vẽ
                if m["batbuoc"] and empty and m["kieu"] not in ("CT", "AUTO") and (not m["buoc"] or any(b in upto for b in m["buoc"])):
                    E.append(f"{sh} dòng {i}: thiếu '{h}'")
                if m["kieu"] == "SO" and not empty and not isinstance(v, (int, float)):
                    E.append(f"{sh} dòng {i}: '{h}' = '{v}' phải là số (không để trong ngoặc kép)")
                if m["kieu"] == "LIST" and not empty and m["ds"] and str(v) not in m["ds"]:
                    E.append(f"{sh} dòng {i}: '{h}' = '{v}' không thuộc {m['ds']}")
    pc = {}
    for s, c, lab, f, pg, sub in REG:
        if not f: E.append(f"Link {s}!{c} '{lab}': thiếu file"); continue
        if not os.path.exists(f): E.append(f"Link {s}!{c} '{lab}': không thấy file {f}"); continue
        if pg and f.lower().endswith(".pdf"):
            if f not in pc:
                try: pc[f] = pymupdf.open(f).page_count
                except Exception: pc[f] = 10 ** 6
            if not str(pg).isdigit() or int(pg) > pc[f] or int(pg) < 1: E.append(f"Link {s}!{c} '{lab}': trang {pg} ngoài file ({pc[f]} trang)")
    T, M, C = merged.get("BOQ_TONGHOP", []), merged.get("BOQ_MAP", []), merged.get("BOQ_CDT", [])
    ma_t, ma_c = {str(r["Mã"]) for r in T}, {str(r.get("Mã")) for r in C if r.get("Mã")}
    if C or T:
        for x in sorted(ma_c - ma_t): E.append(f"Mã {x} có trong BOQ_CDT nhưng thiếu ở BOQ_TONGHOP")
        for x in sorted(ma_t - ma_c): E.append(f"Mã {x} ở BOQ_TONGHOP nhưng không có dòng BOQ_CDT nào")
        mk = {(str(r.get("Mã")), str(r.get("Khu"))) for r in C if r.get("Mã")}
        mm = {(str(r["Mã"]), str(r["Khu"])) for r in M}
        for x in sorted(mk - mm): E.append(f"BOQ_MAP thiếu cặp Mã/Khu {x}")
        for x in sorted(mm - mk): E.append(f"BOQ_MAP có cặp Mã/Khu {x} không có trong BOQ_CDT")
    kl = {}
    for r in C:
        try: kl[str(r.get("Mã"))] = kl.get(str(r.get("Mã")), 0) + float(r.get("KL CĐT") or 0)
        except ValueError: E.append(f"BOQ_CDT {r.get('Sheet')}!{r.get('Dòng')}: KL CĐT không phải số '{r.get('KL CĐT')}'")
    bv = {}
    for r in M:
        if r.get("KL theo BV") not in (None, ""):
            try: bv[str(r["Mã"])] = bv.get(str(r["Mã"]), 0) + float(r["KL theo BV"])
            except ValueError: E.append(f"BOQ_MAP {r['Mã']}/{r['Khu']}: KL theo BV không phải số")
    qs = {}
    for r in merged.get("QUAN_SAT", []):
        qs.setdefault(str(r.get("Mã")), []).append(r)
        if r.get("Nguồn") in ("HÌNH", "TEXT") and not r.get("Tờ · Chi tiết"): E.append(f"QUAN_SAT {r.get('Mã')} · {r.get('Trường')}: nguồn {r.get('Nguồn')} mà không ghi tờ/chi tiết")
    if ngoai: W.append(f"{len(ngoai)} mã ngoài scope nhà máy, không đọc bản vẽ: {', '.join(sorted(ngoai)[:12])}{' …' if len(ngoai) > 12 else ''}")
    for r in T:
        ma = str(r["Mã"])
        if ma in ngoai: continue
        if "banve" in upto and ma not in qs: E.append(f"{ma}: chưa có dòng QUAN_SAT nào (bước banve chưa làm mã này)")
        elif "banve" in upto and not any(x.get("Trường") == "Kiểm chéo Gemini" for x in qs[ma]):
            E.append(f"{ma}: thiếu dòng QUAN_SAT 'Kiểm chéo Gemini' (bước banve phải cho Gemini đọc lại các crop)")
        if ma in bv and abs(kl.get(ma, 0) - bv[ma]) > 1e-6: W.append(f"{ma}: KL BoQ {kl.get(ma, 0):g} ≠ KL theo BV {bv[ma]:g}")
        if r.get("Độ tin cậy") == "THAP": W.append(f"{ma}: độ tin cậy THẤP — {str(r.get('Giả định tính giá') or '')[:60]}")
        for h in ("Bản vẽ đối chứng", "Spec · ML"):
            if isinstance(r.get(h), str) and r[h].upper().startswith(("KHÔNG", "CHƯA")): W.append(f"{ma}: {h} = {r[h]}")
    for ma, rs in qs.items():
        kt = [x.get("Trường") for x in rs if x.get("Trạng thái") == "KHÔNG THẤY"]
        mt = [x.get("Trường") for x in rs if x.get("Trạng thái") == "MÂU THUẪN"]
        if kt: W.append(f"{ma}: KHÔNG THẤY {len(kt)} trường ({', '.join(map(str, kt[:5]))})")
        if mt: W.append(f"{ma}: MÂU THUẪN {len(mt)} trường ({', '.join(map(str, mt[:5]))})")
    rfi = {str(r["RFI"]) for r in merged.get("RFI_LOG", [])}
    for r in merged.get("ISSUES", []):
        for x in re.split(r"[,;\s]+", str(r.get("RFI") or "")):
            if x and x not in rfi: E.append(f"ISSUES {r['ID']}: RFI '{x}' không có trong RFI_LOG")
        if r.get("Trạng thái") in (None, "", "MỞ") and r.get("Mức") == "CAO": W.append(f"Issue {r['ID']} ({r.get('Mã')}) mức CAO đang mở: {str(r.get('Nội dung'))[:70]}")
    for r in merged.get("ASK_USER", []):
        if not r.get("Trả lời"): W.append(f"Hỏi user {r['ID']}: {str(r.get('Câu hỏi'))[:90]}")
    nt = [p for p in merged.get("PAGES", []) if p.get("Lớp text") == "KHÔNG" and not p.get("Đã đọc bằng")]
    if nt: W.append(f"{len(nt)} trang PDF không có lớp text chưa được xem bằng hình (PAGES: Lớp text = KHÔNG)")
    return dict(E=E, W=W)


def report(res):
    print(f"KIỂM: {len(res['E'])} LỖI, {len(res['W'])} ngoại lệ cho user")
    for x in res["E"][:60]: print("  LỖI ", x)
    if len(res["E"]) > 60: print(f"  … còn {len(res['E']) - 60} lỗi")
    for x in res["W"][:60]: print("  XEM ", x)


def cmd_check(goc, dot, args):
    P = paths(goc, dot)
    wb, sheets, cols = load_form()
    raw, errs = load_json(P["data"])
    merged, e2, _ = merge(sheets, cols, raw)
    files_by_id = {r["ID"]: os.path.join(P["inp"], r["Đường dẫn"]) for r in merged.get("FILES", [])}
    REG = []
    for sh, rows in merged.items():
        for i, r in enumerate(rows, 2):
            for h, v in r.items():
                if cols.get(sh, {}).get(h, {}).get("kieu") == "LINK":
                    for l in norm_links(v, files_by_id, P["inp"])[1]: REG.append((sh, f"{h}@{i}") + l)
    res = run_checks(P, sheets, cols, merged, REG, errs + e2, opt(args, "--den"))
    report(res)
    sys.exit(1 if res["E"] else 0)


# ---------------------------------------------------------------- dashboard
def dashboard(wb, sheets, merged, res, dot, notes):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink
    ws = wb["DASHBOARD"]
    pj = {str(r.get("Trường")): r.get("Giá trị") for r in merged.get("PROJECT", [])}
    ws["A1"] = f"{pj.get('Tên dự án', '')} · {pj.get('Gói thầu', '')}".strip(" ·") or dot
    ws["A1"].font = Font(bold=True, size=14)
    info = [("Hạn trả giá", pj.get("Hạn trả giá", "chưa có")), ("Đợt hồ sơ", dot),
            ("Hạng mục", len(merged.get("BOQ_TONGHOP", []))), ("Dòng BoQ", len(merged.get("BOQ_CDT", []))),
            ("File / trang", f"{len(merged.get('FILES', []))} / {len(merged.get('PAGES', []))}"),
            ("Quan sát", len(merged.get("QUAN_SAT", []))), ("Issue / RFI", f"{len(merged.get('ISSUES', []))} / {len(merged.get('RFI_LOG', []))}"),
            ("Kiểm máy", f"{len(res['E'])} LỖI · {len(res['W'])} ngoại lệ"), ("Cập nhật", f"{datetime.datetime.now():%d/%m/%Y %H:%M}")]
    for i, (a, b) in enumerate(info, 3):
        ws.cell(i, 1, a).font = Font(bold=True); ws.cell(i, 2, b)
    LK = Font(color="0563C1", underline="single")
    for i, s in enumerate([x for x in wb.sheetnames if x not in ("DASHBOARD", "_SHEETS", "_COT", "_HUONG_DAN")], 3):
        c = ws.cell(i, 4, s); c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{s}'!A1", display=s); c.font = LK
    r0 = len(info) + 5
    ws.cell(r0, 1, "CẦN XEM (máy đã lọc — chỉ xem các dòng này)").font = Font(bold=True, size=12)
    red = PatternFill("solid", fgColor="FCE4D6")
    for k, x in enumerate([("LỖI", e) for e in res["E"]] + [("XEM", w) for w in res["W"]] + [("USER", n) for n in notes], r0 + 1):
        ws.cell(k, 1, x[0]); ws.cell(k, 2, x[1])
        if x[0] == "LỖI": ws.cell(k, 1).fill = red


def cmd_cot(args):
    """In cột mà 1 bước được ghi (lấy từ form) — subagent dùng thay cho việc nhớ cột."""
    _, sheets, cols = load_form()
    b = opt(args, "--buoc")
    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO": continue
        cs = [(h, m) for h, m in cols[sh].items() if (not b or b in m["buoc"]) and m["kieu"] not in ("CT", "AUTO")]
        if not cs: continue
        print(f"\n## {sh}  ({meta['kieu']}{', khóa: ' + ', '.join(meta['khoa']) if meta['khoa'] else ''})  -> file {sh}.{b or '<buoc>'}.json")
        for h, m in cs:
            ds = f" [{' | '.join(m['ds'])}]" if m["ds"] else ""
            print(f"  - {h}{' (bắt buộc)' if m['batbuoc'] else ''} · {m['kieu']}{ds}{' · ' + m['note'] if m.get('note') else ''}")


# ---------------------------------------------------------------- main
if __name__ == "__main__":
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit(0)
    c, r = a[0], a[1:]
    if c == "kiemke": (cmd_ds if "--ds" in r[2:] else cmd_kiemke)(r[0], r[1])
    elif c == "tim": cmd_tim(r[0], r[1], r[2], r[3:])
    elif c == "ungvien": cmd_ungvien(r[0], r[1], r[2:])
    elif c == "chu": cmd_chu(r[0], r[1], r[2:])
    elif c == "so": cmd_so(r[0], r[1], r[2:])
    elif c == "anh": cmd_anh(r[0], r[1], r[2], r[3:])
    elif c == "xem": cmd_xem(r[0], r[1:])
    elif c == "mail": cmd_mail(r[0])
    elif c == "gemini": cmd_gemini(r)
    elif c == "build": sys.exit(1 if cmd_build(r[0], r[1])["E"] else 0)
    elif c == "check": cmd_check(r[0], r[1], r[2:])
    elif c == "cot": cmd_cot(r)
    else: print(__doc__); sys.exit(2)
