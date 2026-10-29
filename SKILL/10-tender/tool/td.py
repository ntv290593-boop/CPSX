"""td.py — tool cố định của skill 10-tender.

Chạy: py -I td.py <lệnh> ...   (không tham số = xem lệnh)
Model chỉ ghi JSON trong <GOC>\\10_TENDER\\_data\\; tool dựng xlsx, tạo link, kiểm lỗi.
Ưu tiên ghép theo ô: tool (*.<buoc>-tool.json) < LLM (*.<buoc>.json) < user (*.user.json).
"""
import sys, os, re, json, shutil, zipfile, datetime, glob
from collections import OrderedDict, Counter, defaultdict
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "_core"))
from scan_core import chu_ky_file, doc_pdf, trang_trung, chi_muc, bm25_dao, loai_to, toks, ncode, _fold

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")  # thông báo LỖI của sys.exit
HERE = os.path.dirname(os.path.abspath(__file__))
FORM = json.load(open(os.path.join(HERE, "..", "form", "COT.json"), encoding="utf-8"))
SHEETS = FORM["sheet"]
STEPS = FORM["thu_tu_buoc"]
TODAY = datetime.date.today()

LENH = """Lệnh:
  init <GOC> --input <rel> --boq <rel trong input> --tcq <rel> [--takenote <rel>] [--du-an SSP]
  kiemke <GOC>                  FILE_INDEX.kiemke-tool.json (liệt kê cả zip, ID giữ nguyên khi chạy lại)
  boq <GOC>                     SCOPE.boq-tool.json + _boq_map.json
  tcq <GOC>                     TCQ_LOG.tcq-tool.json (file TCQ gộp + bản dịch takenote)
  dwg <GOC>                     DWG.dwg-tool.json (số BV, bản hiệu lực, bản cũ)
  mat <GOC>                     MAT.mat-tool.json (mã VL gặp trong BOQ/TCQ)
  doichieu <GOC>                SCOPE.doichieu-tool.json (nháp Mã TCQ/Trạng thái từ TCQ_LOG)
  thaydoi <GOC>                 sau kiemke đợt mới: file mới/đổi/mất + dòng SCOPE/MAT/TCQ bị ảnh hưởng
  trang <GOC>                   đọc chữ + khung tên mọi trang PDF được dùng -> _data\\_text (chỉ file mới/đổi)
  tim <GOC> <regex> [--file s] [--max n]   trang có chữ khớp regex
  ungvien <GOC> [--phong SK-01 | --sheet S | --khoa K] [--top 6] [--ra f.txt] [--do]
                                mỗi dòng SCOPE -> trang PDF ứng viên (lõi scan: BM25 + mã + phòng); --do = đo trùng với LLM
  zip <GOC> <FID> [--ra <dir>]  giải nén 1 file trong zip ra _data\\_zip để đọc
  cot [--buoc <buoc>]           cột bước được ghi
  loc <GOC> <SHEET> [cột=gt | cột!=gt | cột~regex | cột=] ... [--cot a,b] [--max n]
  check <GOC> [--den <buoc>]
  build <GOC> [--ra <dir>]      03_boq\\<DA>_BOQ_TRACK_<yymmdd>.xlsx + 01_index\\FILE_INDEX.xlsx + 04_out\\<DA>_BOQ_<yymmdd>_CDT.xlsx"""

BO_TEN = re.compile(r"^(~\$|\.~)|^(plot\.log|thumbs\.db|desktop\.ini)$", re.I)
BO_DUOI = {".dwl", ".dwl2", ".bak", ".log", ".tmp", ".lck"}
RX_DWG = r"(?<![A-Z0-9])([A-Z]{1,4}(?:-[A-Z0-9]{1,5}){1,3}-\d{3}[A-Z]?)(?![0-9])"
RX_VL = re.compile(r"(?<![\w-])([A-Z]{2,3}-\d{2}[A-Z]?)(?![\w-])")
RX_FFE = re.compile(r"(?<![\w-])(\d{2}-\d{3}[A-Z]?)(?![\w-])")
RX_ROOM = re.compile(r"^([A-Z]{2,3}-\d{2}[A-Z]?)")
RX_URL = re.compile(r"https?://\S+")
RX_NGU_PHONG = re.compile(r"\b(rooms?|units?|suites?|room type|căn|phòng|loại phòng)\b", re.I)


def ma_vl(text, rc, tien_to):
    """Mã VL trong chữ. Mã trùng mã phòng (SOFITEL: ST-01 = phòng Standard Twin = đá Black Forest) chỉ tính là VL khi
    họ mã có mã VL khác (tien_to: tiền tố các mã VL không trùng phòng, vd ST từ ST-02…) và 25 ký tự quanh nó không nói là phòng."""
    out = set()
    for m in RX_VL.finditer(text):
        c = m.group(1)
        if c in rc and (c.split("-")[0] not in tien_to or RX_NGU_PHONG.search(text[max(0, m.start() - 25):m.end() + 25])):
            continue
        out.add(c)
    return sorted(out)


def phong_of(v):
    """'SK-01:12 | ST-01:4' hoặc 'SK-02; PS-01' -> {mã phòng}."""
    return {m.group(1) for x in re.split(r"[|;,\n]", s(_txt(v))) for m in [RX_ROOM.match(x.strip())] if m}


# ---------------------------------------------------------------- chung
def opt(args, name, default=None):
    if name in args:
        i = args.index(name)
        v = args[i + 1] if i + 1 < len(args) else default
        del args[i:i + 2]
        return v
    return default


class P:
    def __init__(self, goc):
        self.goc = os.path.abspath(goc)
        self.t = os.path.join(self.goc, "10_TENDER")
        self.d = os.path.join(self.t, "_data")
        f = os.path.join(self.d, "_cauhinh.json")
        self.cfg = json.load(open(f, encoding="utf-8")) if os.path.exists(f) else None
        self.inp = os.path.join(self.goc, self.cfg["input"]) if self.cfg else None

    def need(self):
        if not self.cfg:
            sys.exit(f"LỖI: chưa có {self.d}\\_cauhinh.json — chạy init trước.")
        return self

    def f(self, name):
        return os.path.join(self.d, name)


def jload(path, default=None):
    if not os.path.exists(path):
        return default
    return json.load(open(path, encoding="utf-8"))


def jdump(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=1)


def write_tool(p, sheet, buoc, rows):
    path = p.f(f"{sheet}.{buoc}-tool.json")
    rows = [{c: v for c, v in r.items() if v not in (None, "")} for r in rows]
    jdump(path, {"rows": rows})
    print(f"Đã ghi {path} ({len(rows)} dòng)")


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def num(v):
    if isinstance(v, (int, float)):
        return round(v, 3)
    try:
        return round(float(str(v).replace(",", "")), 3)
    except (TypeError, ValueError):
        return None


def ymd(y, m, d):
    try:
        return datetime.date(int(y), int(m), int(d)).strftime("%Y-%m-%d")
    except ValueError:
        return None


def date_in(parts):
    """Ngày issue: yyyymmdd hoặc tiền tố yymmdd_ trong tên thư mục/file, ưu tiên phần trong cùng."""
    for part in reversed(parts):
        for m in re.finditer(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", part):
            d = ymd(*m.groups())
            if d:
                return d
        m = re.match(r"^(\d{2})(\d{2})(\d{2})[_ \-]", part)
        if m:
            d = ymd("20" + m.group(1), m.group(2), m.group(3))
            if d:
                return d
    return None


def dot_of(rel):
    first = rel.replace("\\", "/").split("/")[0]
    if "/" not in rel.replace("\\", "/"):
        return "(gốc)", None, 0
    m = re.match(r"^(\d+)\.(\d{2})(\d{2})(\d{2})", first)
    if m:
        return first, ymd("20" + m.group(2), m.group(3), m.group(4)), int(m.group(1))
    return first, None, 999


def split_list(v):
    if v is None:
        return []
    if isinstance(v, list):
        return [s(x) for x in v if s(x)]
    return [x.strip() for x in re.split(r"[;,\n]", str(v)) if x.strip()]


# ---------------------------------------------------------------- ghép JSON
def file_buoc(fn):
    """'SCOPE.banve-2.json' -> ('SCOPE', 'banve-2', 'banve', prio)."""
    base = fn[:-5]
    if "." not in base:
        return None
    sheet, buoc = base.split(".", 1)
    if buoc == "user":
        return sheet, buoc, "user", 2
    if buoc.endswith("-tool"):
        return sheet, buoc, buoc, 0
    return sheet, buoc, re.sub(r"-\d+$", "", buoc), 1


def data_files(p):
    out = []
    for path in glob.glob(os.path.join(p.d, "*.json")):
        fn = os.path.basename(path)
        if fn.startswith("_"):
            continue
        fb = file_buoc(fn)
        if fb:
            out.append((fb[3], fn, path) + fb[:3])
    out.sort()
    return out


def row_match(row, loc):
    for c, rx in loc.items():
        if not re.search(rx, s(_txt(row.get(c))), re.I):
            return False
    return True


def merge(p, only_tool=False):
    """-> (res, src): res[sheet] = OrderedDict(key->row) | list; src[sheet][key][col] = buoc."""
    res = {sh: (OrderedDict() if d["loai"] == "KHOA" else []) for sh, d in SHEETS.items()}
    src = {sh: defaultdict(dict) for sh in SHEETS}
    for prio, fn, path, sheet, buoc, base in data_files(p):
        if sheet not in SHEETS or (only_tool and prio > 0):
            continue
        rows = (jload(path) or {}).get("rows", [])
        d = SHEETS[sheet]
        if d["loai"] == "NOI":
            for r in rows:
                r = dict(r)
                r.setdefault("Bước", buoc)
                res[sheet].append(r)
            continue
        k = d["khoa"]
        for r in sorted(rows, key=lambda r: 0 if "_loc" in r else 1):
            vals = {c: v for c, v in r.items() if c != "_loc" and v is not None}
            if "_loc" in r:
                keys = [kk for kk, row in res[sheet].items() if row_match(row, r["_loc"])]
            else:
                keys = [s(r.get(k))]
            for kk in keys:
                if not kk:
                    continue
                row = res[sheet].setdefault(kk, {k: kk})
                for c, v in vals.items():
                    if c == k:
                        continue
                    if prio >= 1 and v == "(xóa)":  # LLM/user xóa ô (vd nháp tool sai) = để trống
                        row.pop(c, None)
                    else:
                        old = src[sheet][kk].get(c)
                        if prio == 1 and c in row and old and old != buoc and re.sub(r"-\d+$", "", old) == base:
                            v = gop_nhom(d["cot"].get(c, {}).get("kieu"), c, row[c], v)  # nhóm khác cùng bước: gộp
                        row[c] = v
                    src[sheet][kk][c] = buoc
    return res, src


COT_MA = ("Mã VL", "Mã spec/FF&E")


def gop_nhom(kieu, c, a, b):
    """Gộp ô do 2 nhóm cùng bước ghi (dòng BOQ dùng chung nhiều mã phòng): LINK = hợp các link, cột mã = hợp mã."""
    if kieu == "LINK":
        la, lb = links_of(a), links_of(b)
        if not la and not lb:
            return a if s(a) == s(b) else f"{s(a)}; {s(b)}"
        out, seen = [], set()
        for l in la + lb:
            t = tuple(map(str, l))
            if t not in seen:
                seen.add(t)
                out.append(l)
        return out
    if c in COT_MA:
        ma = [m.strip() for x in (a, b) for m in s(x).split(";") if m.strip()]
        return "; ".join(dict.fromkeys(ma))
    return b


# ---------------------------------------------------------------- link / hiển thị
def is_link(x):
    return isinstance(x, list) and len(x) >= 2 and isinstance(x[0], str) and isinstance(x[1], str)


def links_of(v):
    if isinstance(v, dict):
        return v.get("links", [])
    if is_link(v):
        return [v]
    if isinstance(v, list) and v and all(is_link(x) for x in v):
        return v
    return []


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, dict):
        return v.get("text") or "; ".join(l[0] for l in v.get("links", []))
    ls = links_of(v)
    if ls:
        return "; ".join(l[0] for l in ls)
    if isinstance(v, list):
        return "; ".join(s(x) for x in v)
    return s(v) if not isinstance(v, (int, float)) else v


def unzip(p, fid, row, ra=None):
    """Giải nén 1 file trong zip ra _data\\_zip (giữ lại cho lần sau). Lỗi → None."""
    ra = ra or p.f("_zip")
    out = os.path.join(ra, fid + "_" + os.path.basename(row["Trong zip"]))
    if not os.path.exists(out):
        os.makedirs(ra, exist_ok=True)
        try:
            with zipfile.ZipFile(os.path.join(p.inp, row["Đường dẫn"])) as z, open(out + ".tmp", "wb") as fh:
                fh.write(z.read(row["Trong zip"]))
        except (OSError, KeyError, zipfile.BadZipFile):
            return None
        os.replace(out + ".tmp", out)
    return out


def link_item(p, link, fidx):
    """Link JSON [nhãn, FID, trang, ô đích] → (nhãn, file, trang, ô đích). File trong zip → bản giải nén;
    có p.base (thư mục file xlsx) → đường dẫn tương đối. Không mở được → file None."""
    f = link[1]
    row = fidx.get(f)
    if row:
        rel = s(row.get("Đường dẫn"))
        if rel.startswith("http"):
            return link[0], rel, None, None
        path = unzip(p, f, row) if s(row.get("Trong zip")) else os.path.join(p.inp, rel)
    else:
        path = os.path.join(p.inp, f)
    if path and getattr(p, "base", None):
        try:
            path = os.path.relpath(path, p.base)
        except ValueError:  # khác ổ đĩa
            pass
    pg = link[2] if len(link) > 2 and link[2] else None
    sub = link[3] if len(link) > 3 and link[3] else None
    return link[0], path, pg, sub


# ---------------------------------------------------------------- init
README_DA = """# {du_an} — TENDER (skill 10-tender)

Đọc file này + `01_index\\FILE_INDEX.xlsx` trước khi làm bất cứ việc gì trong `10_TENDER`.

## Đầu vào (chỉ đọc)
- Thư mục hồ sơ: `{input}`
- BOQ CĐT: `{boq}`
- TCQ gộp mới nhất: `{tcq}`
- Bản dịch nội bộ: `{takenote}`

## Đợt hồ sơ hiệu lực
(bước kiemke ghi: danh sách đợt theo thứ tự, đợt nào thay thế gì)

## Thứ tự ưu tiên khi mâu thuẫn
(theo hồ sơ dự án — ghi kèm nguồn, vd "Spec → Bản vẽ → BOQ (TCQ04-3)")

## Bẫy đã biết
(vd: số Area bộ IDD khác bộ FF&E/thảm → luôn gắn theo mã phòng)

## Việc chưa quyết với user
"""


def cmd_init(goc, args):
    p = P(goc)
    cfg = p.cfg or {}
    for k, o in [("input", "--input"), ("boq", "--boq"), ("tcq", "--tcq"), ("takenote", "--takenote"), ("du_an", "--du-an")]:
        v = opt(args, o)
        if v:
            cfg[k] = v
    cfg.setdefault("du_an", os.path.basename(p.goc))
    cfg.setdefault("boq_bo_sheet", [])
    for k in ("input", "boq", "tcq"):
        if not cfg.get(k):
            sys.exit(f"LỖI: thiếu --{k}")
    inp = os.path.join(p.goc, cfg["input"])
    for k in ("boq", "tcq", "takenote"):
        if cfg.get(k) and not os.path.exists(os.path.join(inp, cfg[k])):
            sys.exit(f"LỖI: không thấy {k}: {os.path.join(inp, cfg[k])}")
    for sub in ("01_index", "02_extract", "03_boq", "04_out", "_data"):
        os.makedirs(os.path.join(p.t, sub), exist_ok=True)
    jdump(os.path.join(p.d, "_cauhinh.json"), cfg)
    rd = os.path.join(p.t, "README.md")
    if not os.path.exists(rd) or os.path.getsize(rd) == 0:
        txt = README_DA.format(**dict({"takenote": "chưa có"}, **cfg))
        open(rd, "w", encoding="utf-8").write(txt)
    print(f"OK {p.t}\n{json.dumps(cfg, ensure_ascii=False, indent=1)}")


# ---------------------------------------------------------------- kiemke
def vai_tro_doan(rel, inner, ext):
    t = (rel + "/" + (inner or "")).lower()
    name = os.path.basename(inner or rel).lower()
    if ext == ".msg":
        return "Mail"
    if ("tender quer" in name or re.search(r"\btcq", name)) and ext in (".xlsx", ".pdf"):
        return "Làm rõ"
    if re.search(r"as-?built", t):
        return "Tham khảo"
    if re.search(r"(^|/)(mep|lighting)\.zip|/(elek|elv|fire|hvac|plum)/", t):
        return "Ngoài scope"
    if re.search(r"project schedule|rules and regulations|logistic|annex", t):
        return "Tham khảo"
    return None


def cmd_kiemke(goc):
    p = P(goc).need()
    old = {}
    for r in (jload(p.f("FILE_INDEX.kiemke-tool.json")) or {}).get("rows", []):
        old[(r["Đường dẫn"], r.get("Trong zip", ""))] = r["ID"]
    n_max = max([int(i[1:]) for i in old.values() if re.match(r"F\d+$", i)] or [0])
    rx_dwg = re.compile(p.cfg.get("dwg_regex", RX_DWG))
    ck_cu = jload(p.f("_chuky.json"))  # {FID: chữ ký} lần kiemke trước; None = lần đầu
    ck_by_key = {}
    items = []
    for root, dirs, files in os.walk(p.inp):
        dirs.sort()
        for fn in sorted(files):
            ext = os.path.splitext(fn)[1].lower()
            if BO_TEN.search(fn) or ext in BO_DUOI:
                continue
            rel = os.path.relpath(os.path.join(root, fn), p.inp).replace("\\", "/")
            items.append((rel, "", ext))
            fid_cu = old.get((rel, ""))
            ck_by_key[(rel, "")] = chu_ky_file(os.path.join(root, fn), (ck_cu or {}).get(fid_cu))
            if ext == ".zip":
                try:
                    for zi in zipfile.ZipFile(os.path.join(root, fn)).infolist():
                        nm = zi.filename
                        bn = os.path.basename(nm.rstrip("/"))
                        e2 = os.path.splitext(bn)[1].lower()
                        if nm.endswith("/") or BO_TEN.search(bn) or e2 in BO_DUOI:
                            continue
                        items.append((rel, nm, e2))
                        ck_by_key[(rel, nm)] = {"size": zi.file_size, "crc": zi.CRC}  # CRC có sẵn trong zip, không cần đọc
                except zipfile.BadZipFile:
                    print(f"XEM: zip hỏng {rel}")
    rows, pdf_set = [], set()
    for rel, inner, ext in items:
        m = rx_dwg.search(os.path.basename(inner or rel).upper())
        if m and ext == ".pdf":
            pdf_set.add((m.group(1), date_in((rel + "/" + inner).split("/"))))
    for rel, inner, ext in items:
        key = (rel, inner)
        if key in old:
            fid = old[key]
        else:
            n_max += 1
            fid = f"F{n_max:04d}"
        dot, nd, _ = dot_of(rel)
        r = {"ID": fid, "Đợt": dot, "Đường dẫn": rel, "Loại": ext.lstrip(".") or "?"}
        if nd:
            r["Ngày đợt"] = nd
        if inner:
            r["Trong zip"] = inner
        m = rx_dwg.search(os.path.basename(inner or rel).upper())
        di = date_in((rel + "/" + inner).split("/"))
        if m:
            r["Số BV"] = m.group(1)
        if di:
            r["Ngày issue"] = di
        vt = vai_tro_doan(rel, inner, ext)
        if ext == ".dwg" and m and (m.group(1), di) in pdf_set:
            vt, r["Ghi chú"] = "Bản cũ/Trùng", "DWG trùng tờ PDF cùng ngày issue"
        if vt:
            r["Vai trò"] = vt
        rows.append(r)
    write_tool(p, "FILE_INDEX", "kiemke", rows)
    ck = {r["ID"]: ck_by_key[(r["Đường dẫn"], r.get("Trong zip", ""))] for r in rows}
    if ck_cu is not None:
        sosanh = lambda a, b: {k: v for k, v in a.items() if k != "mtime"} == {k: v for k, v in b.items() if k != "mtime"}
        td = {"ngay": TODAY.strftime("%Y-%m-%d"), "moi": sorted(set(ck) - set(ck_cu)),
              "doi": sorted(f for f in set(ck) & set(ck_cu) if not sosanh(ck[f], ck_cu[f])),
              "mat": sorted(set(ck_cu) - set(ck))}
        jdump(p.f("_thaydoi.json"), td)
        print(f"\nSo với lần kiemke trước: {len(td['moi'])} file mới, {len(td['doi'])} file đổi nội dung, {len(td['mat'])} file mất -> chạy `thaydoi` để xem dòng bị ảnh hưởng")
    else:
        print("\nLần đầu ghi chữ ký file (_chuky.json) — lần kiemke sau sẽ báo file mới/đổi/mất")
    jdump(p.f("_chuky.json"), ck)
    by_dot = Counter(r["Đợt"] for r in rows)
    print("\nĐợt (số file, kể cả trong zip):")
    for d in sorted(by_dot, key=lambda d: dot_of(d + "/x")[2]):
        print(f"  {by_dot[d]:5}  {d}")
    print("Loại:", ", ".join(f"{k}={v}" for k, v in Counter(r["Loại"] for r in rows).most_common()))
    print("Vai trò tool đoán:", ", ".join(f"{k}={v}" for k, v in Counter(r.get("Vai trò", "(trống)") for r in rows).most_common()))
    trong = Counter()
    for r in rows:
        if "Vai trò" not in r:
            if r.get("Trong zip"):
                trong[r["Đường dẫn"] + " :: " + "/".join(r["Trong zip"].split("/")[:-1][:3])] += 1
            else:
                trong[os.path.dirname(r["Đường dẫn"]) or "(gốc)"] += 1
    print(f"\nChưa có vai trò ({sum(trong.values())} file) theo thư mục — bước kiemke ghi bằng dòng _loc:")
    for k, v in sorted(trong.items()):
        print(f"  {v:5}  {k}")


def cmd_zip(goc, fid, args):
    p = P(goc).need()
    res, _ = merge(p)
    r = res["FILE_INDEX"].get(fid)
    if not r or not s(r.get("Trong zip")):
        sys.exit(f"LỖI: {fid} không phải file trong zip")
    ra = opt(args, "--ra")
    out = unzip(p, fid, r, ra)
    if not out:
        sys.exit(f"LỖI: không giải nén được {fid}")
    print(out)


# ---------------------------------------------------------------- boq
def low(v):
    return re.sub(r"\s+", " ", s(v)).lower()


def boq_header(rows):
    for h, r in enumerate(rows[:25]):
        cells = [low(c) for c in r]
        if "item" in cells and any("description" in c for c in cells):
            return h, cells.index("item")
    return None, None


def cmd_boq(goc):
    import openpyxl
    from openpyxl.utils import get_column_letter as L
    p = P(goc).need()
    wb = openpyxl.load_workbook(os.path.join(p.inp, p.cfg["boq"]), read_only=True, data_only=True)
    bo = set(p.cfg.get("boq_bo_sheet", []))
    out, bmap, room_codes = [], {}, set()
    for ws in wb.worksheets:
        if ws.title in bo:
            print(f"bỏ sheet {ws.title} (_cauhinh boq_bo_sheet)")
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        h, ic = boq_header(rows)
        if h is None:
            print(f"bỏ sheet {ws.title}: không thấy hàng tiêu đề Item/Description")
            continue
        width = max(len(r) for r in rows)
        for r in rows:
            r.extend([None] * (width - len(r)))
        hr, sub, top = rows[h], rows[h + 1] if h + 1 < len(rows) else [None] * width, rows[h - 1] if h else [None] * width
        cells = [low(c) for c in hr]
        sub_l = [low(c) for c in sub]
        unit = next((j for j in range(width) if cells[j] == "unit" or sub_l[j] == "unit"), None)
        qty = next((j for j in range(width) if "quantity" in cells[j]), None)
        remarks = [j for j in range(width) if cells[j] == "remark"]
        ve = next((j for j in range(width) if re.match(r"ve\b", low(top[j]))), None)
        last = max(j for j in range(width) if any(s(x[j]) for x in (top, hr, sub))) + 1
        rooms, grp = [], ""
        if qty is not None and unit is not None:
            for j in range(qty + 1, unit):
                grp = s(hr[j]) or grp
                code = re.sub(r"\s+", " ", s(sub[j]))
                if code:
                    m = RX_ROOM.match(code)
                    if m:
                        room_codes.add(m.group(1))
                    rooms.append([j, code if m else f"{grp} {code}".strip()])
        dend = qty if qty is not None else (unit if unit is not None else ic + 4)
        bmap[ws.title] = {"hang_tieu_de": h + 1, "cot_item": L(ic + 1), "cot_cuoi": last,
                          "kl": L(qty + 1) if qty is not None else None, "dv": L(unit + 1) if unit is not None else None,
                          "phong": [[L(j + 1), lab] for j, lab in rooms], "ghi_chu": [L(j + 1) for j in remarks]}
        vung = nhom = muc = None
        n = 0
        for i in range(h + 2, len(rows)):
            r = rows[i]
            item = s(r[ic])
            texts = [s(x) for x in r[ic + 1:dend] if s(x)]
            kl = num(r[qty]) if qty is not None else None
            if not item and texts and kl is not None and unit is not None and s(r[unit]):
                item = "-"  # dòng có KL + ĐV nhưng không đánh số (vd Provisional sum)
            if item:
                if not texts and kl is None:
                    continue
                full = re.sub(r"\s+", " ", " ".join(texts))
                row = {"Khóa": f"{ws.title}!{i + 1}", "Sheet": ws.title, "Dòng": i + 1, "Item": item,
                       "Mô tả BOQ": full[:600], "_full": full}
                for c, v in (("Vùng", vung), ("Nhóm", nhom), ("Mục phụ", muc)):
                    if v:
                        row[c] = v
                if unit is not None and s(r[unit]):
                    row["ĐV"] = s(r[unit])
                if qty is not None and num(r[qty]) is not None:
                    row["KL BOQ"] = num(r[qty])
                ph = [f"{lab}:{s(num(r[j]))}" for j, lab in rooms if num(r[j])]
                if ph:
                    row["Mã phòng"] = " | ".join(ph)
                gc = []
                for j in remarks:
                    if s(r[j]):
                        gc.append(("VE: " if ve is not None and j >= ve else "Base: ") + s(r[j]))
                if gc:
                    row["Ghi chú BOQ"] = " | ".join(gc)
                ffe = sorted(set(RX_FFE.findall(full)))
                if ffe:
                    row["Mã spec/FF&E"] = "; ".join(ffe)
                out.append(row)
                n += 1
            else:
                c1, c2 = s(r[ic + 1]) if ic + 1 < dend else "", s(r[ic + 2]) if ic + 2 < dend else ""
                c3 = " ".join(s(x) for x in r[ic + 3:dend] if s(x))
                if c1 and not c2 and not c3:
                    vung, nhom, muc = c1, None, None
                elif c2:
                    nhom, muc = c2, None
                elif c3 and len(c3) <= 120:
                    muc = c3
        print(f"{ws.title}: tiêu đề hàng {h + 1}, cột Item {L(ic + 1)}, {n} dòng hạng mục, {len(rooms)} cột phòng/khu, cột cuối {L(last)}")
    # Mã VL tính sau khi biết mọi mã phòng (mã phòng có thể nằm ở sheet sau) và mọi họ mã VL
    tien_to = set(p.cfg.get("vl_tien_to", [])) | {c.split("-")[0] for r in out for c in RX_VL.findall(r["_full"]) if c not in room_codes}
    trung = set()
    for r in out:
        vl = ma_vl(r.pop("_full"), room_codes, tien_to)
        trung |= set(vl) & room_codes
        if vl:
            r["Mã VL"] = "; ".join(vl)
    jdump(p.f("_boq_map.json"), {"sheet": bmap, "ma_phong": sorted(room_codes), "vl_tien_to": sorted(tien_to)})
    if trung:
        print(f"XEM: mã trùng mã phòng nhưng tính là VL (họ mã có VL khác, ngữ cảnh không nói phòng): {', '.join(sorted(trung))}")
    write_tool(p, "SCOPE", "boq", out)
    print("Mã phòng:", ", ".join(sorted(room_codes)))


# ---------------------------------------------------------------- tcq
def tcq_parse(path):
    """-> {(dot, stt): {...}} cho mọi sheet có hàng 'No.'"""
    import openpyxl
    from openpyxl.utils import get_column_letter as L
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = OrderedDict()
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        txt = ws.title + " " + " ".join(s(c) for r in rows[:6] for c in r)
        m = re.search(r"No\.?\s*(\d+)", txt, re.I)
        if not m:
            continue
        dot = int(m.group(1))
        h = next((i for i, r in enumerate(rows[:20]) if "no." in [low(c) for c in r]), None)
        if h is None:
            continue
        hc = [low(c) for c in rows[h]]
        c_no = hc.index("no.")
        c_q = next((j for j, c in enumerate(hc) if "query" in c), None)
        c_a = next((j for j, c in enumerate(hc) if "response" in c or "clarification" in c), None)
        c_qv = next((j for j, c in enumerate(hc) if "câu hỏi" in c), None)
        c_av = next((j for j, c in enumerate(hc) if "làm rõ" in c or "phản hồi" in c), None)
        ngay = None
        for r in rows[:h]:
            for j, c in enumerate(r):
                if low(c).startswith("date"):
                    ngay = next((s(x)[:10] for x in r[j + 1:] if s(x)), None)
        k_note = 0
        for i in range(h + 1, len(rows)):
            r = rows[i] + [None] * 20
            stt = s(r[c_no])
            q = s(r[c_q]) if c_q is not None else ""
            a = s(r[c_a]) if c_a is not None else ""
            if not stt:
                if not (q or a):
                    continue
                k_note += 1
                stt = f"N{k_note}"
            out[(dot, stt)] = {"dot": dot, "stt": stt, "ngay": ngay, "q": q, "a": a,
                               "qv": s(r[c_qv]) if c_qv is not None else "", "av": s(r[c_av]) if c_av is not None else "",
                               "o": f"'{ws.title}'!{L((c_a if c_a is not None else c_no) + 1)}{i + 1}", "sheet": ws.title, "dong": i + 1}
    return out


def tham_chieu(text):
    refs = []
    for m in re.finditer(r"(?:refer|ref\.?|see|as per|response to)\D{0,25}?\b(?:TCQ|TC)\s*\.?\s*(?:no\.?)?\s*#?\s*0?(\d{1,2})(?:\s*[-/#]\s*(?:item\s*)?(\d{1,3}))?", text, re.I):
        refs.append(f"TCQ{int(m.group(1)):02d}" + (f"-{m.group(2)}" if m.group(2) else ""))
    for m in re.finditer(r"item\s*(?:no\.?\s*)?(\d{1,3})\s*(?:of|in)\s*(?:the\s*)?(?:TCQ|TC)\s*\.?\s*0?(\d{1,2})", text, re.I):
        refs.append(f"TCQ{int(m.group(2)):02d}-{m.group(1)}")
    return list(OrderedDict.fromkeys(refs))


def cmd_tcq(goc):
    p = P(goc).need()
    path = os.path.join(p.inp, p.cfg["tcq"])
    data = tcq_parse(path)
    tn = tcq_parse(os.path.join(p.inp, p.cfg["takenote"])) if p.cfg.get("takenote") else {}
    fidx, _ = merge(p, only_tool=True)
    by_path = {s(r.get("Đường dẫn")): k for k, r in fidx["FILE_INDEX"].items() if not s(r.get("Trong zip"))}
    fid = by_path.get(p.cfg["tcq"].replace("\\", "/"), p.cfg["tcq"])
    rows = []
    for (dot, stt), d in data.items():
        r = {"Mã TCQ": f"TCQ{dot:02d}-{stt}", "Đợt TCQ": dot, "STT": stt,
             "Câu hỏi (EN)": d["q"], "Trả lời (EN)": d["a"],
             "Nguồn": [f"TCQ gộp {d['sheet']}!{d['dong']}", fid, None, d["o"]]}
        if d["ngay"]:
            r["Ngày"] = d["ngay"]
        t = tn.get((dot, stt), {})
        qv, av = d["qv"] or t.get("qv", ""), d["av"] or t.get("av", "")
        if t and t.get("q") and d["q"] and t["q"][:40] != d["q"][:40]:
            qv = av = ""  # takenote lệch câu -> không dùng bản dịch
            print(f"XEM: TCQ{dot:02d}-{stt} câu trong takenote khác file gộp, bỏ bản dịch")
        if qv:
            r["Câu hỏi (VN)"] = qv
        if av:
            r["Trả lời (VN)"] = av
        refs = tham_chieu(d["a"])
        if refs:
            r["Tham chiếu"] = "; ".join(refs)
        urls = RX_URL.findall(d["a"])
        if urls:
            r["Link ngoài"] = "\n".join(urls)
        a = d["a"].lower()
        if not a:
            r["Tình trạng trả lời"] = "Chưa trả lời"
        elif re.search(r"defer|later|to be (informed|confirmed|advised|provided)|\btbc\b|will be (provided|issued|informed)", a):
            r["Tình trạng trả lời"] = "Hoãn/Chờ"
        elif refs and len(a) < 60:
            r["Tình trạng trả lời"] = "Tham chiếu"
        else:
            r["Tình trạng trả lời"] = "Đã trả lời"
        rows.append(r)
    write_tool(p, "TCQ_LOG", "tcq", rows)
    c = Counter(r["Đợt TCQ"] for r in rows)
    print("Câu theo đợt:", ", ".join(f"TCQ{k:02d}={v}" for k, v in sorted(c.items())))
    print("Tình trạng:", ", ".join(f"{k}={v}" for k, v in Counter(r["Tình trạng trả lời"] for r in rows).items()))
    print(f"Chưa có bản dịch VN: {sum(1 for r in rows if 'Trả lời (VN)' not in r)} câu; có link ngoài: {sum(1 for r in rows if 'Link ngoài' in r)}")


# ---------------------------------------------------------------- dwg
def cmd_dwg(goc):
    p = P(goc).need()
    res, _ = merge(p)
    ver = defaultdict(lambda: defaultdict(lambda: {"pdf": [], "dwg": [], "khu": "", "cu": True}))
    for fid, r in res["FILE_INDEX"].items():
        bv = s(r.get("Số BV"))
        if not bv or s(r.get("Vai trò")) in ("Ngoài scope", "Mail"):
            continue
        # phiên bản = (ngày issue thật | "", số đợt, ngày đợt, tên đợt)
        k = (s(r.get("Ngày issue")), dot_of(s(r.get("Đường dẫn")))[2], s(r.get("Ngày đợt")), s(r.get("Đợt")))
        e = s(r.get("Loại"))
        v = ver[bv][k]
        if e in ("pdf", "dwg"):
            v[e].append(fid)
        v["cu"] = v["cu"] and s(r.get("Vai trò")) == "Bản cũ/Trùng"
        parts = (s(r.get("Đường dẫn")) + "/" + s(r.get("Trong zip"))).strip("/").split("/")
        v["khu"] = parts[-2] if len(parts) > 1 else ""
    rows, nghi = [], []
    for bv in sorted(ver):
        vs = ver[bv]
        # bản LLM đánh "Bản cũ/Trùng" xếp sau. Mọi bản đều có ngày issue thật -> so ngày issue; thiếu -> đợt sau đè đợt trước
        # (không trộn ngày issue của bản này với ngày đợt của bản kia)
        theo_ngay = lambda k: (not vs[k]["cu"], k[0], k[1], k[2])
        theo_dot = lambda k: (not vs[k]["cu"], k[1], k[0] or k[2], k[3])
        ks = sorted(vs, key=theo_ngay if all(k[0] for k in vs) else theo_dot)
        k = ks[-1]
        v = vs[k]
        if len(ks) > 1 and max(vs, key=theo_ngay) != max(vs, key=theo_dot):
            nghi.append(f"{bv}: bản đợt sau ({k[3]}) có ngày issue {k[0] or '?'} cũ hơn/không rõ so với bản khác")
        r = {"Số BV": bv, "Khu (thư mục)": v["khu"], "Đợt hiệu lực": k[3], "Số bản": len(ks), "Trạng thái": "Hiệu lực"}
        if k[0] or k[2]:
            r["Ngày issue hiệu lực"] = k[0] or k[2]
        if v["pdf"]:
            r["File PDF"] = [[f"{bv} PDF", f] for f in v["pdf"]]
        if v["dwg"]:
            r["File DWG"] = [[f"{bv} DWG", f] for f in v["dwg"]]
        if len(ks) > 1:
            r["Bản cũ bị thay"] = "; ".join(f"{x[0] or x[2] or '?'} ({x[3]})" for x in ks[:-1])
        rows.append(r)
    write_tool(p, "DWG", "dwg", rows)
    for x in nghi:
        print("XEM:", x)
    print(f"Có nhiều bản: {sum(1 for r in rows if r['Số bản'] > 1)} số BV; không có PDF: {sum(1 for r in rows if 'File PDF' not in r)}")
    print("Khu (thư mục) — bước dwg lập sheet KHU (khu -> mã phòng):")
    for k, n in sorted(Counter(r["Khu (thư mục)"] for r in rows).items()):
        print(f"  {n:4}  {k}")


# ---------------------------------------------------------------- mat
def cmd_mat(goc):
    p = P(goc).need()
    res, _ = merge(p)
    bm = jload(p.f("_boq_map.json")) or {}
    rc, tien_to = set(bm.get("ma_phong", [])), set(bm.get("vl_tien_to", []))
    boq, tcq = defaultdict(list), defaultdict(list)
    for k, r in res["SCOPE"].items():
        for m in split_list(r.get("Mã VL")):
            boq[m].append(k)
    for k, r in res["TCQ_LOG"].items():
        t = s(r.get("Câu hỏi (EN)")) + " " + s(r.get("Trả lời (EN)"))
        for m in ma_vl(t, rc, tien_to):
            tcq[m].append(k)
    rows = []
    for m in sorted(set(boq) | set(tcq)):
        r = {"Mã VL": m, "Số dòng BOQ": len(boq[m])}
        if boq[m]:
            r["Dòng BOQ (mẫu)"] = "; ".join(boq[m][:4]) + (" …" if len(boq[m]) > 4 else "")
        if tcq[m]:
            r["TCQ nhắc tới"] = "; ".join(tcq[m])
        rows.append(r)
    write_tool(p, "MAT", "mat", rows)
    print(f"Mã chỉ có trong TCQ (không thấy trong mô tả BOQ): {', '.join(m for m in sorted(tcq) if not boq[m]) or '—'}")


# ---------------------------------------------------------------- doichieu (nháp)
MAP_TT = {"Loại trừ": "Loại trừ", "Omit": "Omit", "VE": "VE", "Chờ": "Chờ", "Bao gồm": "Base"}


def tcq_chain(tcq):
    """Mã TCQ -> list mã cuối (còn hiệu lực) theo 'Bị thay bởi', đi hết mọi nhánh (1 câu có thể bị 2 câu thay từng phần).
    [k] = còn hiệu lực; None = chuỗi có vòng."""
    def la(k, seen):
        nxt = [x for x in split_list(tcq.get(k, {}).get("Bị thay bởi")) if x in tcq]
        if not nxt:
            return [k]
        res = []
        for x in nxt:
            r = None if x in seen else la(x, seen | {x})
            if r is None:
                return None
            res += r
        return list(dict.fromkeys(res))
    return {k: la(k, {k}) for k in tcq}


def cmd_doichieu(goc):
    p = P(goc).need()
    res, _ = merge(p)
    tcq, scope = res["TCQ_LOG"], res["SCOPE"]
    ma, tt = defaultdict(list), defaultdict(set)
    n_bo = 0
    for k, r in tcq.items():
        if split_list(r.get("Bị thay bởi")):
            continue
        tp = phong_of(r.get("Mã phòng/Khu"))
        for key in split_list(r.get("Dòng BOQ")):
            if key in scope:
                ma[key].append(k)
                t = MAP_TT.get(s(r.get("Trạng thái phạm vi")))
                if t:
                    tt[key].add(t)
        vls = set(split_list(r.get("Mã VL")))
        if vls:
            for key, sr in scope.items():
                if vls & set(split_list(sr.get("Mã VL"))) and k not in ma[key]:
                    sp = phong_of(sr.get("Mã phòng"))
                    if tp and sp and not tp & sp:  # TCQ chỉ nói phòng khác -> không nối chỉ vì trùng mã VL
                        n_bo += 1
                        continue
                    ma[key].append(k)
    rows = []
    for key in scope:
        r = {"Khóa": key}
        if ma[key]:
            r["Mã TCQ"] = "; ".join(ma[key])
        if len(tt[key]) == 1:
            r["Trạng thái"] = tt[key].pop()
        if len(r) > 1:
            rows.append(r)
    write_tool(p, "SCOPE", "doichieu", rows)
    print(f"Bỏ {n_bo} cặp TCQ–dòng trùng Mã VL nhưng khác mã phòng")
    print(f"Dòng có Mã TCQ: {sum(1 for r in rows if 'Mã TCQ' in r)}; có Trạng thái nháp:{sum(1 for r in rows if 'Trạng thái' in r)} / {len(scope)} dòng BOQ")


# ---------------------------------------------------------------- trang / tim / ungvien / thaydoi (lõi scan: SKILL\_core\scan_core.py)
VT_DOC = ("Chính", "Bổ sung/Thay thế", "Làm rõ", "")  # vai trò được đọc chữ + tìm trang
NHOM_TENDER = [("CHI TIẾT", 6, ("chi_tiet", "mat_dung", "khac", "legend")), ("HOÀN THIỆN", 3, ("hoan_thien",)),
               ("MẶT BẰNG", 3, ("mat_bang",)), ("TRẦN/MEP", 2, ("mep",)), ("THỐNG KÊ", 3, ("thong_ke",)), ("DANH MỤC", 1, ("danh_muc",)),
               ("TÀI LIỆU", 2, ("tai_lieu",))]


def pdf_rows(res):
    return {fid: r for fid, r in res["FILE_INDEX"].items()
            if s(r.get("Loại")) == "pdf" and s(r.get("Vai trò")) in VT_DOC and not s(r.get("Đường dẫn")).startswith("http")}


def cmd_trang(goc):
    """Đọc chữ + khung tên mọi trang PDF được dùng -> _data\\_text\\<FID>.json. Chỉ đọc lại file mới/đổi (theo chữ ký của kiemke)."""
    p = P(goc).need()
    ck = jload(p.f("_chuky.json"))
    if not ck:
        sys.exit("LỖI: chưa có _chuky.json — chạy kiemke trước.")
    res, _ = merge(p)
    rows = pdf_rows(res)
    d = p.f("_text")
    os.makedirs(d, exist_ok=True)
    n_doc = n_trang = 0
    loi, zs = [], {}
    for fid, r in sorted(rows.items()):
        f = os.path.join(d, fid + ".json")
        cu = jload(f)
        if cu and cu.get("ck") == ck.get(fid):
            continue
        try:
            if s(r.get("Trong zip")):
                zp = os.path.join(p.inp, r["Đường dẫn"])
                z = zs[zp] if zp in zs else zs.setdefault(zp, zipfile.ZipFile(zp))
                pages = doc_pdf(data=z.read(r["Trong zip"]))
            else:
                pages = doc_pdf(os.path.join(p.inp, r["Đường dẫn"]))
        except Exception as e:
            loi.append(f"{fid}: {e}"[:150])
            continue
        jdump(f, {"ck": ck.get(fid), "pages": pages})
        n_doc += 1
        n_trang += len(pages)
        if n_doc % 25 == 0:
            print(f"  … {n_doc} file, {n_trang} trang", flush=True)
    for z in zs.values():
        z.close()
    tong = Counter()
    for fid in rows:
        c = jload(os.path.join(d, fid + ".json"))
        for x in (c or {}).get("pages", []):
            tong[x["lop"]] += 1
    print(f"Đọc mới {n_doc} file ({n_trang} trang); cache có {sum(tong.values())} trang / {len(rows)} file PDF được dùng. Lớp chữ: {dict(tong)}")
    for x in loi:
        print("XEM: không đọc được", x)


def nap_trang(p, res):
    """-> (texts {fid: [chữ]}, titles {(fid, tr): 'Số BV tên tờ'}). Tờ 1 trang có Số BV: tên tờ lấy từ DWG (Drawing Register) thay khung tên đoán."""
    texts, titles, bv_pages = {}, {}, set()
    for fid, r in pdf_rows(res).items():
        c = jload(p.f(os.path.join("_text", fid + ".json")))
        if not c:
            continue
        texts[fid] = [x["text"] for x in c["pages"]]
        bv = s(r.get("Số BV"))
        ten = s(res["DWG"].get(bv, {}).get("Tên tờ")) if bv else ""
        for i, x in enumerate(c["pages"], 1):
            titles[(fid, i)] = f"{bv} {ten}".strip() if bv and ten and len(c["pages"]) == 1 else (f'{x["so"]} {x["ten"]}'.strip() or bv)
            if bv or re.search(RX_DWG, s(x["so"]).upper()):
                bv_pages.add((fid, i))
    return texts, titles, bv_pages


def chi_muc_tender(p, res):
    texts, titles, bv_pages = nap_trang(p, res)
    if not texts:
        sys.exit("LỖI: chưa có chữ trang PDF — chạy `trang` trước.")
    dup, tong = trang_trung(texts)
    ix = chi_muc(texts, titles)
    ix.update(texts=texts, titles=titles, dup=dup, tong=set(tong), bv_pages=bv_pages,
              loai={pg: loai_to(titles.get(pg, "")) if pg in bv_pages else "tai_lieu" for pg in ix["pages"]})
    cdao, ffe = defaultdict(dict), defaultdict(dict)
    for pg in ix["pages"]:
        for c, n in ix["codes"][pg].items():
            cdao[c][pg] = n
        for m, n in Counter(RX_FFE.findall(texts[pg[0]][pg[1] - 1])).items():
            ffe[m][pg] = n
    ix["cdao"], ix["ffe"] = cdao, ffe
    fidx, dwg = res["FILE_INDEX"], res["DWG"]
    ix["bv_file"] = {bv: {x[1] for x in links_of(r.get("File PDF"))} for bv, r in dwg.items()}
    hl = set().union(*ix["bv_file"].values()) if ix["bv_file"] else set()
    ix["phong_trang"] = defaultdict(set)
    ix["bo_file"] = set()  # bản cũ của 1 Số BV, hoặc tờ DWG Ngoài scope/Tham khảo
    ix["file_phong"] = {}
    for fid in texts:
        bv = s(fidx[fid].get("Số BV"))
        if bv in dwg and (fid not in hl or s(dwg[bv].get("Trạng thái")) in ("Ngoài scope", "Tham khảo")):
            ix["bo_file"].add(fid)
        # mã phòng của file: DWG.Mã phòng, không có thì KHU của thư mục khu (theo DWG hoặc theo tên thư mục)
        mp = phong_of(dwg.get(bv, {}).get("Mã phòng")) if bv in dwg else set()
        if not mp:
            k = s(dwg.get(bv, {}).get("Khu (thư mục)")) if bv in dwg else ""
            parts = (s(fidx[fid].get("Đường dẫn")) + "/" + s(fidx[fid].get("Trong zip"))).split("/")[:-1]
            for kk in [k] + parts[::-1]:
                if kk in res["KHU"]:
                    mp = phong_of(res["KHU"][kk].get("Mã phòng"))
                    break
        ix["file_phong"][fid] = mp
        for m in mp:
            ix["phong_trang"][m] |= {(fid, i) for i in range(1, len(texts[fid]) + 1) if (fid, i) in bv_pages}
    return ix


def rank_dong(ix, r):
    """-> [(điểm, (fid, tr), số lần mã)] giảm dần cho 1 dòng SCOPE. Tín hiệu như pd.py rank_ma, chỉnh cho mô tả BOQ dài:
    BM25(Nhóm + Mục phụ + Mô tả) chuẩn hóa 0..40 trong dòng (trang BM25 cao nhất = 40) + 3×từ trong tên tờ
    + 8×min(mã VL/FF&E, 3) + 10 mã trong tên tờ + 40 khi mô tả nêu đúng Số BV của tờ + 10 tờ cùng mã phòng; ×0.4 tờ phòng khác.
    Bỏ: trang trùng, file tổng, bản cũ, tờ Ngoài scope/Tham khảo. Trang tài liệu (không phải tờ bản vẽ) xếp riêng nhóm TÀI LIỆU."""
    desc = s(r.get("Mô tả BOQ"))
    q = toks(" ".join(s(r.get(c)) for c in ("Nhóm", "Mục phụ")) + " " + desc)
    sc = bm25_dao(ix, q)
    qset = set(q)
    codes = {c for c in (ncode(m) for m in split_list(r.get("Mã VL"))) if c and ix["cdf"].get(c, 0) <= max(8, 0.1 * ix["N"])}
    ffe = [m for m in split_list(r.get("Mã spec/FF&E")) if m in ix["ffe"]]
    bv_nhac = set(re.findall(RX_DWG, desc.upper()))
    cand = set(sc)
    for c in codes:
        cand |= set(ix["cdao"].get(c, {}))
    for m in ffe:
        cand |= set(ix["ffe"][m])
    for bv in bv_nhac:
        cand |= {pg for pg in ix["pages"] if pg[0] in ix["bv_file"].get(bv, ())}
    rp = phong_of(r.get("Mã phòng"))
    for m in rp:  # tờ của phòng dòng này luôn là ứng viên (tờ mặt đứng/FF&E plan chung ít chữ trùng mô tả ngắn)
        cand |= ix["phong_trang"].get(m, set())
    mx = max([v for pg, v in sc.items() if pg in ix["bv_pages"]] or [0]) or 1
    mx_tl = max([v for pg, v in sc.items() if pg not in ix["bv_pages"]] or [0]) or 1
    out = []
    for pg in cand:
        fid = pg[0]
        if pg in ix["dup"] or fid in ix["tong"] or fid in ix["bo_file"]:
            continue
        diem = 40 * sc.get(pg, 0.0) / (mx if pg in ix["bv_pages"] else mx_tl) + 3.0 * sum(1 for w in qset if w in ix["title"][pg])
        mh = sum(ix["codes"][pg].get(c, 0) for c in codes) + sum(ix["ffe"][m].get(pg, 0) for m in ffe)
        diem += 8 * min(mh, 3) + (10 if codes & ix["tit_codes"][pg] else 0)
        if any(fid in ix["bv_file"].get(bv, ()) for bv in bv_nhac):
            diem += 40
        fp = ix["file_phong"].get(fid)
        if rp and fp:
            diem = diem + 10 if rp & fp else diem * 0.4
        if diem > 0:
            out.append((diem, pg, mh))
    out.sort(key=lambda x: -x[0])
    return out


def cmd_ungvien(goc, args):
    """Mỗi dòng SCOPE -> trang PDF ứng viên chia nhóm theo loại tờ. Lọc: --phong SK-01 | --sheet <tên> | --khoa <khóa>.
    --do: so với 'Bản vẽ hiệu lực' LLM đã gán (độ trùng @5/@12/@40) — dùng để đo trước/sau khi sửa lõi."""
    p = P(goc).need()
    ph, sh, kh = opt(args, "--phong"), opt(args, "--sheet"), opt(args, "--khoa")
    top, ra, do = int(opt(args, "--top", "6")), opt(args, "--ra"), "--do" in args
    res, _ = merge(p)
    ix = chi_muc_tender(p, res)
    rows = [(k, r) for k, r in res["SCOPE"].items()
            if (not ph or ph in phong_of(r.get("Mã phòng"))) and (not sh or s(r.get("Sheet")) == sh) and (not kh or k == kh)]
    lab = lambda pg: f"{pg[0]} p{pg[1]}"
    out, tk = [], Counter()
    for k, r in rows:
        rk = rank_dong(ix, r)
        if do:
            gold = {(l[1], int(l[2]) if len(l) > 2 and l[2] else 1) for l in links_of(r.get("Bản vẽ hiệu lực"))}
            gold = {g for g in gold if g[0] in ix["texts"]}
            if gold:
                tk["dòng"] += 1
                tk["trang"] += len(gold)
                pos = [pg for _, pg, _ in rk if pg in ix["bv_pages"]]
                for n in (5, 12, 40):
                    tk[f"@{n}"] += len(gold & set(pos[:n]))
                    tk[f"d@{n}"] += bool(gold & set(pos[:n]))
            continue
        out.append(f"## {k} · {s(r.get('Mô tả BOQ'))[:90]}" + (f"   [{', '.join(sorted(phong_of(r.get('Mã phòng'))))}]" if phong_of(r.get("Mã phòng")) else ""))
        if not rk:
            out.append("   KHÔNG CÓ ỨNG VIÊN — tìm theo tên tờ (loc DWG) hoặc `tim`")
            continue
        for ten, n, loais in NHOM_TENDER:
            hit = [x for x in rk if ix["loai"][x[1]] in loais]
            if hit:
                out.append(f"   {ten:10s} " + ", ".join(f"{lab(x[1])} {ix['titles'].get(x[1], '')[:38]}" + (f" ×{x[2]}" if x[2] else "") for x in hit[:n])
                           + (f"  (+{len(hit) - n})" if len(hit) > n else ""))
    if do:
        n = max(tk["trang"], 1)
        print(f"Độ trùng với 'Bản vẽ hiệu lực' LLM đã gán: {tk['dòng']} dòng, {tk['trang']} trang · "
              + " · ".join(f"@{x} {tk[f'@{x}'] * 100 // n}%" for x in (5, 12, 40))
              + " | dòng có ≥1 trang trùng: " + " · ".join(f"@{x} {tk[f'd@{x}'] * 100 // max(tk['dòng'], 1)}%" for x in (5, 12, 40)))
        print(f"Chỉ mục: {ix['N']} trang, {len(ix['dup'])} trang trùng, file tổng {sorted(ix['tong']) or 'không'}, bỏ {len(ix['bo_file'])} file bản cũ/ngoài scope")
        return
    txt = "\n".join(out)
    if ra:
        open(ra, "w", encoding="utf-8").write(txt + "\n")
        print(f"{len(rows)} dòng -> {ra}")
    else:
        print(txt)


def cmd_tim(goc, args):
    """Tìm trang có chữ khớp regex (trên cache `trang`). --file <chuỗi trong tên file> --max n."""
    p = P(goc).need()
    fl, mx = opt(args, "--file"), int(opt(args, "--max", "100"))
    if not args:
        sys.exit("LỖI: tim <GOC> <regex>")
    rx = re.compile(args[0], re.I)
    res, _ = merge(p)
    texts, titles, _ = nap_trang(p, res)
    n = 0
    for fid in sorted(texts):
        fr = res["FILE_INDEX"][fid]
        ten = os.path.basename(s(fr.get("Trong zip")) or s(fr.get("Đường dẫn")))
        if fl and fl.lower() not in ten.lower():
            continue
        for i, t in enumerate(texts[fid], 1):
            hits = sorted(set(m.group(0) for m in rx.finditer(t)))
            if hits:
                n += 1
                if n <= mx:
                    print(f"{fid} p{i}\t{titles.get((fid, i), '')[:50]}\t{', '.join(hits[:12])}")
    print(f"-- {n} trang")


def cmd_thaydoi(goc):
    """Sau kiemke (+ dwg) của đợt mới: file mới/đổi/mất và dòng SCOPE/MAT/TCQ bị ảnh hưởng -> giao lại đúng phần cho bước banve/mat/doichieu."""
    p = P(goc).need()
    td = jload(p.f("_thaydoi.json"))
    if not td:
        sys.exit("Chưa có _thaydoi.json — kiemke phải chạy ít nhất 2 lần (lần đầu chỉ ghi chữ ký).")
    res, _ = merge(p)
    fidx = res["FILE_INDEX"]
    doi, moi = set(td["doi"]) | set(td["mat"]), set(td["moi"])
    print(f"kiemke {td['ngay']}: {len(moi)} file mới · {len(td['doi'])} đổi nội dung · {len(td['mat'])} mất")
    for nhom, ids in (("MỚI", sorted(moi)), ("ĐỔI", td["doi"]), ("MẤT", td["mat"])):
        if ids:
            vt = Counter(s(fidx.get(f, {}).get("Vai trò")) or "(chưa có vai trò)" for f in ids)
            print(f"  {nhom}: {dict(vt)}")
            for f in ids[:10]:
                r = fidx.get(f, {})
                print(f"     {f} {s(r.get('Đường dẫn'))[-60:]} {s(r.get('Trong zip'))[-50:]}")
    bv_moi = {s(fidx[f].get("Số BV")) for f in moi | set(td["doi"]) if f in fidx and s(fidx[f].get("Số BV"))}
    sc = {}
    for k, r in res["SCOPE"].items():
        ly = []
        for c in ("Bản vẽ hiệu lực", "Nguồn"):
            for l in links_of(r.get(c)):
                if l[1] in doi:
                    ly.append(f"{l[1]} đổi/mất")
                bv = re.findall(RX_DWG, s(l[0]).upper())
                if bv and bv[0] in bv_moi:
                    ly.append(f"{bv[0]} có bản mới")
        if ly:
            sc[k] = sorted(set(ly))
    mat = [k for k, r in res["MAT"].items() if any(l[1] in doi for l in links_of(r.get("Nguồn spec")))]
    tcq = [k for k, r in res["TCQ_LOG"].items() if not s(r.get("Chủ đề"))]
    cu = link_cu(res)
    phong = Counter(m for k in sc for m in phong_of(res["SCOPE"][k].get("Mã phòng")))
    print(f"Số BV có bản mới: {len(bv_moi)}")
    print(f"SCOPE bị ảnh hưởng: {len(sc)} dòng; theo mã phòng: {dict(phong.most_common())}")
    for k in list(sc)[:15]:
        print(f"     {k}: {'; '.join(sc[k])}")
    print(f"MAT có nguồn spec đổi/mất: {mat or '—'}")
    print(f"TCQ chưa phân loại (câu mới): {len(tcq)} {tcq[:12]}")
    print(f"Link bản vẽ trỏ bản cũ (sau khi chạy dwg): {len(cu)}")
    td["anh_huong"] = {"SCOPE": sc, "MAT": mat, "TCQ_LOG": tcq}
    jdump(p.f("_thaydoi.json"), td)


# ---------------------------------------------------------------- cot / loc
def cmd_cot(args):
    b = opt(args, "--buoc")
    for sh, d in SHEETS.items():
        cols = [(c, x) for c, x in d["cot"].items() if not b or b in x["buoc"] or "*" in x["buoc"]]
        if not cols:
            continue
        print(f"\n[{sh}] {d['loai']}" + (f" khóa = {d['khoa']}" if d["loai"] == "KHOA" else ""))
        for c, x in cols:
            ds = f"  ds: {' | '.join(x['ds'])}" if "ds" in x else ""
            bb = "  BẮT BUỘC" if x.get("bat_buoc") in (b, "*") else ""
            print(f"  {c} [{x['kieu']}]{bb}{ds}")


def cmd_loc(goc, sheet, args):
    p = P(goc).need()
    cols = opt(args, "--cot")
    mx = int(opt(args, "--max", "200"))
    res, _ = merge(p)
    data = res[sheet]
    rows = list(data.values()) if isinstance(data, dict) else data
    for f in args:
        m = re.match(r"^(.+?)(!=|=|~)(.*)$", f)
        if not m:
            sys.exit(f"LỖI lọc: {f}")
        c, op, v = m.groups()
        if op == "=":
            rows = [r for r in rows if s(_txt(r.get(c))) == v]
        elif op == "!=":
            rows = [r for r in rows if s(_txt(r.get(c))) != v]
        else:
            rows = [r for r in rows if re.search(v, s(_txt(r.get(c))), re.I)]
    keep = cols.split(",") if cols else None
    print(f"{len(rows)} dòng")
    for r in rows[:mx]:
        print(json.dumps({c: r.get(c) for c in keep} if keep else r, ensure_ascii=False))


# ---------------------------------------------------------------- check
def run_check(p, den=None):
    errs, xem = [], []
    lim = STEPS.index(den) + 1 if den in STEPS else len(STEPS)
    res, src = merge(p)
    fidx = res["FILE_INDEX"]
    tool_keys = {sh: set(merge_keys) for sh, merge_keys in _tool_keys(p).items()}
    rc = set((jload(p.f("_boq_map.json")) or {}).get("ma_phong", []))
    for prio, fn, path, sheet, buoc, base in data_files(p):
        if sheet not in SHEETS:
            errs.append(f"{fn}: sheet lạ '{sheet}' (có: {', '.join(SHEETS)})")
            continue
        if prio == 1 and base not in STEPS:
            errs.append(f"{fn}: bước lạ '{base}'")
            continue
        try:
            rows = json.load(open(path, encoding="utf-8")).get("rows")
            assert isinstance(rows, list)
        except Exception as e:
            errs.append(f"{fn}: JSON hỏng hoặc thiếu 'rows' ({e})")
            continue
        d = SHEETS[sheet]
        for i, r in enumerate(rows):
            tag = f"{fn} dòng {i + 1}"
            for c, v in r.items():
                if c == "_loc":
                    if d["loai"] != "KHOA" or not isinstance(v, dict):
                        errs.append(f"{tag}: _loc chỉ dùng cho sheet KHOA, dạng {{\"cột\": \"regex\"}}")
                        continue
                    try:
                        [re.compile(x) for x in v.values()]
                    except re.error as e:
                        errs.append(f"{tag}: regex _loc lỗi {e}")
                    continue
                if c not in d["cot"]:
                    errs.append(f"{tag}: cột lạ '{c}'")
                    continue
                x = d["cot"][c]
                is_key = d["loai"] == "KHOA" and c == d["khoa"]
                if prio == 0 and buoc not in x["buoc"] and not is_key:
                    errs.append(f"{tag}: tool '{buoc}' không được ghi cột '{c}'")
                if prio == 1 and base not in x["buoc"] and "*" not in x["buoc"] and not is_key:
                    errs.append(f"{tag}: bước '{base}' không được ghi cột '{c}' (được: {', '.join(x['buoc'])})")
                if "build" in x["buoc"] and prio < 2:
                    errs.append(f"{tag}: cột '{c}' do build tự tính")
                if prio >= 1 and v == "(xóa)":
                    continue
                errs.extend(f"{tag}: {e}" for e in check_val(c, x, v, fidx))
            if d["loai"] == "KHOA" and prio == 1 and "_loc" not in r:
                k = s(r.get(d["khoa"]))
                if not k:
                    errs.append(f"{tag}: thiếu khóa '{d['khoa']}'")
                elif sheet in ("SCOPE", "TCQ_LOG", "DWG") and k not in tool_keys.get(sheet, set()):
                    errs.append(f"{tag}: khóa '{k}' không có trong nháp tool (không được thêm dòng mới vào {sheet})")
                elif sheet == "FILE_INDEX" and k not in tool_keys.get(sheet, set()):
                    if not re.match(r"L\d+$", k) or not r.get("Đường dẫn") or not r.get("Vai trò"):
                        errs.append(f"{tag}: dòng FILE_INDEX mới phải có ID 'L01…' + Đường dẫn + Vai trò")
            if "_loc" in r and not any(row_match(row, r["_loc"]) for row in (res[sheet].values() if d["loai"] == "KHOA" else [])):
                xem.append(f"{tag}: _loc không khớp dòng nào {r['_loc']}")
    tcq, scope = res["TCQ_LOG"], res["SCOPE"]
    for k, r in tcq.items():
        for t in split_list(r.get("Bị thay bởi")):
            if t not in tcq:
                errs.append(f"TCQ_LOG {k}: 'Bị thay bởi' {t} không tồn tại")
            elif t == k:
                errs.append(f"TCQ_LOG {k}: tự thay chính nó")
        for b in split_list(r.get("Dòng BOQ")):
            if b not in scope:
                errs.append(f"TCQ_LOG {k}: Dòng BOQ '{b}' không có (khóa dạng 'Sheet!dòng', xem: loc SCOPE)")
        for m in split_list(r.get("Mã phòng/Khu")):
            if RX_ROOM.match(m) and rc and RX_ROOM.match(m).group(1) not in rc:
                xem.append(f"TCQ_LOG {k}: mã phòng '{m}' không có trong BOQ")
    ch = tcq_chain(tcq)
    for k, r in tcq.items():
        nx = split_list(r.get("Bị thay bởi"))
        if nx and ch[k] is None:
            errs.append(f"TCQ_LOG {k}: chuỗi 'Bị thay bởi' bị vòng")
    for k, r in scope.items():
        for t in split_list(r.get("Mã TCQ")):
            if t not in tcq:
                errs.append(f"SCOPE {k}: Mã TCQ '{t}' không có trong TCQ_LOG")
        for m in split_list(r.get("Mã VL")):
            if res["MAT"] and m not in res["MAT"]:
                xem.append(f"SCOPE {k}: Mã VL '{m}' chưa có trong MAT")
    xem.extend(link_cu(res))
    for st in STEPS[:lim]:
        for sh, d in SHEETS.items():
            for c, x in d["cot"].items():
                if x.get("bat_buoc") != st:
                    continue
                if sh == "SCOPE" and st == "doichieu":
                    rows = list(scope.items())
                else:
                    rows = list(res[sh].items()) if d["loai"] == "KHOA" else []
                miss = [k for k, r in rows if not s(_txt(r.get(c)))]
                if miss:
                    errs.append(f"{sh}: {len(miss)} dòng thiếu '{c}' (bước {st}) vd: {', '.join(miss[:6])}")
    for sh, d in SHEETS.items():  # sheet NỐI: cột bắt buộc "*" áp cho mọi dòng
        if d["loai"] != "NOI":
            continue
        for c, x in d["cot"].items():
            if x.get("bat_buoc") != "*":
                continue
            miss = [f"{s(r.get('Bước'))} #{i}" for i, r in enumerate(res[sh], 1) if not s(_txt(r.get(c)))]
            if miss:
                errs.append(f"{sh}: {len(miss)} dòng thiếu '{c}' vd: {', '.join(miss[:6])}")
    return errs, xem, res


def link_cu(res):
    """Link 'Bản vẽ hiệu lực' trỏ vào file của 1 Số BV nhưng không phải file của bản hiệu lực (DWG đã đổi bản sau khi LLM gán)."""
    dwg, fidx, out = res["DWG"], res["FILE_INDEX"], []
    for k, r in res["SCOPE"].items():
        for l in links_of(r.get("Bản vẽ hiệu lực")):
            bv = s(fidx.get(l[1], {}).get("Số BV"))
            if bv not in dwg:
                continue
            ok = {x[1] for c in ("File PDF", "File DWG") for x in links_of(dwg[bv].get(c))}
            if l[1] not in ok:
                out.append(f"SCOPE {k}: link {l[1]} là bản cũ của {bv} (hiệu lực: {', '.join(sorted(ok)) or '?'})")
    return out


def _tool_keys(p):
    out = defaultdict(set)
    for prio, fn, path, sheet, buoc, base in data_files(p):
        if prio == 0 and sheet in SHEETS and SHEETS[sheet]["loai"] == "KHOA":
            k = SHEETS[sheet]["khoa"]
            out[sheet] |= {s(r.get(k)) for r in jload(path)["rows"]}
    return out


def check_val(c, x, v, fidx):
    e = []
    k = x["kieu"]
    if v is None or v == "":
        return [f"cột '{c}' rỗng — bỏ cột thay vì ghi rỗng"]
    if k == "SO" and not isinstance(v, (int, float)):
        e.append(f"'{c}' phải là số, đang là {v!r}")
    elif k == "DS" and s(v) not in x["ds"]:
        e.append(f"'{c}' = '{v}' không thuộc: {' | '.join(x['ds'])}")
    elif k == "DSN":
        bad = [t for t in split_list(v) if t not in x["ds"]]
        if bad:
            e.append(f"'{c}' có giá trị lạ {bad} (được: {' | '.join(x['ds'])})")
    elif k == "NGAY" and not re.match(r"^\d{4}-\d{2}-\d{2}$", s(v)):
        e.append(f"'{c}' phải dạng yyyy-mm-dd")
    elif k == "LINK" and not isinstance(v, str):
        ls = links_of(v)
        if not ls:
            e.append(f"'{c}' link sai dạng: [\"nhãn\", \"FID\", trang, \"ô\"] hoặc list các link")
        for l in ls:
            if re.match(r"^[FL]\d+$", l[1]) and l[1] not in fidx:
                e.append(f"'{c}' file {l[1]} không có trong FILE_INDEX")
    return e


def cmd_check(goc, args):
    p = P(goc).need()
    den = opt(args, "--den")
    errs, xem, _ = run_check(p, den)
    for x in xem[:60]:
        print("XEM ", x)
    if len(xem) > 60:
        print(f"XEM  … còn {len(xem) - 60} dòng")
    for x in errs[:120]:
        print("LỖI ", x)
    if len(errs) > 120:
        print(f"LỖI  … còn {len(errs) - 120} dòng")
    print(f"\n{len(errs)} LỖI, {len(xem)} XEM")
    sys.exit(1 if errs else 0)


# ---------------------------------------------------------------- build
def cols_of(sheet):
    return list(SHEETS[sheet]["cot"])


def latest_track(p):
    fs = sorted(glob.glob(os.path.join(p.t, "03_boq", f"{p.cfg['du_an']}_BOQ_TRACK_*.xlsx")))
    return fs[-1] if fs else None


def pull_user(p):
    """So file đã build lần trước với _snapshot.json -> ghi chỗ user sửa vào *.user.json."""
    import openpyxl
    snap = jload(p.f("_snapshot.json"))
    if not snap:
        return 0
    n = 0
    users = defaultdict(lambda: OrderedDict())
    for sheet in SHEETS:
        u = jload(p.f(f"{sheet}.user.json"))
        if u and SHEETS[sheet]["loai"] == "KHOA":
            for r in u["rows"]:
                users[sheet][s(r.get(SHEETS[sheet]["khoa"]))] = r

    def note(sheet, key, col, val):
        nonlocal n
        k = SHEETS[sheet]["khoa"]
        row = users[sheet].setdefault(key, {k: key})
        x = SHEETS[sheet]["cot"][col]
        if x["kieu"] == "SO":
            val = num(val) if num(val) is not None else val
        elif val is not None:
            val = s(val)
        row[col] = val if val not in ("", None) else "(xóa)"
        n += 1

    for fkey in ("boq_track", "file_index"):
        path = snap.get(fkey)
        if not path or not os.path.exists(path):
            continue
        lock = os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))
        if os.path.exists(lock):
            sys.exit(f"LỖI: {path} đang mở trong Excel — đóng file rồi build lại.")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for sheet, rows in snap["reg"].items():
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            it = ws.iter_rows(values_only=True)
            hdr = [s(c) for c in next(it)]
            for r in it:
                key = s(r[0])
                if not key:
                    continue
                old = rows.get(key)
                for j, col in enumerate(hdr):
                    if col not in SHEETS[sheet]["cot"] or "build" in SHEETS[sheet]["cot"][col]["buoc"] or j >= len(r):
                        continue
                    cur = s(r[j])
                    if old is None:
                        if cur and j > 0:
                            note(sheet, key, col, r[j])
                    elif cur != s(old.get(col, "")):
                        note(sheet, key, col, r[j])
        lech = []
        for sh, info in snap.get("ct", {}).items():
            if sh not in wb.sheetnames:
                continue
            idx = {col: column_index(letter) - 1 for col, letter in info["cols"].items()}
            ji = column_index(info["item"]) - 1 if info.get("item") else None
            want = {int(k.rsplit("!", 1)[1]): k for k in info["rows"]}
            for i, r in enumerate(wb[sh].iter_rows(values_only=True), 1):  # read_only: không truy cập ô ngẫu nhiên
                key = want.get(i)
                if not key:
                    continue
                exp = info["rows"][key].get("_sig")
                if ji is not None and exp is not None:
                    cur = row_sig(r[ji + 1:ji + 5])
                    if cur != exp:
                        lech.append(f"{sh} hàng {i}: '{cur}' ≠ '{exp}'")
                        continue
                for col, j in idx.items():
                    cur = r[j] if j < len(r) else None
                    if s(cur) != s(info["rows"][key].get(col, "")):
                        note("SCOPE", key, col, cur)
        if lech:
            sys.exit(f"LỖI: hàng trong sheet BOQ của {path} bị dời (chèn/xóa hàng?) — chưa nhận sửa tay nào:\n  "
                     + "\n  ".join(lech[:10]) + ("\n  …" if len(lech) > 10 else "")
                     + "\nĐưa hàng về như bản build, hoặc sửa 3 cột chính thức ở sheet SCOPE rồi build lại.")
    for sheet, rows in users.items():
        if rows:
            jdump(p.f(f"{sheet}.user.json"), {"rows": list(rows.values())})
    return n


def derive(p, res, snap):
    tcq, scope, dwg = res["TCQ_LOG"], res["SCOPE"], res["DWG"]
    ch = tcq_chain(tcq)
    for k, r in tcq.items():
        r["Hiệu lực"] = "Hiệu lực" if ch[k] == [k] else ("Chuỗi bị vòng" if ch[k] is None else f"Bị thay → {'; '.join(ch[k])}")
    for k, r in dwg.items():
        khu = res["KHU"].get(s(r.get("Khu (thư mục)")))
        if khu and not s(_txt(r.get("Mã phòng"))):
            r["Mã phòng"] = khu.get("Mã phòng")
    for k, r in scope.items():
        bvs = set()
        for l in links_of(r.get("Bản vẽ hiệu lực")) or []:
            bvs |= set(re.findall(RX_DWG, l[0].upper()))
        bvs |= set(re.findall(RX_DWG, s(_txt(r.get("Bản vẽ hiệu lực"))).upper()))
        info = [f"{b} {s(dwg[b].get('Ngày issue hiệu lực'))} ({s(dwg[b].get('Đợt hiệu lực'))})" for b in sorted(bvs) if b in dwg]
        if info:
            r["Rev/ngày/đợt"] = "; ".join(info)
    old = (snap or {}).get("reg", {}).get("SCOPE", {})
    for k, r in scope.items():
        cur = {c: s(_txt(r.get(c))) for c in cols_of("SCOPE") if c != "Ngày cập nhật"}
        o = old.get(k)
        same = o and all(s(o.get(c, "")) == v for c, v in cur.items())
        r["Ngày cập nhật"] = s(o.get("Ngày cập nhật")) if same else TODAY.strftime("%Y-%m-%d")


def set_links(ws, c, p, items):
    """items = [(nhãn, file, trang, ô đích)]; file None = ô trong chính workbook này.
    Ô trỏ về chính nó, đích ghi ở _LINKS → add-in SRC_LINKS (_MASTER\\vba) mở PDF đúng trang, ô nhiều link hiện menu chọn.
    Riêng 1 đích trong workbook / link web → hyperlink thường (không cần add-in)."""
    from openpyxl.styles import Font
    from openpyxl.worksheet.hyperlink import Hyperlink
    items = [x for x in items if x[1] or x[3]]
    if not items:
        return
    if len(items) == 1 and not items[0][1]:
        c.hyperlink = f"#{items[0][3]}"
    elif s(items[0][1]).startswith("http"):
        c.hyperlink = items[0][1]
    else:
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{ws.title}'!{c.coordinate}", display=s(c.value))
        for lab, f, pg, sub in items:
            if not s(f).startswith("http"):
                p.links.append((ws.title, c.coordinate, lab, f or p.self_name, pg, sub))
    c.font = Font(color="0563C1", underline="single")


def put(ws, row, col, v, p, fidx, kieu=None):
    c = ws.cell(row=row, column=col)
    t = _txt(v)
    c.value = t if t != "" else None
    if kieu == "LINK":
        set_links(ws, c, p, [link_item(p, l, fidx) for l in links_of(v)])
    return c


def write_links(wb, p):
    """Sheet _LINKS đúng dạng add-in SRC_LINKS đọc: A Key(Sheet!Ô) B Sheet C Ô D Nhãn E File(tương đối) F Trang G Ô đích.
    Các dòng cùng Key phải liền nhau. H/I = mở thẳng / về ô khi không có add-in."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    ws = wb.create_sheet("_LINKS")
    hdr = [("Key", 26), ("Sheet", 20), ("Ô", 7), ("Nhãn", 55), ("File (tương đối)", 60), ("Trang", 7), ("Ô đích", 16),
           ("Mở (dự phòng)", 9), ("Về ô", 7)]
    for j, (c, w) in enumerate(hdr, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[L(j)].width = w
    lk = Font(color="0563C1", underline="single")
    for i, (sh, coord, lab, f, pg, sub) in enumerate(p.links, 2):
        for j, v in enumerate([f"{sh}!{coord}", sh, coord, lab, f, pg, sub, "Mở", "←"], 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=8).hyperlink = f"#{sub}" if f == p.self_name else f + (f"#{sub}" if sub else "")
        ws.cell(row=i, column=9).hyperlink = f"#'{sh}'!{coord}"
        ws.cell(row=i, column=8).font = ws.cell(row=i, column=9).font = lk
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(len(p.links) + 1, 2)}"


def write_reg(wb, sheet, rows, p, fidx, snap_reg):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L
    ws = wb.create_sheet(sheet)
    cols = cols_of(sheet) if sheet != "VIEC_MO" else ["ID"] + cols_of(sheet)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        long = c.startswith(("Mô tả", "Câu hỏi", "Trả lời", "Nội dung", "Đường dẫn", "Trong zip", "Việc", "Ghi chú"))
        ws.column_dimensions[L(j)].width = 45 if long else 16
    kinds = SHEETS[sheet]["cot"]
    snap = {}
    for i, r in enumerate(rows, 2):
        rec = {}
        for j, c in enumerate(cols, 1):
            cell = put(ws, i, j, r.get(c), p, fidx, kinds.get(c, {}).get("kieu"))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            rec[c] = s(cell.value)
        snap[rec[cols[0]]] = rec
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{L(len(cols))}{max(len(rows) + 1, 2)}"
    if SHEETS[sheet]["loai"] == "KHOA":
        snap_reg[sheet] = snap
    return ws


def add_boq_cols(ws, info, scope, p, fidx, scope_letters, chinh_thuc_only=False):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L
    h = info["hang_tieu_de"]
    start = info["cot_cuoi"] + 2
    ct = [c for c, x in SHEETS["SCOPE"]["cot"].items() if x.get("boq") == "chinh_thuc"]
    nb = [] if chinh_thuc_only else [c for c, x in SHEETS["SCOPE"]["cot"].items() if x.get("boq") == "noi_bo"]
    green, grey = PatternFill("solid", fgColor="E2EFDA"), PatternFill("solid", fgColor="D9D9D9")
    letters = {}
    for j, c in enumerate(ct + nb):
        col = start + j
        letters[c] = L(col)
        cell = ws.cell(row=h, column=col, value=c)
        cell.font, cell.fill = Font(bold=True), (green if c in ct else grey)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[L(col)].width = 40 if c.startswith(("Mô tả", "Nội dung", "Việc")) else 16
        if c in nb:
            ws.column_dimensions[L(col)].outlineLevel = 1
            ws.column_dimensions[L(col)].hidden = True
    if h > 1:
        ws.cell(row=h - 1, column=start, value="NHÀ THẦU" if chinh_thuc_only else "NHÀ THẦU – CHÍNH THỨC").font = Font(bold=True)
        if nb:
            ws.cell(row=h - 1, column=start + len(ct), value="NỘI BỘ – bỏ khi gửi CĐT (tra từ sheet SCOPE)").font = Font(bold=True, color="C00000")
    rows_snap = {}
    for key, r in scope.items():
        if r.get("Sheet") != ws.title:
            continue
        row = int(r["Dòng"])
        ic = column_index(info["cot_item"])
        rec = {"_sig": row_sig([ws.cell(row=row, column=ic + j).value for j in range(1, 5)])}  # pull_user kiểm hàng bị dời
        for c in ct:
            cell = put(ws, row, column_index(letters[c]), r.get(c), p, fidx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            rec[c] = s(cell.value)
        for c in nb:
            if SHEETS["SCOPE"]["cot"][c]["kieu"] == "LINK":  # công thức tra không mang theo link → ghi thẳng
                put(ws, row, column_index(letters[c]), r.get(c), p, fidx, "LINK").alignment = Alignment(wrap_text=True, vertical="top")
                continue
            sl = scope_letters[c]
            ws[f"{letters[c]}{row}"] = f'=IFERROR(INDEX(SCOPE!${sl}:${sl},MATCH("{key}",SCOPE!$A:$A,0))&"","")'
        rows_snap[key] = rec
    return {"cols": {c: letters[c] for c in ct}, "item": info["cot_item"], "rows": rows_snap}


def row_sig(vals):
    """Chữ thường (không phải công thức) đầu tiên sau cột Item — nhận ra hàng BOQ khi hàng bị dời."""
    for v in vals:
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return re.sub(r"\s+", " ", v).strip()[:60]
    return ""


def column_index(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


TONGHOP = [("Dòng BOQ gốc", 20), ("Item", 7), ("Vùng", 14), ("Nhóm", 14), ("Mô tả BOQ", 50), ("ĐV", 6), ("KL BOQ", 10),
           ("Mã phòng", 28), ("Trạng thái", 11), ("Nhóm thực hiện", 11), ("NTP chi tiết", 14), ("Bản vẽ hiệu lực", 40),
           ("Rev/ngày/đợt", 18), ("Mã VL", 12), ("Spec · VL", 50), ("Mã spec/FF&E", 14), ("Mã TCQ", 16),
           ("Nội dung cuối (VN)", 45), ("Việc còn mở", 35), ("Mô tả NT – Base", 40)]


def boq_tonghop(wb, res, p, fidx):
    """1 dòng = 1 dòng BOQ CĐT (mọi sheet). Nhìn ngang: mô tả – phạm vi – bản vẽ – vật liệu/spec – TCQ, ô nào cũng bấm được.
    Chỉ để xem: dựng lại mỗi lần build, sửa ở SCOPE/MAT/TCQ_LOG."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L
    ws = wb.create_sheet("BOQ_TONGHOP", 0)
    cols = [c for c, _ in TONGHOP]
    for j, (c, w) in enumerate(TONGHOP, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(j)].width = w
    mat, tcq = res["MAT"], res["TCQ_LOG"]
    row_of = {sh: {k: i for i, k in enumerate(res[sh], 2)} for sh in ("MAT", "TCQ_LOG")}
    for i, r in enumerate(res["SCOPE"].values(), 2):
        vl = [m for m in split_list(r.get("Mã VL")) if m in mat]
        tc = [m for m in split_list(r.get("Mã TCQ")) if m in tcq]
        spec = []
        for m in vl:
            x = mat[m]
            st = s(x.get("Trạng thái cuối"))
            spec.append(f"{m}: {s(x.get('Mô tả'))}" + (f" [{st}]" if st and st != "Theo spec" else ""))
        for j, c in enumerate(cols, 1):
            if c == "Dòng BOQ gốc":
                cell = ws.cell(row=i, column=j, value=s(r.get("Khóa")))
                set_links(ws, cell, p, [(cell.value, None, None, f"'{r['Sheet']}'!A{int(r['Dòng'])}")])
            elif c == "Spec · VL":
                cell = ws.cell(row=i, column=j, value="\n".join(spec) or None)
                set_links(ws, cell, p, [link_item(p, l, fidx) for m in vl for l in links_of(mat[m].get("Nguồn spec"))])
            elif c == "Mã VL":
                cell = put(ws, i, j, r.get(c), p, fidx)
                set_links(ws, cell, p, [(m, None, None, f"'MAT'!A{row_of['MAT'][m]}") for m in vl])
            elif c == "Mã TCQ":
                cell = put(ws, i, j, r.get(c), p, fidx)
                set_links(ws, cell, p, [(m, None, None, f"'TCQ_LOG'!A{row_of['TCQ_LOG'][m]}") for m in tc])
            else:
                cell = put(ws, i, j, r.get(c), p, fidx, SHEETS["SCOPE"]["cot"].get(c, {}).get("kieu"))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{L(len(cols))}{max(len(res['SCOPE']) + 1, 2)}"
    return ws


def dashboard(wb, res, errs, xem, p, n_user):
    from openpyxl.styles import Font
    ws = wb.create_sheet("DASHBOARD", 0)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 22
    ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 90, 14
    r = 1
    ws.cell(row=r, column=1, value=f"{p.cfg['du_an']} — BOQ TRACK {TODAY:%d/%m/%Y}").font = Font(bold=True, size=14)
    r += 2

    def line(*vals, bold=False, link=None):
        nonlocal r
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(bold=bold)
        if link:
            c = ws.cell(row=r, column=4, value="→ xem")
            c.hyperlink, c.font = link, Font(color="0563C1", underline="single")
        r += 1

    scope, tcq = res["SCOPE"], res["TCQ_LOG"]
    line("TỔNG QUAN", bold=True)
    line("Dòng BOQ", len(scope))
    for k, v in Counter(s(x.get("Trạng thái")) or "(chưa gán)" for x in scope.values()).most_common():
        line("  Trạng thái", k, v)
    for k, v in Counter(s(x.get("Nhóm thực hiện")) or "(chưa gán)" for x in scope.values()).most_common():
        line("  Nhóm thực hiện", k, v)
    line("Câu TCQ", len(tcq), f"bị thay: {sum(1 for x in tcq.values() if s(x.get('Hiệu lực')).startswith('Bị'))}")
    line("Mã VL (MAT)", len(res["MAT"]))
    line("Số BV (DWG)", len(res["DWG"]))
    line("Sửa tay user giữ lại", n_user)
    r += 1
    line("CẦN XEM", bold=True)
    line("Loại", "Liên quan", "Nội dung", bold=True)
    keys = list(scope)
    for x in errs:
        line("LỖI kiểm", "", x)
    for k, t in tcq.items():
        tt = s(t.get("Tình trạng trả lời"))
        if tt in ("Chưa trả lời", "Hoãn/Chờ") and s(t.get("Hiệu lực")) == "Hiệu lực":
            line("TCQ " + tt, k, s(t.get("Nội dung cuối (VN)")) or s(t.get("Câu hỏi (EN)"))[:200],
                 link=f"#'TCQ_LOG'!A{list(tcq).index(k) + 2}")
    for k, x in scope.items():
        if s(x.get("Trạng thái")) == "Chờ" or s(x.get("Việc còn mở")):
            line("BOQ " + (s(x.get("Trạng thái")) or "mở"), k, s(x.get("Việc còn mở")) or s(x.get("Mô tả BOQ"))[:150],
                 link=f"#'SCOPE'!A{keys.index(k) + 2}")
    for i, v in enumerate(res["VIEC_MO"]):
        line("Việc mở", s(v.get("Liên quan")), s(v.get("Nội dung")), link=f"#'VIEC_MO'!A{i + 2}")
    for k, f in res["FILE_INDEX"].items():
        if s(f.get("Vai trò")) == "Chưa có (link)":
            line("Chưa có file", k, f"{s(f.get('Ghi chú'))} {s(f.get('Đường dẫn'))}")
    miss = sum(1 for x in scope.values() if not s(x.get("Trạng thái")))
    if miss:
        line("BOQ chưa gán", f"{miss} dòng", "Trạng thái trống — lọc SCOPE cột Trạng thái")
    for x in xem[:40]:
        line("XEM", "", x)
    return ws


PART_GOC = re.compile(r"^xl/(externalLinks|drawings)/")


def tra_part_goc(src, out):
    """openpyxl ghi hỏng externalLink (mất Id rels) và drawing (mất extLst) → Excel báo 'found a problem'.
    Chép nguyên văn các part đó từ BOQ gốc (cùng tên ở cả 2 file) vào file ra."""
    with zipfile.ZipFile(src) as a:
        goc = {n: a.read(n) for n in a.namelist() if PART_GOC.match(n)}
    tmp = out + ".tmp"
    with zipfile.ZipFile(out) as b, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as c:
        for n in b.namelist():
            c.writestr(n, goc[n] if n in goc else b.read(n))
    os.replace(tmp, out)


def cmd_build(goc, args):
    import openpyxl
    p = P(goc).need()
    ra = opt(args, "--ra")
    bmap = jload(p.f("_boq_map.json"))
    if not bmap:
        sys.exit("LỖI: chưa chạy boq")
    snap_old = jload(p.f("_snapshot.json"))
    n_user = pull_user(p) if not ra else 0
    errs, xem, res = run_check(p)
    derive(p, res, snap_old)
    fidx = res["FILE_INDEX"]
    stamp = TODAY.strftime("%y%m%d")
    d_boq = ra or os.path.join(p.t, "03_boq")
    d_idx = ra or os.path.join(p.t, "01_index")
    d_out = ra or os.path.join(p.t, "04_out")
    for d in (d_boq, d_idx, d_out):
        os.makedirs(d, exist_ok=True)
    out = os.path.join(d_boq, f"{p.cfg['du_an']}_BOQ_TRACK_{stamp}.xlsx")
    for f in (out, os.path.join(d_idx, "FILE_INDEX.xlsx")):
        if os.path.exists(os.path.join(os.path.dirname(f), "~$" + os.path.basename(f))):
            sys.exit(f"LỖI: {f} đang mở trong Excel — đóng file rồi build lại.")
    if os.path.exists(out):
        bk = os.path.join(d_boq, "_backup")
        os.makedirs(bk, exist_ok=True)
        shutil.copy2(out, os.path.join(bk, f"{os.path.basename(out)[:-5]}_{datetime.datetime.now():%H%M%S}.xlsx"))
    src_boq = os.path.join(p.inp, p.cfg["boq"])
    snap = {"reg": {}, "ct": {}}
    scope_letters = {c: openpyxl.utils.get_column_letter(j) for j, c in enumerate(cols_of("SCOPE"), 1)}
    # 1. BOQ_TRACK
    p.base, p.links, p.self_name = d_boq, [], os.path.basename(out)  # link tương đối theo thư mục file ra; đích → sheet _LINKS
    wb = openpyxl.load_workbook(src_boq)
    for sh, info in bmap["sheet"].items():
        if sh in wb.sheetnames:
            snap["ct"][sh] = add_boq_cols(wb[sh], info, res["SCOPE"], p, fidx, scope_letters)
    for sh in ("SCOPE", "TCQ_LOG", "MAT", "DWG", "KHU"):
        write_reg(wb, sh, list(res[sh].values()), p, fidx, snap["reg"])
    vm = [dict(v, ID=f"V{i:03d}") for i, v in enumerate(res["VIEC_MO"], 1)]
    write_reg(wb, "VIEC_MO", vm, p, fidx, snap["reg"])
    boq_tonghop(wb, res, p, fidx)
    dashboard(wb, res, errs, xem, p, n_user)
    write_links(wb, p)
    wb.active = 0
    wb.save(out)
    tra_part_goc(src_boq, out)
    # 2. FILE_INDEX
    p.base, p.links, p.self_name = d_idx, [], "FILE_INDEX.xlsx"
    wf = openpyxl.Workbook()
    wf.remove(wf.active)
    write_reg(wf, "FILE_INDEX", list(fidx.values()), p, fidx, snap["reg"])
    f_idx = os.path.join(d_idx, "FILE_INDEX.xlsx")
    wf.save(f_idx)
    # 3. Bản gửi CĐT: BOQ gốc + chỉ cột chính thức
    p.base, p.links, p.self_name = d_out, [], ""
    wc = openpyxl.load_workbook(src_boq)
    for sh, info in bmap["sheet"].items():
        if sh in wc.sheetnames:
            add_boq_cols(wc[sh], info, res["SCOPE"], p, fidx, scope_letters, chinh_thuc_only=True)
    f_cdt = os.path.join(d_out, f"{p.cfg['du_an']}_BOQ_{stamp}_CDT.xlsx")
    wc.save(f_cdt)
    tra_part_goc(src_boq, f_cdt)
    with zipfile.ZipFile(src_boq) as a, zipfile.ZipFile(f_cdt) as b:
        ma = len([n for n in a.namelist() if "/media/" in n])
        mb = len([n for n in b.namelist() if "/media/" in n])
        if mb < ma:
            print(f"XEM: bản gửi CĐT mất hình ({mb}/{ma}) — kiểm sheet Cover trước khi gửi")
    if not ra:
        snap.update(boq_track=out, file_index=f_idx)
        jdump(p.f("_snapshot.json"), snap)
    print(f"Đã ghi:\n  {out}\n  {f_idx}\n  {f_cdt}")
    print(f"Dòng BOQ {len(res['SCOPE'])} · TCQ {len(res['TCQ_LOG'])} · MAT {len(res['MAT'])} · DWG {len(res['DWG'])} · việc mở {len(vm)} · sửa tay user giữ {n_user}")
    print(f"{len(errs)} LỖI, {len(xem)} XEM (chi tiết ở sheet DASHBOARD)")


# ---------------------------------------------------------------- main
def main(a):
    if not a:
        print(LENH)
        return
    c, rest = a[0], a[1:]
    if c == "cot":
        return cmd_cot(rest)
    if not rest:
        sys.exit(LENH)
    goc, rest = rest[0], rest[1:]
    fn = {"init": lambda: cmd_init(goc, rest), "kiemke": lambda: cmd_kiemke(goc), "boq": lambda: cmd_boq(goc),
          "tcq": lambda: cmd_tcq(goc), "dwg": lambda: cmd_dwg(goc), "mat": lambda: cmd_mat(goc),
          "doichieu": lambda: cmd_doichieu(goc), "zip": lambda: cmd_zip(goc, rest[0], rest[1:]),
          "loc": lambda: cmd_loc(goc, rest[0], rest[1:]), "check": lambda: cmd_check(goc, rest),
          "build": lambda: cmd_build(goc, rest), "trang": lambda: cmd_trang(goc), "tim": lambda: cmd_tim(goc, rest),
          "ungvien": lambda: cmd_ungvien(goc, rest), "thaydoi": lambda: cmd_thaydoi(goc)}.get(c)
    if not fn:
        sys.exit(LENH)
    fn()


if __name__ == "__main__":
    main(sys.argv[1:])
