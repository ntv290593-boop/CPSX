# -*- coding: utf-8 -*-
"""pd.py — công cụ cố định của skill 01-scan. Model chỉ ghi JSON; file PROJECT_DATA do script này dựng.

Đường dẫn: <goc> = thư mục dự án (vd D:\\COSTING EXCEL\\_ACTIVE\\NORTON), <dot> = tên thư mục trong 00_input.
  đầu vào : <goc>\\00_input\\<dot>
  dữ liệu : <goc>\\01_scan\\_data\\<dot>\\<SHEET>.<buoc>.json
  kết quả : <goc>\\01_scan\\PROJECT_DATA_<dot>.xlsx

Lệnh (chạy: py -I pd.py <lệnh> ...):
  kiemke <goc> <dot> --ds                    chỉ liệt kê gọn theo thư mục (không đọc nội dung) để chọn file cần quét
  kiemke <goc> <dot>                      kiểm kê file + trang PDF -> FILES/PAGES .kiemke.json, cache text
  tim    <goc> <dot> <regex> [--file chuoi]  tìm trang có chữ khớp regex (từ cache text)
  ungvien <goc> <dot> [--top 8] [--ra f.txt] sau bước boq: mỗi Mã -> trang PDF ứng viên chia nhóm (chi tiết/thống kê/mặt bằng/danh mục) -> _ungvien.txt
  chu    <pdf> <trang> [--clip x0,y0,x1,y1]  in lớp text của trang (clip theo tỉ lệ 0..1 của trang)
  so     <pdf> <trang> [--clip ...]          in các số trong lớp text kèm toạ độ (đọc kích thước)
  anh    <pdf> <trang> <png> [--dpi 150] [--clip ...]  xuất ảnh trang/vùng để xem bằng Read
  xem    <xlsx> [sheet] [A1:Z60]             in giá trị ô (đọc BoQ, Material List)
  gemini <a.png> [b.png …] [--claude "900,450,450"] [--quansat "Tủ lavabo (V01)"] [--hoi "câu hỏi"]
                                             Gemini (agy) đọc độc lập: số kích thước / bộ trường quan sát / câu hỏi riêng
  mail   <file.msg>                          in tiêu đề, người gửi/nhận, đính kèm, nội dung mail Outlook
  build  <goc> <dot>                         ghép JSON theo form -> PROJECT_DATA_<dot>.xlsx (giữ sửa tay của user)
  check  <goc> <dot> [--den buoc]            kiểm; exit 1 nếu có LỖI. --den: chỉ đòi cột bắt buộc của các bước tới buoc
  cot    [--buoc boq]                        in các cột bước đó được ghi (từ form) + ghi chú
  boq-tach <goc> <dot> [--file F002] [--sheet a,b] [--ra DIR]   bước boq: tách dòng BoQ -> BOQ_CDT.boq-tool.json (nháp) + issue BT… (KL chữ)
  scope  <goc> <dot> [--ra DIR]              bước boq: phân loại Phạm vi nhà máy bằng luật từ khóa -> QUAN_SAT.boq-tool.json (chỉ mã khớp chắc)
  doichieu <goc> <dot> [--ra DIR]            bước doichieu: KL BoQ↔BV, Độ tin cậy, Issue từng mã, issue DT… -> *.doichieu-tool.json (nháp)
  loc    <goc> <dot> <SHEET> cột=gt [cột~regex] [--cot a,b]   lọc dòng trên dữ liệu đã ghép (thay cho đọc cả file JSON)
  do     <goc> <dot>                         đo tỉ lệ ô do tool ghi (file *-tool.json) / tổng ô
"""
import sys, os, re, json, glob, shutil, datetime
sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
FORM = os.path.join(HERE, "..", "form", "PROJECT_DATA_FORM.xlsx")
BUOC = ["kiemke", "boq", "spec", "banve", "doichieu", "user"]
SKIP_EXT = {".dwl", ".dwl2", ".bak", ".tmp", ".log", ".plt", ".ctb", ".stb"}
# L0: file nhiễu theo tên / thư mục (khung tên CAD, xref, bản lưu) — bỏ khỏi kiểm kê, in số lượng khi chạy
SKIP_NAME = re.compile(r"^(frame|xref)[-_ ]|^thumbs\.db$", re.I)
SKIP_DIR = re.compile(r"(^|[\\/])(archive|backup|old|_old|cũ)([\\/]|$)", re.I)


# ---------------------------------------------------------------- đường dẫn
def paths(goc, dot):
    goc = os.path.abspath(goc)
    return dict(goc=goc, inp=os.path.join(goc, "00_input", dot), data=os.path.join(goc, "01_scan", "_data", dot),
                out=os.path.join(goc, "01_scan", f"PROJECT_DATA_{dot}.xlsx"), bak=os.path.join(goc, "01_scan", "_backup"))


def parse_clip(args, page):
    if "--clip" not in args: return None
    import pymupdf
    x0, y0, x1, y1 = [float(v) for v in args[args.index("--clip") + 1].split(",")]
    r = page.rect
    return pymupdf.Rect(r.x0 + x0 * r.width, r.y0 + y0 * r.height, r.x0 + x1 * r.width, r.y0 + y1 * r.height)


def opt(args, name, default=None):
    return args[args.index(name) + 1] if name in args else default


# ---------------------------------------------------------------- form
def load_form():
    import openpyxl
    wb = openpyxl.load_workbook(FORM)
    sheets = {}
    for r in wb["_SHEETS"].iter_rows(min_row=2, values_only=True):
        if r[0]: sheets[r[0]] = dict(kieu=r[1], khoa=[k.strip() for k in (r[2] or "").split(",") if k.strip()])
    cols = {}
    for r in wb["_COT"].iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        sh, c, buoc, bb, kieu, ds, rong, nhom, ct, note = r[:10]
        cols.setdefault(sh, {})[c] = dict(buoc=[b.strip() for b in (buoc or "").split(",") if b.strip()], batbuoc=bb == "x",
                                          kieu=kieu, ds=[x.strip() for x in (ds or "").split("|") if x.strip()], ct=ct, note=note)
    for sh in sheets:
        if sh in wb.sheetnames and sheets[sh]["kieu"] != "AUTO":
            hdr = [c.value for c in wb[sh][1] if c.value is not None]
            sheets[sh]["hdr"] = hdr
            for h in hdr:
                if h not in cols.get(sh, {}): raise SystemExit(f"Form lỗi: {sh}!{h} không có trong _COT")
    return wb, sheets, cols


# ---------------------------------------------------------------- dữ liệu JSON
def load_json(data):
    """-> {sheet: [(buoc, row_dict), ...]} theo thứ tự BUOC."""
    out, errs = {}, []
    files = glob.glob(os.path.join(data, "*.json"))
    def order(f):
        b = os.path.basename(f).split(".")[1].split("-")[0] if os.path.basename(f).count(".") >= 2 else "?"
        return (BUOC.index(b) if b in BUOC else 99, 0 if os.path.basename(f).split(".")[1].endswith("-tool") else 1)
    for f in sorted(files, key=order):
        name = os.path.basename(f)
        if name.startswith("_"): continue
        parts = name.split(".")
        if len(parts) != 3 or parts[1].split("-")[0] not in BUOC:
            errs.append(f"Tên file JSON sai '{name}' (đúng: <SHEET>.<buoc>.json hoặc <SHEET>.<buoc>-<nhóm>.json, buoc ∈ {BUOC})"); continue
        try:
            obj = json.load(open(f, encoding="utf-8"))
        except Exception as e:
            errs.append(f"{name}: JSON hỏng: {e}"); continue
        rows = obj["rows"] if isinstance(obj, dict) else obj
        for r in rows: out.setdefault(parts[0], []).append((parts[1], r))
    return out, errs


def norm_links(v, files_by_id, inp):
    """Giá trị ô LINK -> (text, [(nhan, abs_file, trang, o_dich)]). Chấp nhận chuỗi, list link, hoặc {text, links}."""
    if v is None or v == "": return "", []
    if isinstance(v, (str, int, float)): return str(v), []
    text, links = "", v
    if isinstance(v, dict): text, links = v.get("text", ""), v.get("links", [])
    res = []
    for l in links:
        if isinstance(l, (list, tuple)):
            l = dict(zip(["nhan", "file", "trang", "o"], l))
        f = str(l.get("file", ""))
        if re.fullmatch(r"F\d{3,4}", f): f = files_by_id.get(f, f)
        elif f and not os.path.isabs(f): f = os.path.join(inp, f)
        res.append((str(l.get("nhan") or os.path.basename(f)), os.path.normpath(f) if f else "", l.get("trang") or "", l.get("o") or ""))
    if not text and res:
        text = res[0][0] if len(res) == 1 else f"{res[0][0]}  ▾{len(res)}"
    return text, res


OWN = {}  # sheet -> chủ từng ô (KHOA: {(khóa, cột): bước}; NOI: [bước mỗi dòng]) của lần merge gần nhất — để đo tool ghi bao nhiêu


def merge(sheets, cols, raw):
    """Ghép dữ liệu các bước. KHOA: theo khóa, mỗi ô 1 bước; NOI: nối dòng. user ghi đè sau cùng."""
    errs, merged, owner = [], {}, {}
    for sh, items in raw.items():
        items = [(b, r) for b, r in items]
        if sh not in sheets or sheets[sh]["kieu"] == "AUTO":
            errs.append(f"Sheet '{sh}' không có trong form (hoặc là AUTO)"); continue
        meta, keys = cols[sh], sheets[sh]["khoa"]
        if sheets[sh]["kieu"] == "NOI":
            rows = []
            if sh == "QUAN_SAT":  # nháp tool 'Phạm vi nhà máy' nhường dòng LLM/user cùng mã
                co = {r.get("Mã") for bf, r in items if not bf.endswith("-tool") and r.get("Trường") == "Phạm vi nhà máy"}
                items = [(bf, r) for bf, r in items if not (bf.endswith("-tool") and r.get("Trường") == "Phạm vi nhà máy" and r.get("Mã") in co)]
            for bf, r in items:
                b = bf.split("-")[0]
                for c in r:
                    if c not in meta: errs.append(f"{sh}: cột lạ '{c}' (bước {bf})")
                    elif b != "user" and meta[c]["buoc"] and b not in meta[c]["buoc"]: errs.append(f"{sh}!{c}: bước {b} không được ghi cột này")
                rr = dict(r); rr.setdefault("Bước", bf) if "Bước" in meta else None
                rows.append(rr)
            merged[sh] = rows; OWN[sh] = [bf for bf, _ in items]; continue
        rows, own = {}, {}
        for bf, r in items:
            b = bf.split("-")[0]
            miss = [k for k in keys if r.get(k) in (None, "")]
            if miss: errs.append(f"{sh}: dòng thiếu khóa {miss} (bước {b}): {str(r)[:80]}"); continue
            key = tuple(str(r[k]).strip() for k in keys)
            row = rows.setdefault(key, {k: r[k] for k in keys})
            for c, v in r.items():
                if c in keys: continue
                if c not in meta: errs.append(f"{sh}: cột lạ '{c}' (bước {b})"); continue
                if b != "user" and meta[c]["buoc"] and b not in meta[c]["buoc"]:
                    errs.append(f"{sh}!{c}: bước {b} không được ghi cột này (chỉ {meta[c]['buoc']})"); continue
                if v in (None, ""): continue
                prev = own.get((key, c))
                if bf.endswith("-tool") and prev and not prev.endswith("-tool"): continue  # nháp của tool không đè dữ liệu LLM/user
                if prev and prev.endswith("-tool") and not bf.endswith("-tool") and bf != "user":  # LLM sửa nháp tool: LLM thắng
                    row[c] = v; own[(key, c)] = bf; continue
                if bf == "user" and prev and row.get(c) != v:
                    owner.setdefault("user_de", []).append(f"{sh} {'/'.join(key)} · {c}: user '{str(v)[:40]}' thay model '{str(row.get(c))[:40]}'")
                elif prev and prev != bf and prev != "kiemke" and row.get(c) != v:  # giá trị máy đoán (kiemke) được sửa đè
                    errs.append(f"{sh} {'/'.join(key)} · {c}: 2 bước ghi khác nhau ({prev} vs {bf})"); continue
                row[c] = v; own[(key, c)] = bf
        merged[sh] = list(rows.values()); OWN[sh] = own
    return merged, errs, owner.get("user_de", [])


# ---------------------------------------------------------------- kiểm kê
def cmd_ds(goc, dot):
    """kiemke --ds: liệt kê gọn theo thư mục (không đọc nội dung, không ghi file) để model xem phạm vi rồi quyết định bỏ gì vào _bo_qua.txt."""
    import pymupdf
    P = paths(goc, dot)
    if not os.path.isdir(P["inp"]): raise SystemExit(f"Không có thư mục {P['inp']}")
    dirs = {}
    for dp, dn, fn in os.walk(P["inp"]):
        for f in sorted(fn):
            if f.startswith("~$") or os.path.splitext(f)[1].lower() in SKIP_EXT: continue
            dirs.setdefault(os.path.relpath(dp, P["inp"]), []).append(f)
    for d in sorted(dirs):
        fs, ext, mb, pg = dirs[d], {}, 0.0, 0
        for f in fs:
            e = os.path.splitext(f)[1].lower(); ext[e] = ext.get(e, 0) + 1
            full = os.path.join(P["inp"], d, f); mb += os.path.getsize(full) / 1048576
            if e == ".pdf":
                try: pg += pymupdf.open(full).page_count
                except Exception: pass
        print(f"[{d}]  {len(fs)} file · {mb:.1f} MB · " + " ".join(f"{k}×{v}" for k, v in sorted(ext.items())) + (f" · PDF {pg} trang" if pg else ""))
        if len(fs) <= 8: print("    " + " | ".join(fs))
        else: print("    " + " | ".join(fs[:3]) + f" … {fs[-1]}")


def cmd_kiemke(goc, dot):
    import pymupdf
    P = paths(goc, dot)
    if not os.path.isdir(P["inp"]): raise SystemExit(f"Không có thư mục {P['inp']}")
    os.makedirs(os.path.join(P["data"], "_text"), exist_ok=True)
    allf, skipped = [], []
    # L0 riêng dự án: <goc>\01_scan\_bo_qua.txt — mỗi dòng 1 regex (so với đường dẫn tương đối, không phân biệt hoa/thường); dòng bắt đầu # là ghi chú
    rf = os.path.join(goc, "01_scan", "_bo_qua.txt"); user_rx = None
    if os.path.isfile(rf):
        pats = [l.strip() for l in open(rf, encoding="utf-8") if l.strip() and not l.lstrip().startswith("#")]
        if pats: user_rx = re.compile("|".join(f"(?:{p})" for p in pats), re.I); print(f"L0 luật riêng dự án: {len(pats)} dòng trong {rf}")
    for dp, dn, fn in os.walk(P["inp"]):
        dn.sort()
        for f in sorted(fn):
            rel = os.path.relpath(os.path.join(dp, f), P["inp"])
            if f.startswith("~$") or os.path.splitext(f)[1].lower() in SKIP_EXT: continue
            if SKIP_NAME.search(f) or SKIP_DIR.search(rel) or (user_rx and user_rx.search(rel)): skipped.append(rel); continue
            allf.append(rel)
    if skipped: print(f"L0 bỏ {len(skipped)} file nhiễu (frame/xref/Archive/log…): " + "; ".join(os.path.basename(s) for s in skipped[:6]) + (" …" if len(skipped) > 6 else ""))
    dup_report(P["inp"], allf)
    FILES, PAGES, npg = [], [], 0
    tok = re.compile(r"\b[A-Z0-9]{1,6}(?:[.\-_][A-Z0-9]{1,6}){2,}\b")
    for i, rel in enumerate(allf, 1):
        fid, ext = f"F{i:03d}", os.path.splitext(rel)[1].lower()
        full = os.path.join(P["inp"], rel)
        row = {"ID": fid, "Loại": loai_file(rel), "Tên file": os.path.basename(rel), "Đường dẫn": rel, "Ext": ext,
               "MB": round(os.path.getsize(full) / 1048576, 1), "Mở": [[os.path.basename(rel), rel]]}
        if ext == ".pdf":
            try:
                d = pymupdf.open(full)
            except Exception as e:
                row["Số trang"] = f"LỖI MỞ: {e}"[:60]; FILES.append(row); continue
            row["Số trang"] = d.page_count; texts = []
            for p in d:
                t = p.get_text(); texts.append(t); n = len(t.strip())
                W, H = p.rect.width, p.rect.height
                so, ten = title_guess(p, tok)
                PAGES.append({"File ID": fid, "Trang": p.number + 1, "Số BV": so, "Tên BV": ten,
                              "Lớp text": "CÓ" if n > 200 else ("ÍT" if n else "KHÔNG"),
                              "Mở": [[f"{fid} p{p.number + 1}", fid, p.number + 1]]})
            json.dump(texts, open(os.path.join(P["data"], "_text", fid + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
            npg += d.page_count
        FILES.append(row)
    dump(P["data"], "FILES.kiemke.json", FILES); dump(P["data"], "PAGES.kiemke.json", PAGES)
    by = {}
    for f in FILES: by[f["Loại"]] = by.get(f["Loại"], 0) + 1
    tx = build_timkiem(P, load_json(P["data"])[0])
    if tx["dup"]: print(f"L1 TRANG TRÙNG (bỏ ngày): {len(tx['dup'])} trang trùng file khác; giữ bản trong file ít trang hơn (hòa: file sau) -> _timkiem.json")
    if tx["tong"]: print("L1 FILE TỔNG (≥85% trang trùng file khác, không đưa vào tìm kiếm): " + ", ".join(sorted(tx["tong"])))
    notext = sum(1 for p in PAGES if p["Lớp text"] == "KHÔNG")
    print(f"{len(FILES)} file, {npg} trang PDF ({notext} trang không có lớp text). Theo loại: {by}")
    print("Loại 'Khác' cần model xem tên/đường dẫn rồi ghi lại vào FILES.kiemke.json nếu đoán sai.")


def dup_report(inp, rels):
    """L1: in nhóm file trùng nội dung (cùng hash) và nhóm cùng tên khác ngày/REV (chỉ báo, không xóa) — model quyết giữ bản nào."""
    import hashlib
    by_size = {}
    for r in rels: by_size.setdefault(os.path.getsize(os.path.join(inp, r)), []).append(r)
    for sz, g in by_size.items():
        if len(g) < 2 or sz == 0: continue
        h = {}
        for r in g: h.setdefault(hashlib.md5(open(os.path.join(inp, r), "rb").read()).hexdigest(), []).append(r)
        for v in h.values():
            if len(v) > 1: print("L1 TRÙNG NỘI DUNG: " + " = ".join(v))
    stem = {}
    for r in rels:
        b, e = os.path.splitext(os.path.basename(r))
        k = re.sub(r"[\s_\-]*rev[\s._-]*\d+\w*|^\d{6}(?!\d)[\s_\-]*|[\s_\-]+\d{6}$", "", b.lower()).strip(" _-")
        stem.setdefault((os.path.dirname(r), k, e.lower()), []).append(r)
    for v in stem.values():
        if len(v) > 1: print("L1 CÙNG TÊN KHÁC NGÀY/REV (giữ bản mới nhất, ghi NOTES): " + " | ".join(os.path.basename(x) for x in v))


RAC = re.compile(r"PERMISSION|COPYRIGHT|WWW\.|@|TƯ VẤN|CONSULTANT|ARCHITECT|DRAWN|CHECKED|APPROVED|SCALE|TỶ LỆ|TỈ LỆ|"
                 r"HẠNG MỤC|ITEM:|BY:|DIRECTOR|DESIGNER|NGÀY|DATE|REV|CHỦ ĐẦU TƯ|CLIENT|PROJECT|DỰ ÁN|SỐ BẢN VẼ|DRAWING NO|TÊN BẢN VẼ|DRAWING TITLE|SHEET|ISSUE", re.I)


def title_guess(p, tok):
    """Đoán số tờ + tên tờ trong khung tên (góc phải-dưới theo chiều nhìn): số = mã gần nhãn DRAWING NO (+ mảnh tiền tố);
    tên = khối chữ to nhất gần nhãn DRAWING NAME. Chỉ là đoán — bước bản vẽ sửa khi sai."""
    import pymupdf
    W, H, M = p.rect.width, p.rect.height, p.rotation_matrix
    bl = []
    for b in p.get_text("dict")["blocks"]:
        if b.get("type") != 0: continue
        r = pymupdf.Rect(b["bbox"]) * M  # toạ độ theo chiều nhìn (trang xoay 90/270)
        if r.x0 < W * 0.6 or r.y0 < H * 0.55: continue
        sp = [s for l in b["lines"] for s in l["spans"]]
        t = "\n".join("".join(s["text"] for s in l["spans"]) for l in b["lines"])
        bl.append((r.x0, r.y0, r.x1, r.y1, t, max((s["size"] for s in sp), default=0)))
    tx = lambda b: " ".join(b[4].split())
    nos = lambda b: b[4].replace(" ", "").replace("\n", "").upper()
    ctr = lambda b: ((b[0] + b[2]) / 2, (b[1] + b[3]) / 2)
    dist = lambda a, b: ((ctr(a)[0] - ctr(b)[0]) ** 2 + (ctr(a)[1] - ctr(b)[1]) ** 2) ** 0.5
    frag = re.compile(r"^[A-Z0-9][A-Z0-9.\-_]*[-.]$")
    lab_no = [b for b in bl if re.search(r"KÝHIỆUBẢNVẼ|DRAWINGNO|SỐBẢNVẼ|SHEETNO", nos(b))]
    lab_ten = [b for b in bl if re.search(r"TÊNBẢNVẼ|DRAWINGNAME|DRAWINGTITLE", nos(b))]
    num = [(b, m) for b in bl for m in tok.findall(b[4]) if not RAC.search(tx(b)) and not frag.match(tx(b))]
    if not num: return "", ""
    nb, so = min(num, key=lambda x: dist(x[0], lab_no[0])) if lab_no else num[-1]
    fr = [b for b in bl if frag.match(tx(b))]
    if fr:
        f = min(fr, key=lambda b: dist(b, nb))
        if dist(f, nb) < W * 0.08: so = tx(f) + so
    anchor = lab_ten[0] if lab_ten else nb
    ok = [b for b in bl if b is not nb and b is not anchor and 6 <= len(tx(b)) <= 120 and not RAC.search(tx(b))
          and not tok.fullmatch(tx(b)) and not frag.match(tx(b)) and re.search(r"[A-ZÀ-Ỹ]{3}", tx(b))
          and not re.search(r"CO\.,? ?LTD|STREET|WARD|TEL[: ]", tx(b), re.I)]
    ok = [b for b in ok if dist(b, anchor) < W * 0.15] or ok
    ok.sort(key=lambda b: (-round(b[5]), dist(b, anchor)))  # tên tờ = chữ to nhất gần nhãn
    return so, " / ".join(tx(b) for b in ok[:2])[:120]


def loai_file(rel):
    s, ext = rel.lower(), os.path.splitext(rel)[1].lower()
    if ext in (".msg", ".eml"): return "Mail"
    if ext in (".dwg", ".dxf", ".dwt"): return "CAD"
    if re.search(r"boq|bill|khoi luong|khối lượng", s) or (ext in (".xlsx", ".xls", ".xlsm") and re.search(r"(^|\\)\d+_", s)): return "BoQ"
    if re.search(r"material|vat lieu|vật liệu|\bml\b", s): return "Material List"
    if re.search(r"spec|chỉ dẫn|chi dan|tieu chuan|tiêu chuẩn|\d\d \d\d \d\d", s): return "Spec"
    if re.search(r"drawing|ban ve|bản vẽ|banve|\\bv\\|mass units|plan", s) or ext == ".pdf": return "Bản vẽ"
    if ext in (".xlsx", ".xls", ".xlsm", ".csv"): return "Excel khác"
    return "Khác"


def dump(d, name, rows):
    os.makedirs(d, exist_ok=True)
    json.dump({"rows": rows}, open(os.path.join(d, name), "w", encoding="utf-8"), ensure_ascii=False, indent=0)


# ---------------------------------------------------------------- công cụ đọc
def cmd_tim(goc, dot, rx, args):
    P = paths(goc, dot)
    files = {r["ID"]: r for _, r in load_json(P["data"])[0].get("FILES", [])}
    fl, pat, n = opt(args, "--file"), re.compile(rx, re.I), 0
    for f in sorted(glob.glob(os.path.join(P["data"], "_text", "F*.json"))):
        fid = os.path.basename(f)[:-5]; name = files.get(fid, {}).get("Tên file", "")
        if fl and fl.lower() not in name.lower(): continue
        for i, t in enumerate(json.load(open(f, encoding="utf-8")), 1):
            hits = sorted(set(m.group(0) for m in pat.finditer(t)))
            if hits:
                n += 1; print(f"{fid} p{i}\t{name[:50]}\t{', '.join(hits[:12])}")
    print(f"-- {n} trang")


def _fold(s):
    import unicodedata
    s = unicodedata.normalize("NFD", str(s or "").lower().replace("đ", "d"))
    return re.sub(r"\s+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn"))


STOP = set("va voi cho cua the loai dang kieu bang theo tren duoi trong ngoai khong co mot cac cai bo and the with for of in on".split())


def ngoai_pham_vi(qs_rows):
    """Mã bước boq đã phân loại ngoài scope nhà máy (5 xưởng BOHO): QUAN_SAT Trường 'Phạm vi nhà máy', Giá trị 'KHÔNG THUỘC SCOPE NHÀ MÁY…'."""
    return {str(r.get("Mã")) for r in qs_rows
            if r.get("Trường") == "Phạm vi nhà máy" and str(r.get("Giá trị") or "").strip().upper().startswith("KHÔNG THUỘC SCOPE")}


# ---------------------------------------------------------------- tìm trang ứng viên
LOAI_TO = [  # (loại, regex trên tên tờ đã _fold) — theo thứ tự; "chi tiết" thắng mep/hoan_thien
    ("danh_muc", r"danh muc|drawing list|muc luc|\bindex\b|cover"),
    ("thong_ke", r"thong ke|schedule|bang ke|statistic"),
    ("mep", r"\bdien\b|den tran|\bden\b|cong tac|o cam|cap thoat|\bpccc\b|\bmep\b|lighting|electrical|plumbing|ceiling plan|\bm&e\b"),
    ("mat_bang", r"mat bang|layout|floor plan|furniture plan"),
    ("hoan_thien", r"hoan thien|\bsan\b|\btran\b|\btuong\b|finish|flooring|ceiling"),
    ("chi_tiet", r"chi tiet|detail|\bdt\b|section|mat cat"),
    ("mat_dung", r"mat dung|elevation"),
    ("legend", r"legend|ky hieu|ghi chu chung"),
]
NHOM_RA = [("chi_tiet", "CHI TIẾT", 8), ("thong_ke", "THỐNG KÊ", 3), ("mat_bang", "MẶT BẰNG", 4), ("danh_muc", "DANH MỤC", 2)]
TOK = re.compile(r"[a-z0-9]+")
STOP2 = STOP | set("cung cap lap dat hoan thien ban ve hoan cong phong".split())
CODE_IN_TEXT = re.compile(r"(?<![a-z0-9])([a-z]{1,4})[\s._-]?0*(\d{1,3})(?![0-9])")


def loai_to(title):
    s = _fold(title)
    has_ct = bool(re.search(r"chi tiet|detail", s))
    for loai, rx in LOAI_TO:
        if loai in ("mep", "hoan_thien") and has_ct: continue
        if re.search(rx, s): return loai
    return "khac"


def toks(s): return [w for w in TOK.findall(_fold(s)) if len(w) >= 2 and w not in STOP2]


def ncode(s):
    """Chuẩn hóa mã: bỏ . - _ khoảng trắng, bỏ số 0 đầu: LF-04 / LF04 / lf.4 -> lf4. Tiền tố 'ID-' bỏ."""
    s = re.sub(r"^id[\s._-]+", "", _fold(s).strip())
    m = re.fullmatch(r"([a-z]{1,4})[\s._-]*0*(\d{1,3})", s)
    return f"{m.group(1)}{m.group(2)}" if m else None


def base_code(ma):
    """Mã BoQ gốc: bỏ hậu tố tự đặt _ST/_1BR/_2BR, -DT/-HM/-TL (hạng mục con), chữ a/b cuối."""
    s = re.sub(r"_(st|1br|2br)$", "", _fold(ma))
    s = re.sub(r"-(dt|hm|tl)$", "", s)
    s = re.sub(r"(\d)[a-d]$", r"\1", s)
    return ncode(s)


def _area_rx(a):
    """'blk01' khớp 'BLK 01' / 'blk-01': cho phép khoảng trắng/gạch giữa chữ và số."""
    return r"(?<![a-z0-9])" + re.sub(r"(?<=[a-z])(?=\d)", lambda m: r"[\s._-]*", re.escape(a)) + r"(?![a-z0-9])"


def _area_of(rel, areas):
    """Khu ghi ở phần sâu nhất của đường dẫn (tên file, rồi thư mục gần nhất); nhiều khu cùng chỗ -> bỏ qua."""
    for part in reversed(re.split(r"[\\/]", rel)):
        hit = {a for a in areas if re.search(_area_rx(a), _fold(part))}
        if len(hit) == 1: return hit
        if len(hit) > 1: return set()
    return set()


def khu_key(k):
    """Khóa khu: từ đầu của Khu; nếu từ đầu chỉ là chữ và từ sau là số (BLK 01, TMDV 01) thì ghép: blk01."""
    w = _fold(k).split()
    return w[0] + w[1] if len(w) > 1 and w[0].isalpha() and w[1].isdigit() else (w[0] if w else "")


def build_timkiem(P, sheets, write=True):
    """Chỉ mục tìm kiếm từ cache text + FILES/PAGES (máy làm): trang trùng, file tổng, loại tờ. Ghi _timkiem.json."""
    import hashlib
    files = {r["ID"]: r for _, r in sheets.get("FILES", [])}
    titles = {}
    for b, r in sheets.get("PAGES", []):
        if r.get("Trang"): titles[(r["File ID"], int(r["Trang"]))] = f'{r.get("Số BV") or ""} {r.get("Tên BV") or ""}'.strip()
    texts = {}
    for f in sorted(glob.glob(os.path.join(P["data"], "_text", "F*.json"))):
        texts[os.path.basename(f)[:-5]] = json.load(open(f, encoding="utf-8"))
    h = {}
    for fid, pages in texts.items():
        for i, t in enumerate(pages, 1):
            k = re.sub(r"\s+", "", re.sub(r"\d{1,2}[/.]\d{1,2}[/.]\d{2,4}", "", t))  # bỏ ngày: bản 29/09 vs 30/09 chỉ khác ngày
            if len(k) > 50: h.setdefault(hashlib.md5(k.encode("utf-8")).hexdigest(), []).append((fid, i))
    npg = {fid: len(p) for fid, p in texts.items()}
    dup = {}
    for v in h.values():
        if len({fid for fid, _ in v}) < 2: continue
        keep = min(v, key=lambda x: (npg[x[0]], -int(x[0][1:]), x[1]))  # file ít trang hơn (cụ thể hơn); hòa -> file sau
        for p in v:
            if p != keep: dup[p] = keep
    tong = []
    for fid, pages in texts.items():
        real = [(fid, i) for i, t in enumerate(pages, 1) if len(re.sub(r"\s+", "", t)) > 50]
        if real and sum(p in dup for p in real) >= 0.85 * len(real): tong.append(fid)
    tk = {"trang_trung": {f"{a} p{b}": f"{c} p{d}" for (a, b), (c, d) in sorted(dup.items())},
          "file_tong": tong,
          "loai": {f"{fid} p{i}": loai_to(titles.get((fid, i), "")) for fid, pages in texts.items() for i in range(1, len(pages) + 1)}}
    if write: json.dump(tk, open(os.path.join(P["data"], "_timkiem.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    return dict(texts=texts, files=files, titles=titles, dup=dup, tong=set(tong), tk=tk)


def read_aux(P, name):
    f = os.path.join(P["data"], name)
    return json.load(open(f, encoding="utf-8")) if os.path.isfile(f) else {}


def build_index(P, sheets, write=True):
    """Dựng chỉ mục xếp hạng dùng chung cho `ungvien` và bench. LLM ghi: _timkiem.kiemke.json (du_phong), _timkiem.boq.json (khu_file, ma_cha, tu_khoa)."""
    import collections
    ix = build_timkiem(P, sheets, write)
    kk, bq = read_aux(P, "_timkiem.kiemke.json"), read_aux(P, "_timkiem.boq.json")
    ix["du_phong"] = set(kk.get("du_phong", []))
    ix["bq"] = bq
    pages = [(fid, i) for fid in ix["texts"] for i in range(1, len(ix["texts"][fid]) + 1)]
    ix["pages"] = pages
    ix["body"] = {p: collections.Counter(toks(ix["texts"][p[0]][p[1] - 1])) for p in pages}
    ix["title"] = {p: collections.Counter(toks(ix["titles"].get(p, ""))) for p in pages}
    ix["codes"] = {p: collections.Counter(a + b for a, b in CODE_IN_TEXT.findall(_fold(ix["texts"][p[0]][p[1] - 1]))) for p in pages}
    ix["N"] = len(pages)
    ix["cdf"] = collections.Counter(c for cn in ix["codes"].values() for c in cn)
    ix["avgdl"] = (sum(sum(c.values()) for c in ix["body"].values()) / len(pages)) if pages else 1
    ix["df"] = collections.Counter(w for c in ix["body"].values() for w in c)
    ix["loai"] = {p: ix["tk"]["loai"].get(f"{p[0]} p{p[1]}", "khac") for p in pages}
    ix["tit_codes"] = {p: {ncode(x) for x in re.findall(r"[A-Za-z]{1,4}[-._]?\d{1,3}", ix["titles"].get(p, ""))} for p in pages}
    return ix


def _bm25(ix, q, c, k1=1.2, b=0.75):
    import math
    dl = sum(c.values()); s = 0.0
    for w in q:
        if w in c:
            idf = math.log(1 + (ix["N"] - ix["df"][w] + 0.5) / (ix["df"][w] + 0.5))
            s += idf * c[w] * (k1 + 1) / (c[w] + k1 * (1 - b + b * dl / ix["avgdl"]))
    return s


def file_allowed(ix, fid, khus, areas, want):
    """Lọc khu: _timkiem.boq.json khu_file (LLM ghi) ưu tiên; không có thì đoán Khu theo đường dẫn file. True = giữ."""
    kf = ix["bq"].get("khu_file", {})
    given = [kf[k] for k in khus if k in kf]
    if given: return any(fid == x or fid in x for x in sum(given, []))
    if not want: return True
    fa = _area_of(ix["files"].get(fid, {}).get("Đường dẫn", ""), areas)
    return not fa or bool(fa & want)


def rank_ma(ix, ma, r, khus, areas):
    """-> list (điểm, (fid, trang), số lần mã xuất hiện) xếp giảm dần cho 1 mã BoQ. r: dòng BOQ_TONGHOP đã ghép."""
    kw = ix["bq"].get("tu_khoa", {}).get(ma)
    q_ten = set(toks(kw if kw else re.split(r"\s*@", r.get("Hạng mục") or "")[0]))  # "Gối nằm @ phòng ngủ": phần sau @ là vị trí, không là tên món
    q_mota = set(toks((r.get("Mô tả BoQ") or "")[:300]))
    qs = q_ten | q_mota
    # BM25 chỉ bật khi ≥ 50% từ của Mô tả có trong lớp text BV (cùng ngôn ngữ); không thì dùng từ khóa Hạng mục kiểu cũ
    bm_on = bool(qs) and sum(1 for w in qs if ix["df"].get(w)) / len(qs) >= 0.5
    bc = base_code(ma)
    # ma_bv (LLM ghi): mã trên BV của món có mã tự đặt (V01 -> VC1); chữ không phải mã (GUONG) thêm vào từ khóa
    xtra = ix["bq"].get("ma_bv", {}).get(ma, [])
    bcs = ({bc} if bc else set()) | {ncode(x) for x in xtra if ncode(x)}
    qs = qs | {w for x in xtra if not ncode(x) for w in toks(x)}
    # mã trùng nhiều trang (ô lưới F1, cao độ…) không phân biệt được món: bỏ khỏi tín hiệu mã
    bcs = {c for c in bcs if ix["cdf"].get(c, 0) <= max(8, 0.1 * ix["N"])}
    want = {khu_key(k) for k in khus}
    res = []
    for p in ix["pages"]:
        fid = p[0]
        if p in ix["dup"] or fid in ix["tong"] or fid in ix["du_phong"]: continue
        if ix["loai"][p] in ("mep", "hoan_thien"): continue
        lech = not file_allowed(ix, fid, khus, areas, want)  # khác khu: không loại hẳn (file chung kiểu GENERAL vẫn dùng được), chỉ hạ điểm
        if bm_on:
            sc = _bm25(ix, qs, ix["body"][p]) + 3.0 * sum(1 for w in qs if w in ix["title"][p])
        else:
            kh = [w for w in q_ten if w in ix["body"][p]]
            sc = len(kh) if kh and len(kh) >= min(2, len(q_ten)) else 0
        mh = sum(ix["codes"][p].get(c, 0) for c in bcs)
        sc += 8 * min(mh, 3) + (10 if bcs & ix["tit_codes"][p] else 0)
        if lech: sc *= 0.4
        if sc > 0: res.append((sc, p, mh))
    res.sort(key=lambda x: -x[0])
    return res


def ungvien_all(ix, sheets):
    """-> ({mã: {row, khu, res, cha, res_cha}}, số mã ngoài scope). Đã bỏ mã ngoài scope nhà máy."""
    boq = {}
    for _, r in sheets.get("BOQ_TONGHOP", []): boq.setdefault(str(r["Mã"]), {}).update({k: v for k, v in r.items() if v not in (None, "")})
    ngoai = ngoai_pham_vi([r for _, r in sheets.get("QUAN_SAT", [])])
    bmap = {}
    for _, r in sheets.get("BOQ_MAP", []):
        if r.get("Khu"): bmap.setdefault(str(r["Mã"]), set()).add(r["Khu"])
    areas = {khu_key(k) for v in bmap.values() for k in v}
    cha = ix["bq"].get("ma_cha", {})
    out = {}
    for ma, r in boq.items():
        if ma in ngoai: continue
        khus = sorted(bmap.get(ma, []))
        out[ma] = dict(row=r, khu=khus, res=rank_ma(ix, ma, r, khus, areas), cha=cha.get(ma, []))
    for ma, d in out.items():  # trang của mã cha (phụ kiện -> cửa): chỉ để TÌM, bước 4 vẫn đọc riêng từng mã
        d["res_cha"] = [x for c in d["cha"] if c in out for x in out[c]["res"][:6]]
    return out, len(ngoai)


def cmd_ungvien(goc, dot, args):
    """Sau bước boq: mỗi Mã trong BOQ_TONGHOP -> trang PDF ứng viên, chia nhóm theo loại tờ. Ghi _ungvien.txt (hoặc --ra <file>).
    Tín hiệu: mã chuẩn hóa · BM25(từ khóa + Mô tả + tên tờ) khi cùng ngôn ngữ · lọc khu. Loại trước: trang trùng, file tổng, file DỰ PHÒNG, tờ MEP/hoàn thiện."""
    P = paths(goc, dot); sheets = load_json(P["data"])[0]
    if not sheets.get("BOQ_TONGHOP"): raise SystemExit("Chưa có BOQ_TONGHOP (chạy bước boq trước)")
    ix = build_index(P, sheets, write="--ra" not in args)
    allm, nngoai = ungvien_all(ix, sheets)
    top = int(opt(args, "--top", 8))
    out, lab = [], (lambda p: f"{p[0]} p{p[1]}")
    for ma, d in allm.items():
        out.append(f"## {ma} · {d['row'].get('Hạng mục') or ''}" + (f"   [khu: {', '.join(d['khu'])}]" if d["khu"] else ""))
        if not d["res"]:
            out.append("   KHÔNG CÓ ỨNG VIÊN — tìm theo tên tờ / Drawing List, hoặc thật sự chưa có tờ"); continue
        for loai, ten, n in NHOM_RA:
            if loai == "chi_tiet":
                n = top; hit = [x for x in d["res"] if ix["loai"][x[1]] in ("chi_tiet", "mat_dung", "khac", "legend")]
            else:  # thống kê / mặt bằng / danh mục: chỉ lấy trang có mã
                hit = [x for x in d["res"] if ix["loai"][x[1]] == loai and x[2]]
            if not hit: continue
            tt = "; ".join(dict.fromkeys(ix["titles"].get(x[1], "")[:40] for x in hit[:3] if ix["titles"].get(x[1])))
            out.append(f"   {ten:9s} " + ", ".join(f"{lab(x[1])}{' ×' + str(x[2]) if x[2] else ''}" for x in hit[:n])
                       + (f"  (+{len(hit) - n})" if len(hit) > n else "") + (f"   | {tt}" if tt and loai == "chi_tiet" else ""))
        if d["res_cha"]:
            out.append("   CỦA MÃ CHA " + ", ".join(lab(x[1]) for x in d["res_cha"][:6]) + f"   (mã cha: {', '.join(d['cha'])})")
    txt = "\n".join(out)
    dst = opt(args, "--ra", os.path.join(P["data"], "_ungvien.txt"))
    open(dst, "w", encoding="utf-8").write(txt + "\n")
    print(txt)
    print(f"-- {len(allm)} mã -> {dst}; bỏ {nngoai} mã ngoài scope; {len(ix['dup'])} trang trùng; file tổng: {sorted(ix['tong']) or 'không'}; dự phòng: {sorted(ix['du_phong']) or 'không'}")


def cmd_chu(pdf, pg, args):
    import pymupdf
    p = pymupdf.open(pdf)[int(pg) - 1]
    print(p.get_text(clip=parse_clip(args, p), sort=True))


def cmd_so(pdf, pg, args):
    import pymupdf
    p = pymupdf.open(pdf)[int(pg) - 1]
    W, H = p.rect.width, p.rect.height
    for w in p.get_text("words", clip=parse_clip(args, p), sort=True):
        if re.fullmatch(r"[\d.,x×/+\-]+", w[4]) and re.search(r"\d", w[4]):
            print(f"{w[4]}\t x={w[0] / W:.3f} y={w[1] / H:.3f}")


def cmd_anh(pdf, pg, png, args):
    import pymupdf
    p = pymupdf.open(pdf)[int(pg) - 1]
    pix = p.get_pixmap(dpi=int(opt(args, "--dpi", 150)), clip=parse_clip(args, p))
    os.makedirs(os.path.dirname(os.path.abspath(png)), exist_ok=True)
    pix.save(png); print(png, pix.width, "x", pix.height)


GEM_MODEL = "gemini-3.8-flash-low"  # benchmark 24/10: low đủ chính xác (~13 s/crop); medium/high quá chậm
GEM_SO = ("It is a crop of an interior/joinery drawing (dimensions in millimetres). List EVERY dimension number printed along "
          "dimension lines in the crop, including vertical dimension lines, one entry per printed number, in reading order. "
          "Do NOT include detail numbers inside circles, drawing numbers, scale text, level markers (FFL/FCL/AFFL) or note text. "
          "If a number is unreadable write \"?\". Never guess. Output ONLY a JSON array of strings, nothing else.")


GEM_TRUONG = ["Lắp đặt (gắn tường/đứng rời/âm tường)", "Kích thước BV (R×S×C mm, chỉ số in trên bản vẽ)", "Có cánh (số cánh, mở/lùa)",
              "Ngăn kéo (số)", "Thanh treo", "Kệ/đợt (số)", "Hốc/ô tivi", "Mặt đá", "Gương", "Kính", "Kim loại",
              "Giấy/vải", "Mã hoàn thiện trên BV (nhãn WD/PT/LAM/SS… chỉ vào món này)", "Khung xương", "Cong/bo",
              "Phần ngoài phạm vi (chậu, vòi, thiết bị… thấy trên hình)", "Số chi tiết/hình đang xem"]
GEM_QS = ("It is a page or crop of an interior/joinery drawing. Inspect ONLY this item: {vat}. "
          "For each field below, report what is VISIBLE in the image. Output ONLY a JSON object: "
          "{{\"<field>\": {{\"value\": \"…\", \"status\": \"YES|NO|NOT VISIBLE\", \"where\": \"view/detail number + position in the image\"}}}}. "
          "YES = present, NO = clearly absent, NOT VISIBLE = cannot tell from this image. Count only what you can see; never guess, "
          "never infer from typical designs. Fields (keep these exact Vietnamese keys): {fields}")


def agy_exe():
    import shutil as sh
    p = sh.which("agy")
    if p: return p
    g = glob.glob(os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "WinGet", "Packages", "Google.AntigravityCLI*", "**", "agy.exe"), recursive=True)
    if not g: raise SystemExit("Không thấy agy.exe (Antigravity CLI). Bỏ qua kiểm chéo Gemini, ghi vào kết quả trả về.")
    return g[0]


def cmd_gemini(args):
    """Gemini đọc độc lập ảnh crop (người đọc thứ 2). Mặc định: liệt kê số kích thước; --hoi: câu hỏi riêng.
    --quansat "Tủ lavabo (V01)": trả JSON đủ bộ trường quan sát (GEM_TRUONG) để so từng trường với Claude.
    --claude "900,450,450": so với số Claude đọc -> KHỚP / chỉ Claude / chỉ Gemini. Kết quả lưu cạnh ảnh: <ảnh>.gemini.txt"""
    import subprocess, concurrent.futures as cf
    from collections import Counter
    pngs = [a for a in args if a.lower().endswith(".png")]
    hoi, mine, par = opt(args, "--hoi"), opt(args, "--claude"), int(opt(args, "--par", 4))
    vat = opt(args, "--quansat")
    if vat: hoi = GEM_QS.format(vat=vat, fields=json.dumps(GEM_TRUONG, ensure_ascii=False))
    exe = agy_exe()
    def run(png):
        png = os.path.abspath(png)
        out = png + ".gemini.txt"
        if os.path.exists(out) and os.path.getsize(out) > 2 and not hoi: return png, open(out, encoding="utf-8").read()
        prompt = (f"Do NOT run shell commands; use your file-viewing tool directly on this absolute path: {png} . "
                  + ((hoi if vat else hoi + " Answer only from what is visible in the image; if not visible say NOT VISIBLE. Never guess.") if hoi else GEM_SO))
        try:
            r = subprocess.run([exe, "--model", GEM_MODEL, "--print", prompt, "--print-timeout", "280s"], cwd=os.path.dirname(png),
                               capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=320)
            txt = r.stdout.strip() or f"LỖI agy: {r.stderr.strip()[:200]}"
        except subprocess.TimeoutExpired:
            txt = "LỖI agy: quá 320 s"
        if not txt.startswith("LỖI") and not hoi: open(out, "w", encoding="utf-8").write(txt)
        return png, txt
    with cf.ThreadPoolExecutor(par) as ex:
        res = list(ex.map(run, pngs))
    allg = []
    for png, txt in res:
        print(f"== {os.path.basename(png)}\n{txt}")
        if not hoi:
            m = re.search(r"\[.*\]", txt, re.S)
            try: allg += [str(x).strip() for x in json.loads(m.group(0))] if m else []
            except ValueError: print("  (không đọc được JSON của Gemini)")
    if mine and not hoi:
        a = Counter(x.strip() for x in mine.split(",") if x.strip()); g = Counter(x for x in allg if x != "?")
        print(f"-- SO SÁNH: khớp {sum((a & g).values())} · chỉ Claude {sorted((a - g).elements())} · chỉ Gemini {sorted((g - a).elements())}")


def cmd_mail(msg):
    """In tiêu đề, người gửi, người nhận, tệp đính kèm, nội dung của file Outlook .msg."""
    import olefile
    o = olefile.OleFileIO(msg)
    def get(tag, base=""):
        for enc, suf in (("utf-16-le", "001F"), ("cp1252", "001E")):
            p = f"{base}__substg1.0_{tag}{suf}"
            if o.exists(p): return o.openstream(p).read().decode(enc, "replace").strip("\x00")
        return ""
    print("Tiêu đề :", get("0037")); print("Từ      :", get("0C1A"), get("0C1F"))
    print("Tới     :", get("0E04")); print("Cc      :", get("0E03"))
    att = sorted({e[0] for e in o.listdir() if e[0].startswith("__attach")})
    print("Đính kèm:", [get("3707", a + "/") or get("3704", a + "/") for a in att])
    print("-" * 60); print(get("1000"))


def cmd_xem(xlsx, args):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if not args:
        for ws in wb.worksheets: print(f"[{ws.title}] {ws.max_row}x{ws.max_column}")
        return
    ws = wb[args[0]]
    rng = args[1] if len(args) > 1 else f"A1:{openpyxl.utils.get_column_letter(min(ws.max_column or 26, 30))}{min(ws.max_row or 60, 80)}"
    from openpyxl.utils import range_boundaries
    c0, r0, c1, r1 = range_boundaries(rng)
    for i, row in enumerate(ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True), r0):
        v = [f"{openpyxl.utils.get_column_letter(c0 + j)}={str(x)[:60]}" for j, x in enumerate(row) if x not in (None, "")]
        if v: print(i, " | ".join(v))


# ---------------------------------------------------------------- build
def snapshot_rows(ws, hdr, cols):
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v in (None, "") for v in r): continue
        rows.append({h: r[i] for i, h in enumerate(hdr) if i < len(r) and cols[h]["kieu"] not in ("CT", "AUTO", "LINK")})
    return rows


def pull_user(P, sheets, cols):
    """So file xlsx hiện có với ảnh chụp lần build trước -> sửa tay của user ghi vào <SHEET>.user.json."""
    snap_f = os.path.join(P["data"], "_snapshot.json")
    if not (os.path.exists(P["out"]) and os.path.exists(snap_f)): return []
    import openpyxl
    snap, wb, notes = json.load(open(snap_f, encoding="utf-8")), openpyxl.load_workbook(P["out"]), []
    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO" or sh not in wb.sheetnames or sh not in snap: continue
        ws = wb[sh]; hdr = [c.value for c in ws[1]]
        if [h for h in hdr if h] != meta["hdr"]: notes.append(f"{sh}: user đổi cột trong file — không lấy sửa tay sheet này"); continue
        cur, old = snapshot_rows(ws, meta["hdr"], cols[sh]), snap[sh]
        uf = os.path.join(P["data"], f"{sh}.user.json")
        user = json.load(open(uf, encoding="utf-8"))["rows"] if os.path.exists(uf) else []
        if meta["kieu"] == "KHOA":
            k = meta["khoa"]; kk = lambda r: tuple(str(r.get(x)).strip() for x in k)
            om = {kk(r): r for r in old}; um = {kk(r): r for r in user}
            for r in cur:
                o = om.get(kk(r))
                ch = {c: v for c, v in r.items() if c not in k and (o is None or norm(v) != norm(o.get(c))) and not (o is None and v in (None, ""))}
                if ch:
                    um.setdefault(kk(r), {x: r[x] for x in k}).update(ch)
                    notes.append(f"{sh} {'/'.join(kk(r))}: user sửa {list(ch)}")
            user = list(um.values())
        else:
            olds = [json.dumps({c: norm(v) for c, v in r.items()}, ensure_ascii=False, sort_keys=True) for r in old]
            seen = [json.dumps({c: norm(v) for c, v in r.items()}, ensure_ascii=False, sort_keys=True) for r in user]
            for r in cur:
                s = json.dumps({c: norm(v) for c, v in r.items()}, ensure_ascii=False, sort_keys=True)
                if s in olds: olds.remove(s)
                elif s not in seen:
                    user.append({c: v for c, v in r.items() if v not in (None, "")}); seen.append(s)
                    notes.append(f"{sh}: user thêm dòng {str(list(r.values())[:3])[:60]}")
        if user: dump(P["data"], f"{sh}.user.json", user)
    return notes


def norm(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()


def cmd_build(goc, dot):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.hyperlink import Hyperlink
    P = paths(goc, dot)
    lock = os.path.join(os.path.dirname(P["out"]), "~$" + os.path.basename(P["out"]))
    if os.path.exists(lock): raise SystemExit(f"DỪNG: {os.path.basename(P['out'])} đang mở trong Excel. Đóng file rồi chạy lại.")
    wb, sheets, cols = load_form()
    notes = pull_user(P, sheets, cols)
    raw, errs = load_json(P["data"])
    merged, e2, user_de = merge(sheets, cols, raw)
    errs += e2
    files_by_id = {r["ID"]: os.path.join(P["inp"], r["Đường dẫn"]) for r in merged.get("FILES", [])}
    LK = Font(color="0563C1", underline="single")
    REG, snap = [], {}
    pos = {sh: {h: i + 1 for i, h in enumerate(m.get("hdr", []))} for sh, m in sheets.items()}

    def fx(tpl, sh, r):
        def rep(m):
            s, c = (m.group(1).split("!") if "!" in m.group(1) else (sh, m.group(1)))
            col = L(pos[s][c])
            return f"'{s}'!${col}:${col}" if "!" in m.group(1) else f"{col}{r}"
        return "=" + re.sub(r"\{([^{}]+)\}", rep, tpl)

    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO": continue
        ws, hdr, rows = wb[sh], meta["hdr"], merged.get(sh, [])
        if sh == "FILES": auto_file_status(rows, merged, files_by_id)
        for i, row in enumerate(rows, 2):
            for j, h in enumerate(hdr, 1):
                m, v, cell = cols[sh][h], row.get(h), ws.cell(i, j)
                if m["kieu"] == "CT":
                    cell.value = fx(m["ct"], sh, i)
                elif h == "STT":
                    cell.value = i - 1
                elif m["kieu"] == "LINK":
                    text, links = norm_links(v, files_by_id, P["inp"])
                    cell.value = text or None
                    if links:
                        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sh}'!{cell.coordinate}", display=text); cell.font = LK
                        for l in links: REG.append((sh, cell.coordinate) + l)
                else:
                    if isinstance(v, (list, dict)): v = json.dumps(v, ensure_ascii=False)
                    cell.value = v
                    if isinstance(v, str) and v.startswith("="): cell.data_type = "s"
                if m["kieu"] in ("TEXT",) and isinstance(cell.value, str) and len(cell.value) > 40:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        if rows: ws.auto_filter.ref = f"A1:{L(len(hdr))}{len(rows) + 1}"
        snap[sh] = [{h: r.get(h) for h in hdr if cols[sh][h]["kieu"] not in ("CT", "AUTO", "LINK")} for r in rows]
    # _LINKS
    ws, outdir = wb["_LINKS"], os.path.dirname(P["out"])
    for i, (s, c, lab, f, pg, sub) in enumerate(REG, 2):
        rel = os.path.relpath(f, outdir) if f else ""
        ws.append([f"{s}!{c}", s, c, lab, rel, pg, sub, "Mở", 1.0])
        if rel:
            ws.cell(i, 8).hyperlink = rel + (("#" + sub) if sub else ""); ws.cell(i, 8).font = LK
    res = run_checks(P, sheets, cols, merged, REG, errs, None)
    dashboard(wb, sheets, merged, res, dot, notes + [f"Ghi đè: {u}" for u in user_de])
    for s in ("_SHEETS", "_COT", "_HUONG_DAN"):
        if s in wb.sheetnames: wb[s].sheet_state = "hidden"
    if os.path.exists(P["out"]):
        os.makedirs(P["bak"], exist_ok=True)
        shutil.copy2(P["out"], os.path.join(P["bak"], f"PROJECT_DATA_{dot}_{datetime.datetime.now():%y%m%d-%H%M%S}.xlsx"))
    wb.save(P["out"])
    json.dump(snap, open(os.path.join(P["data"], "_snapshot.json"), "w", encoding="utf-8"), ensure_ascii=False, default=str)
    print(f"Đã ghi {P['out']}  ({sum(len(v) for v in merged.values())} dòng, {len(REG)} link)")
    for n in notes + user_de: print("  user:", n)
    report(res)
    return res


def auto_file_status(rows, merged, files_by_id):
    used = {}
    for sh, rs in merged.items():
        if sh in ("FILES", "PAGES"): continue
        for r in rs:
            for v in r.values():
                if isinstance(v, (list, dict)):
                    for l in (v.get("links", []) if isinstance(v, dict) else v):
                        f = l[1] if isinstance(l, (list, tuple)) and len(l) > 1 else (l.get("file") if isinstance(l, dict) else None)
                        if f: used[f] = used.get(f, 0) + 1
    read = {}
    for p in merged.get("PAGES", []):
        if p.get("Đã đọc bằng"): read[p["File ID"]] = read.get(p["File ID"], 0) + 1
    for r in rows:
        n = used.get(r["ID"], 0) + used.get(r["Đường dẫn"], 0)
        pg = r.get("Số trang")
        s = f"đọc {read.get(r['ID'], 0)}/{pg} trang" if isinstance(pg, int) else ""
        r["Trạng thái scan"] = "; ".join(x for x in [s, f"{n} link dẫn tới" if n else "chưa có link dẫn tới"] if x)


# ---------------------------------------------------------------- check
def run_checks(P, sheets, cols, merged, REG, errs, den):
    import pymupdf
    E, W = list(errs), []
    upto = BUOC[:BUOC.index(den) + 1] if den else BUOC
    ngoai = ngoai_pham_vi(merged.get("QUAN_SAT", []))
    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO": continue
        for i, r in enumerate(merged.get(sh, []), 2):
            for h, m in cols[sh].items():
                v = r.get(h)
                empty = v in (None, "") or (isinstance(v, (list, dict)) and not v)
                if str(r.get("Mã")) in ngoai and m["buoc"] and "banve" in m["buoc"]: continue  # mã ngoài scope: không đọc bản vẽ
                if m["batbuoc"] and empty and m["kieu"] not in ("CT", "AUTO") and (not m["buoc"] or any(b in upto for b in m["buoc"])):
                    E.append(f"{sh} dòng {i}: thiếu '{h}'")
                if m["kieu"] == "SO" and not empty and not isinstance(v, (int, float)):
                    E.append(f"{sh} dòng {i}: '{h}' = '{v}' phải là số (không để trong ngoặc kép)")
                if m["kieu"] == "LIST" and not empty and m["ds"] and str(v) not in m["ds"]:
                    E.append(f"{sh} dòng {i}: '{h}' = '{v}' không thuộc {m['ds']}")
    pc = {}
    for s, c, lab, f, pg, sub in REG:
        if not f: E.append(f"Link {s}!{c} '{lab}': thiếu file"); continue
        if not os.path.exists(f): E.append(f"Link {s}!{c} '{lab}': không thấy file {f}"); continue
        if pg and f.lower().endswith(".pdf"):
            if f not in pc:
                try: pc[f] = pymupdf.open(f).page_count
                except Exception: pc[f] = 10 ** 6
            if not str(pg).isdigit() or int(pg) > pc[f] or int(pg) < 1: E.append(f"Link {s}!{c} '{lab}': trang {pg} ngoài file ({pc[f]} trang)")
    T, M, C = merged.get("BOQ_TONGHOP", []), merged.get("BOQ_MAP", []), merged.get("BOQ_CDT", [])
    ma_t, ma_c = {str(r["Mã"]) for r in T}, {str(r.get("Mã")) for r in C if r.get("Mã")}
    if C or T:
        for x in sorted(ma_c - ma_t): E.append(f"Mã {x} có trong BOQ_CDT nhưng thiếu ở BOQ_TONGHOP")
        for x in sorted(ma_t - ma_c): E.append(f"Mã {x} ở BOQ_TONGHOP nhưng không có dòng BOQ_CDT nào")
        mk = {(str(r.get("Mã")), str(r.get("Khu"))) for r in C if r.get("Mã")}
        mm = {(str(r["Mã"]), str(r["Khu"])) for r in M}
        for x in sorted(mk - mm): E.append(f"BOQ_MAP thiếu cặp Mã/Khu {x}")
        for x in sorted(mm - mk): E.append(f"BOQ_MAP có cặp Mã/Khu {x} không có trong BOQ_CDT")
    # _timkiem.*.json do LLM ghi (phục vụ ungvien): mã có thật, file có thật
    fids = {str(r.get("ID")) for r in merged.get("FILES", [])}
    kk, bq = read_aux(P, "_timkiem.kiemke.json"), read_aux(P, "_timkiem.boq.json")
    for f in kk.get("du_phong", []):
        if f not in fids: E.append(f"_timkiem.kiemke.json du_phong: file {f} không có trong FILES")
    for ma, cha in bq.get("ma_cha", {}).items():
        for x in [ma] + list(cha):
            if x not in ma_t: E.append(f"_timkiem.boq.json ma_cha: mã {x} không có trong BOQ_TONGHOP")
    for ma in bq.get("tu_khoa", {}):
        if ma not in ma_t: E.append(f"_timkiem.boq.json tu_khoa: mã {ma} không có trong BOQ_TONGHOP")
    khus = {str(r.get("Khu")) for r in M}
    for k, fl in bq.get("khu_file", {}).items():
        if k not in khus: E.append(f"_timkiem.boq.json khu_file: Khu '{k}' không có trong BOQ_MAP")
        for f in fl:
            if f not in fids: E.append(f"_timkiem.boq.json khu_file '{k}': file {f} không có trong FILES")
    kl = {}
    for r in C:
        try: kl[str(r.get("Mã"))] = kl.get(str(r.get("Mã")), 0) + float(r.get("KL CĐT") or 0)
        except ValueError: E.append(f"BOQ_CDT {r.get('Sheet')}!{r.get('Dòng')}: KL CĐT không phải số '{r.get('KL CĐT')}'")
    bv = {}
    for r in M:
        if r.get("KL theo BV") not in (None, ""):
            try: bv[str(r["Mã"])] = bv.get(str(r["Mã"]), 0) + float(r["KL theo BV"])
            except ValueError: E.append(f"BOQ_MAP {r['Mã']}/{r['Khu']}: KL theo BV không phải số")
    qs = {}
    for r in merged.get("QUAN_SAT", []):
        qs.setdefault(str(r.get("Mã")), []).append(r)
        if r.get("Nguồn") in ("HÌNH", "TEXT") and not r.get("Tờ · Chi tiết"): E.append(f"QUAN_SAT {r.get('Mã')} · {r.get('Trường')}: nguồn {r.get('Nguồn')} mà không ghi tờ/chi tiết")
    for ma_, tr_, ch_, tong_ in kt_chuoi_sai(merged.get("QUAN_SAT", [])): E.append(f"QUAN_SAT {ma_} · {tr_}: chuỗi cộng sai '{ch_}' (tổng thật {tong_:g})")
    if ngoai: W.append(f"{len(ngoai)} mã ngoài scope nhà máy, không đọc bản vẽ: {', '.join(sorted(ngoai)[:12])}{' …' if len(ngoai) > 12 else ''}")
    for r in T:
        ma = str(r["Mã"])
        if ma in ngoai: continue
        if "banve" in upto and ma not in qs: E.append(f"{ma}: chưa có dòng QUAN_SAT nào (bước banve chưa làm mã này)")
        elif "banve" in upto and not any(x.get("Trường") == "Kiểm chéo Gemini" for x in qs[ma]):
            E.append(f"{ma}: thiếu dòng QUAN_SAT 'Kiểm chéo Gemini' (bước banve phải cho Gemini đọc lại các crop)")
        if ma in bv and abs(kl.get(ma, 0) - bv[ma]) > 1e-6: W.append(f"{ma}: KL BoQ {kl.get(ma, 0):g} ≠ KL theo BV {bv[ma]:g}")
        if r.get("Độ tin cậy") == "THAP": W.append(f"{ma}: độ tin cậy THẤP — {str(r.get('Giả định tính giá') or '')[:60]}")
        for h in ("Bản vẽ đối chứng", "Spec · ML"):
            if isinstance(r.get(h), str) and r[h].upper().startswith(("KHÔNG", "CHƯA")): W.append(f"{ma}: {h} = {r[h]}")
    for ma, rs in qs.items():
        kt = [x.get("Trường") for x in rs if x.get("Trạng thái") == "KHÔNG THẤY"]
        mt = [x.get("Trường") for x in rs if x.get("Trạng thái") == "MÂU THUẪN"]
        if kt: W.append(f"{ma}: KHÔNG THẤY {len(kt)} trường ({', '.join(map(str, kt[:5]))})")
        if mt: W.append(f"{ma}: MÂU THUẪN {len(mt)} trường ({', '.join(map(str, mt[:5]))})")
    rfi = {str(r["RFI"]) for r in merged.get("RFI_LOG", [])}
    for r in merged.get("ISSUES", []):
        for x in re.split(r"[,;\s]+", str(r.get("RFI") or "")):
            if x and x not in rfi: E.append(f"ISSUES {r['ID']}: RFI '{x}' không có trong RFI_LOG")
        if r.get("Trạng thái") in (None, "", "MỞ") and r.get("Mức") == "CAO": W.append(f"Issue {r['ID']} ({r.get('Mã')}) mức CAO đang mở: {str(r.get('Nội dung'))[:70]}")
    for r in merged.get("ASK_USER", []):
        if not r.get("Trả lời"): W.append(f"Hỏi user {r['ID']}: {str(r.get('Câu hỏi'))[:90]}")
    nt = [p for p in merged.get("PAGES", []) if p.get("Lớp text") == "KHÔNG" and not p.get("Đã đọc bằng")]
    if nt: W.append(f"{len(nt)} trang PDF không có lớp text chưa được xem bằng hình (PAGES: Lớp text = KHÔNG)")
    return dict(E=E, W=W)


def report(res):
    print(f"KIỂM: {len(res['E'])} LỖI, {len(res['W'])} ngoại lệ cho user")
    for x in res["E"][:60]: print("  LỖI ", x)
    if len(res["E"]) > 60: print(f"  … còn {len(res['E']) - 60} lỗi")
    for x in res["W"][:60]: print("  XEM ", x)


def cmd_check(goc, dot, args):
    P = paths(goc, dot)
    wb, sheets, cols = load_form()
    raw, errs = load_json(P["data"])
    merged, e2, _ = merge(sheets, cols, raw)
    files_by_id = {r["ID"]: os.path.join(P["inp"], r["Đường dẫn"]) for r in merged.get("FILES", [])}
    REG = []
    for sh, rows in merged.items():
        for i, r in enumerate(rows, 2):
            for h, v in r.items():
                if cols.get(sh, {}).get(h, {}).get("kieu") == "LINK":
                    for l in norm_links(v, files_by_id, P["inp"])[1]: REG.append((sh, f"{h}@{i}") + l)
    res = run_checks(P, sheets, cols, merged, REG, errs + e2, opt(args, "--den"))
    report(res)
    sys.exit(1 if res["E"] else 0)


# ---------------------------------------------------------------- dashboard
def dashboard(wb, sheets, merged, res, dot, notes):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink
    ws = wb["DASHBOARD"]
    pj = {str(r.get("Trường")): r.get("Giá trị") for r in merged.get("PROJECT", [])}
    ws["A1"] = f"{pj.get('Tên dự án', '')} · {pj.get('Gói thầu', '')}".strip(" ·") or dot
    ws["A1"].font = Font(bold=True, size=14)
    info = [("Hạn trả giá", pj.get("Hạn trả giá", "chưa có")), ("Đợt hồ sơ", dot),
            ("Hạng mục", len(merged.get("BOQ_TONGHOP", []))), ("Dòng BoQ", len(merged.get("BOQ_CDT", []))),
            ("File / trang", f"{len(merged.get('FILES', []))} / {len(merged.get('PAGES', []))}"),
            ("Quan sát", len(merged.get("QUAN_SAT", []))), ("Issue / RFI", f"{len(merged.get('ISSUES', []))} / {len(merged.get('RFI_LOG', []))}"),
            ("Kiểm máy", f"{len(res['E'])} LỖI · {len(res['W'])} ngoại lệ"), ("Cập nhật", f"{datetime.datetime.now():%d/%m/%Y %H:%M}")]
    for i, (a, b) in enumerate(info, 3):
        ws.cell(i, 1, a).font = Font(bold=True); ws.cell(i, 2, b)
    LK = Font(color="0563C1", underline="single")
    for i, s in enumerate([x for x in wb.sheetnames if x not in ("DASHBOARD", "_SHEETS", "_COT", "_HUONG_DAN")], 3):
        c = ws.cell(i, 4, s); c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{s}'!A1", display=s); c.font = LK
    r0 = len(info) + 5
    ws.cell(r0, 1, "CẦN XEM (máy đã lọc — chỉ xem các dòng này)").font = Font(bold=True, size=12)
    red = PatternFill("solid", fgColor="FCE4D6")
    for k, x in enumerate([("LỖI", e) for e in res["E"]] + [("XEM", w) for w in res["W"]] + [("USER", n) for n in notes], r0 + 1):
        ws.cell(k, 1, x[0]); ws.cell(k, 2, x[1])
        if x[0] == "LỖI": ws.cell(k, 1).fill = red


def cmd_cot(args):
    """In cột mà 1 bước được ghi (lấy từ form) — subagent dùng thay cho việc nhớ cột."""
    _, sheets, cols = load_form()
    b = opt(args, "--buoc")
    for sh, meta in sheets.items():
        if meta["kieu"] == "AUTO": continue
        cs = [(h, m) for h, m in cols[sh].items() if (not b or b in m["buoc"]) and m["kieu"] not in ("CT", "AUTO")]
        if not cs: continue
        print(f"\n## {sh}  ({meta['kieu']}{', khóa: ' + ', '.join(meta['khoa']) if meta['khoa'] else ''})  -> file {sh}.{b or '<buoc>'}.json")
        for h, m in cs:
            ds = f" [{' | '.join(m['ds'])}]" if m["ds"] else ""
            print(f"  - {h}{' (bắt buộc)' if m['batbuoc'] else ''} · {m['kieu']}{ds}{' · ' + m['note'] if m.get('note') else ''}")


# ---------------------------------------------------------------- phần B: tool trích / so / lọc (LLM chỉ quyết phần mơ hồ)
# Nháp của tool ghi ra <SHEET>.<buoc>-tool.json: LLM sửa thì LLM thắng, user thắng tất cả (xem merge). ID issue của tool: BT01…, DT01…
def load_merged(P, steps=None):
    wb, sheets, cols = load_form()
    raw, errs = load_json(P["data"])
    if steps: raw = {sh: [(b, r) for b, r in it if b.split("-")[0] in steps] for sh, it in raw.items()}
    merged, e2, _ = merge(sheets, cols, raw)
    return sheets, cols, merged, errs + e2


def write_tool(P, ra, sheet, buoc, rows):
    d = ra or P["data"]; os.makedirs(d, exist_ok=True)
    f = os.path.join(d, f"{sheet}.{buoc}-tool.json")
    if rows:
        with open(f, "w", encoding="utf-8") as fh:
            fh.write('{"rows": [\n' + ",\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + "\n]}\n")
    elif os.path.exists(f): os.remove(f)
    return f if rows else None


def num_or_none(v):
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): return v
    return None


# ---- B2 boq-tach
HDR_ROLES = [  # (vai, regex trên chữ đã _fold) — vai đầu khớp thắng; cột đầu tiên của mỗi vai được dùng
    ("tech", r"mo ta vat lieu|mo ta ky thuat|thong so|specification|material desc|technical|\bspec\b"),
    ("qty", r"^(khoi luong|kl|so luong|sl|qty|quantity|q'ty)\b"),
    ("unit", r"^(dvt|don vi|unit|uom)\b"),
    ("ma", r"^(ma sp|ma so|ma hang muc|ma\b|code|ref\b|item code)"),
    ("stt", r"^(stt|so tt|no\.?|item|#)$"),
    ("dai", r"^(dai|length|l)$"), ("rong", r"^(rong|width|w|sau|depth|d)$"), ("cao", r"^(cao|height|h)$"),
    ("desc", r"hang muc|description|mo ta|noi dung|cong viec|work item|ten "),
]
TONG_RX = re.compile(r"^(tong|total|cong\b|sub-?total|grand|thanh tien)")
ROMAN = re.compile(r"^[IVXLC]+$")


def sheet_header(ws, maxr=30):
    """-> (dòng tiêu đề, dòng cuối của tiêu đề, {vai: cột 1-based}) hoặc None. Tiêu đề có thể 2-3 dòng (R/S/C ở dòng dưới)."""
    for r in range(1, min(ws.max_row, maxr) + 1):
        roles, last = {}, r
        for rr in range(r, min(r + 3, ws.max_row + 1)):
            for c in range(1, min(ws.max_column, 30) + 1):
                v = ws.cell(rr, c).value
                if not isinstance(v, str) or not v.strip(): continue
                t = _fold(v).strip()
                if len(t) > 60: continue  # câu dài = tiêu đề bảng / ghi chú, không phải đầu cột
                for vai, rx in HDR_ROLES:
                    dim = vai in ("dai", "rong", "cao")
                    if dim != (rr > r): continue  # R/S/C chỉ nhận ở dòng dưới tiêu đề; các vai còn lại chỉ ở dòng tiêu đề
                    if re.search(rx, t):
                        if vai not in roles: roles[vai] = c
                        if vai in ("dai", "rong", "cao"): last = max(last, rr)
                        break
        if "desc" in roles and "qty" in roles and "unit" in roles:
            return r, last, roles
    return None


def cmd_boq_tach(goc, dot, args):
    import openpyxl
    P = paths(goc, dot); ra = opt(args, "--ra")
    sheets, cols, merged, errs = load_merged(P, ("kiemke", "user"))
    boqs = [r for r in merged.get("FILES", []) if r.get("Loại") == "BoQ"]
    fsel = opt(args, "--file")
    if fsel: boqs = [r for r in boqs if fsel in (r["ID"], r["Tên file"])]
    shsel = [s.strip() for s in opt(args, "--sheet", "").split(",") if s.strip()]
    if not boqs: raise SystemExit("Không có file Loại=BoQ trong FILES.kiemke (hoặc --file sai).")
    out, iss, side = [], [], {}
    for fr in boqs:
        wb = openpyxl.load_workbook(os.path.join(P["inp"], fr["Đường dẫn"]), data_only=True)
        for ws in wb.worksheets:
            if shsel and ws.title not in shsel: continue
            tag = f"[{fr['ID']}] {ws.title}"
            if not shsel and re.match(r"^(th|tong hop|summary|cover|bia)\b", _fold(ws.title)):
                print(f"{tag}: BỎ QUA (sheet tổng hợp; dùng --sheet để ép)"); continue
            h = sheet_header(ws)
            if not h: print(f"{tag}: KHÔNG THẤY dòng tiêu đề (cần Mô tả / ĐV / KL) — LLM tự đọc sheet này"); continue
            hr, hlast, R = h
            n = nt = nk = ntong = ngroup = 0; bill = ""

            def cv(r, vai):
                c = R.get(vai)
                return ws.cell(r, c).value if c else None

            for r in range(hlast + 1, ws.max_row + 1):
                desc, unit, q, stt = cv(r, "desc"), cv(r, "unit"), cv(r, "qty"), cv(r, "stt")
                d = str(desc).strip() if desc not in (None, "") else ""
                if not d: continue
                if TONG_RX.match(_fold(d)): ntong += 1; print(f"{tag}!{r}: dòng tổng '{d[:40]}' = {q}"); continue
                qn = num_or_none(q)
                if qn is None and q in (None, "") and unit in (None, ""):
                    if str(stt or "").strip() and ROMAN.match(str(stt).strip()): bill = d.replace("\n", " ")
                    continue  # dòng tiêu đề nhóm / ghi chú
                if qn is not None and not str(stt or "").strip() and cv(r, "tech") in (None, "") and cv(r, "ma") in (None, "")                         and all(cv(r, v) in (None, "", " ") for v in ("dai", "rong", "cao")):
                    ngroup += 1; print(f"{tag}!{r}: bỏ dòng nhóm có KL nhưng không STT/mã/mô tả/kích thước: '{d[:40]}' ({unit} {q})"); continue
                if qn is not None: qn = round(qn, 6)
                row = {"Sheet": ws.title, "Dòng": r, "Mô tả": desc if isinstance(desc, str) else d, "Link": [[f"BoQ {ws.title}!{r}", fr["ID"], "", f"'{ws.title}'!A{r}"]]}
                if bill: row["Bill"] = bill
                if str(stt or "").strip(): row["Item"] = str(stt).strip()
                if unit not in (None, ""): row["ĐV"] = str(unit).strip()
                if qn is not None: row["KL CĐT"] = qn; n += 1
                elif q not in (None, ""):
                    nt += 1
                    iss.append({"ID": f"BT{len(iss) + 1:02d}", "Loại": "BoQ: KL ghi chữ", "Mã": f"{ws.title}!{r}", "Mức": "TB",
                                "Nội dung": f"KL ở BoQ ghi chữ \"{str(q).strip()[:30]}\" (không phải số): \"{d[:60]}\"", "Nguồn": row["Link"]})
                else: nk += 1
                t = cv(r, "tech")
                if t not in (None, ""): row["Mô tả kỹ thuật"] = t if isinstance(t, str) else str(t)
                out.append(row)
                ex = {}
                if cv(r, "ma") not in (None, ""): ex["ma_boq"] = str(cv(r, "ma")).strip()
                dm = [cv(r, v) for v in ("dai", "rong", "cao")]
                if any(x not in (None, "", " ") for x in dm): ex["dai_rong_cao"] = [str(x).strip() if x not in (None, "") else "" for x in dm]
                if ex: side[f"{ws.title}!{r}"] = ex
            cm = ", ".join(f"{k}={openpyxl.utils.get_column_letter(v)}" for k, v in R.items())
            print(f"{tag}: tiêu đề dòng {hr}-{hlast} ({cm}) -> {n} dòng có KL, {nt} KL chữ, {nk} thiếu KL, {ntong} dòng tổng, {ngroup} dòng nhóm bỏ")
    f1 = write_tool(P, ra, "BOQ_CDT", "boq", out); write_tool(P, ra, "ISSUES", "boq", iss)
    d = ra or P["data"]
    if side: json.dump(side, open(os.path.join(d, "_boq_tach.json"), "w", encoding="utf-8"), ensure_ascii=False)
    print(f"Ghi {len(out)} dòng BOQ_CDT nháp, {len(iss)} issue BT… -> {f1}")
    print("LLM còn phải đặt: Khu, Mã (cột bắt buộc), Hạng mục; mã BoQ (ma_boq) và R/S/C của BoQ ở _boq_tach.json")


# ---- B3 scope: phân loại Phạm vi nhà máy bằng luật từ khóa (2-boq.md mục 4b). Chỉ gán khi khớp chắc; còn lại để LLM/user.
VT_XUONG = re.compile(r"\b(van ep|van mdf|mdf|veneer|melamine|laminate|picomat|nem\b|van nen|go (tram|ash|tu nhien|dac|tap|oc cho)|son pu)\b")
SC_PK_CUA = r"^(ban le|tay nam|tay gat|than khoa|ruot khoa|chan cua|dem khi|khoa)\b"
SC_OUT = [  # (nhóm, regex trên Hạng mục+Mô tả BoQ đã _fold, kiểm vật tư xưởng lẫn vào?)
    ("Khóa điện tử", r"khoa (dien tu|tu|the)|digital lock", True),
    ("Sơn nước", r"^(thi cong )?son\b", False),
    ("Thạch cao", r"tran (chim|thach cao|phang|giat cap)|hoc den|mat dung|nap tham|thanh [vz]\b|khoet lo|mieng gio|hop den|hop man|shadowline|^thach cao", True),
    ("Đá", r"^lat da\b|ngach cua|^op da\b", False),
    ("Sàn", r"^lat van san|^san\b|san go|^lat san", False),
    ("Nẹp kim loại", r"^(nep|u|v)\b.*(inox|zenta|nhom|dong)|^nep\b", True),
    ("Lam nhựa/nhôm", r"\blam (nhua|nhom|composite)|composite", True),
    ("Cửa thép", r"cua thep", True),
    ("Cửa nhôm kính", r"cua nhom|vach nhom", True),
    ("Vách compact", r"compact", True),
    ("Kim loại trang trí", r"^op (inox|dong)|lam (inox|kim loai)", True),
    ("Lan can", r"lan can", True),
    ("Đèn", r"^(den|led)\b|\bden (tha|ban|tran|trang tri)", True),
    ("Điện", r"o cam|cong tac|di day", True),
    ("Thiết bị", r"^(chau|voi|bep|hut mui|lo nuong|tv|ket sat|may|thiet bi)\b", True),
    ("Nệm", r"^nem\b", True), ("Chăn ga gối", r"^(chan|ga giuong|drap)\b", True),
    ("Rèm", r"^rem\b", True), ("Thảm", r"^tham\b", True),
    ("Decor", r"^(tranh|khung tranh|decor|binh hoa|do trang tri)\b", True), ("Signage", r"signage|bien hieu", True),
]
SC_IN = [  # (xưởng, regex) — theo thứ tự; trước đó đã loại các nhóm ngoài scope
    ("mộc", r"picomat"),
    ("cửa", r"^vach kinh|vach kinh|^cua\b|^cung cap va lap dat vach kinh"),
    ("mộc", r"guong"),
    ("cabinet", r"\btu\b|he tu|^bang|^hop che|^dot tu|^khung che|^he (bang|khung)|bench|vanity|ban console|ban lam viec|locker|minibar|quay (le tan|bar)|\bke bep"),
    ("loose", r"^(goi|sofa|giuong|ghe|ban|ke|bục|buc)\b|goi (trang tri|nam|tay)"),
]


def scope_phan_loai(hm, mt, kt):
    """-> (giá trị 'Phạm vi nhà máy' hoặc None, luật). None = tool không chắc -> LLM quyết."""
    S = _fold(f"{hm} {mt}"); T = _fold(kt)
    if re.search(SC_PK_CUA, S) and not re.search(r"khoa (dien tu|tu|the)", S): return "TRONG SCOPE – cửa (phụ kiện cửa)", "pk_cua"
    if re.search(r"picomat", S + " " + T[:60]): return "TRONG SCOPE – mộc (picomat CNC)", "picomat"
    if re.search(r"hop man", S) and re.search(r"van ep|van mdf", T): return "TRONG SCOPE – mộc (hộp màn lót ván ép)", "hop_man_van"
    if re.search(r"^(len chan|chi go|chi nhua)|\blen chan\b", S):
        if re.search(r"mdf|van (ep|nen)|laminate", T): return "TRONG SCOPE – mộc (len chân MDF/ván)", "len_van"
        if re.search(r"nhua", T): return "KHÔNG THUỘC SCOPE NHÀ MÁY – Len/chỉ nhựa", "len_nhua"
        return None, "len_mo_ho"
    if (re.search(r"\bgach\b", S) and not re.match(r"^(nep|u|v)\b", S)) or (re.match(r"^san\b", S) and re.search(r"\bgach\b", T[:30])): return "KHÔNG THUỘC SCOPE NHÀ MÁY – Gạch", "gach"
    for nhom, rx, kiem in SC_OUT:
        if re.search(rx, S):
            if kiem:
                vt = sorted(set(m.group(0).strip() for m in VT_XUONG.finditer(S + " " + T)))
                if vt: return f"CHƯA RÕ – {nhom} nhưng có {', '.join(vt[:3])}", "lan_vt"
            return f"KHÔNG THUỘC SCOPE NHÀ MÁY – {nhom}", "out_" + nhom
    for xuong, rx in SC_IN:
        if re.search(rx, S): return f"TRONG SCOPE – {xuong}", "in_" + xuong
    return None, "khong_khop"


def cmd_scope(goc, dot, args):
    from collections import Counter
    P = paths(goc, dot); ra = opt(args, "--ra")
    raw, _ = load_json(P["data"])
    sheets, cols, merged, errs = load_merged(P, ("boq",))
    da_co = {str(r.get("Mã")) for b, r in raw.get("QUAN_SAT", []) if not b.endswith("-tool") and r.get("Trường") == "Phạm vi nhà máy"}
    rows, bo, kk = [], [], []
    for t in merged.get("BOQ_TONGHOP", []):
        ma = str(t.get("Mã"))
        if ma in da_co: continue
        v, rule = scope_phan_loai(t.get("Hạng mục") or "", t.get("Mô tả BoQ") or "", t.get("Mô tả kỹ thuật BoQ") or "")
        if v is None: kk.append((ma, rule, t.get("Hạng mục"))); continue
        rows.append({"Mã": ma, "Trường": "Phạm vi nhà máy", "Giá trị": v, "Trạng thái": "CÓ", "Nguồn": "BOQ", "Model": "pd.py scope"})
    f = write_tool(P, ra, "QUAN_SAT", "boq", rows)
    c = Counter(r["Giá trị"].split("–")[0].strip() for r in rows)
    print(f"Ghi {len(rows)} dòng QUAN_SAT nháp ({dict(c)}); {len(da_co)} mã LLM đã phân loại (bỏ qua); {len(kk)} mã tool không chắc -> LLM quyết -> {f}")
    for ma, rule, hm in kk: print(f"  KHÔNG CHẮC {ma} [{rule}] {hm}")
    for r in rows:
        if r["Giá trị"].startswith("CHƯA RÕ"): print(f"  CHƯA RÕ {r['Mã']}: {r['Giá trị']}")


# ---- B4 loc
def cmd_loc(goc, dot, args):
    P = paths(goc, dot)
    sheets, cols, merged, errs = load_merged(P)
    sh = args[0]
    if sh not in merged: raise SystemExit(f"Sheet '{sh}' không có dữ liệu. Có: {', '.join(merged)}")
    conds, show = [], opt(args, "--cot")
    for a in args[1:]:
        if a.startswith("--"): break
        m = re.match(r"^(.+?)(!=|~|=)(.*)$", a)
        if not m: raise SystemExit(f"Điều kiện sai '{a}' (dạng cột=giá_trị | cột!=giá_trị | cột~regex)")
        conds.append(m.groups())

    def txt(v):
        if isinstance(v, (list, dict)): return norm_links(v, {}, "")[0] or json.dumps(v, ensure_ascii=False)
        return "" if v is None else str(v)

    def ok(r):
        for c, op, val in conds:
            x = txt(r.get(c))
            if op == "=" and x != val: return False
            if op == "!=" and x == val: return False
            if op == "~" and not re.search(val, x, re.I): return False
        return True
    rows = [r for r in merged[sh] if ok(r)]
    keys, cc = sheets[sh]["khoa"], [c for c, _, _ in conds]
    want = [c.strip() for c in show.split(",")] if show else None
    w = int(opt(args, "--rong", 70))
    for r in rows:
        cs = (keys + [c for c in want if c not in keys]) if want else [c for c in r if c in keys or c in cc] + [c for c in r if c not in keys and c not in cc][:4]
        print(" | ".join(f"{c}={txt(r.get(c))[:w].replace(chr(10), ' ')}" for c in cs if r.get(c) not in (None, "")))
    print(f"-- {len(rows)}/{len(merged[sh])} dòng")


# ---- B1 doichieu (+ B5 kiểm chuỗi cộng)
PRICE_FIELDS = ["Kích thước BV", "Có cánh", "Ngăn kéo", "Kệ/đợt", "Mặt đá", "Gương", "Kính", "Kim loại", "Giấy/vải"]
MUC = {"CAO": 3, "TB": 2, "THAP": 1}
GIA_RX = re.compile(r"dot|canh|ngan|kich thuoc|sau|rong|cao|day|mat da|da|kinh|guong|kim loai|vat lieu|giay|vai")
CHUOI = re.compile(r"(?<![+×x*\d.,])(\d+(?:[.,]\d+)?(?:\s*\+\s*\d+(?:[.,]\d+)?)+)\s*=\s*(\d+(?:[.,]\d+)?)")


def worse(a, b): return a if MUC[a] <= MUC[b] else b


def nat_key(s): return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]


def kt_chuoi_sai(qs_rows):
    """QUAN_SAT Giá trị có chuỗi 'a+b=c' mà tổng sai (B5). -> [(mã, trường, chuỗi, tổng thật)]"""
    bad = []
    for r in qs_rows:
        for m in CHUOI.finditer(str(r.get("Giá trị") or "")):
            try:
                parts = [float(x.replace(",", ".")) for x in re.split(r"\s*\+\s*", m.group(1))]
                tot = float(m.group(2).replace(",", "."))
            except ValueError: continue
            if abs(sum(parts) - tot) > 0.51: bad.append((str(r.get("Mã")), str(r.get("Trường")), m.group(0), sum(parts)))
    return bad


def gemini_lech(gem):
    """Dòng 'Kiểm chéo Gemini' -> đoạn sau 'lệch:' nếu khác 'không', else None."""
    m = re.search(r"lệch:\s*([^;]*)", str(gem[0].get("Giá trị") or "")) if gem else None
    return m.group(1).strip() if m and not _fold(m.group(1)).strip().startswith("khong") else None


def cmd_doichieu(goc, dot, args):
    P = paths(goc, dot); ra = opt(args, "--ra")
    sheets, cols, M, errs = load_merged(P, ("kiemke", "boq", "spec", "banve", "user"))
    T, MAP, C, QS, IS = (M.get(k, []) for k in ("BOQ_TONGHOP", "BOQ_MAP", "BOQ_CDT", "QUAN_SAT", "ISSUES"))
    if not T: raise SystemExit("Chưa có BOQ_TONGHOP (bước boq).")
    ngoai = ngoai_pham_vi(QS)
    ma_set = {str(r["Mã"]) for r in T}
    qs, kl, bvkl = {}, {}, {}
    for r in QS: qs.setdefault(str(r.get("Mã")), []).append(r)
    for r in C:
        n = num_or_none(r.get("KL CĐT")); k = (str(r.get("Mã")), str(r.get("Khu")))
        if n is not None: kl[k] = kl.get(k, 0) + n
    iss_of = {}
    rx_ma = {m: re.compile(r"(?<![\w-])" + re.escape(m) + r"(?![\w-])") for m in ma_set}
    for r in IS:
        if r.get("Trạng thái") in ("ĐÃ CHỐT", "ĐÓNG"): continue
        for m, rx in rx_ma.items():  # chỉ khớp theo cột Mã (Mã ghi "Nhiều mã…" / "Tất cả" thì LLM tự gắn)
            if rx.search(str(r.get("Mã") or "")): iss_of.setdefault(m, []).append(str(r["ID"]))
    map_state = {}
    for r in MAP:
        k = (str(r["Mã"]), str(r["Khu"])); bv = num_or_none(r.get("KL theo BV"))
        if k[0] in ngoai: continue
        bvkl[k] = bv
        map_state[k] = "CHƯA ĐỌC" if bv is None else ("KHỚP" if abs(kl.get(k, 0) - bv) < 1e-6 else "LỆCH")
    muc, ly_do, moi = {}, {}, []
    new_id = lambda: f"DT{len(moi) + 1:02d}"
    tdict = {str(r["Mã"]): r for r in T}
    for ma, tr in tdict.items():
        if ma in ngoai: muc[ma] = "CAO"; continue
        rows, why, lv = qs.get(ma, []), [], "CAO"
        by = {}
        for r in rows: by.setdefault(r.get("Trường"), []).append(r)
        if not [r for r in rows if r.get("Nguồn") in ("HÌNH", "TEXT")]: lv = worse(lv, "THAP"); why.append("chưa có quan sát từ hình/text")
        kt = [f for f in PRICE_FIELDS if f in by and not any(x.get("Trạng thái") in ("CÓ", "KHÔNG CÓ") for x in by[f]) and any(x.get("Trạng thái") == "KHÔNG THẤY" for x in by[f])]
        if kt: lv = worse(lv, "THAP"); why.append("KHÔNG THẤY: " + ", ".join(kt))
        mt = sorted({str(r.get("Trường")) for r in rows if r.get("Trạng thái") == "MÂU THUẪN"})
        mtg = [f for f in mt if f in PRICE_FIELDS and any(x.get("Nguồn") == "GEMINI" and x.get("Trạng thái") == "MÂU THUẪN" for x in by[f])]
        if mtg: lv = worse(lv, "THAP"); why.append("Claude↔Gemini lệch ở: " + ", ".join(mtg))
        if mt: lv = worse(lv, "TB"); why.append("MÂU THUẪN: " + ", ".join(mt))
        dc = tr.get("Bản vẽ đối chứng")
        if dc in (None, "") or (isinstance(dc, str) and dc.upper().startswith(("KHÔNG", "CHƯA"))): lv = worse(lv, "THAP"); why.append("thiếu Bản vẽ đối chứng")
        cd = [str(x.get("Giá trị") or "") for x in by.get("Tờ chưa đọc", [])]
        cd = [x for x in cd if x.strip() and not re.match(r"^(khong|het|none|-)", _fold(x).strip())]
        if cd: lv = worse(lv, "THAP"); why.append("còn tờ chưa đọc")
        gem = by.get("Kiểm chéo Gemini", [])
        gl = gemini_lech(gem)
        if not gem or all(x.get("Trạng thái") == "KHÔNG THẤY" for x in gem): lv = worse(lv, "TB"); why.append("Gemini chưa chạy")
        elif gl:
            lv = worse(lv, "THAP" if GIA_RX.search(_fold(gl)) else "TB"); why.append("Gemini lệch: " + gl[:80])
        st = {k[1]: s for k, s in map_state.items() if k[0] == ma}
        lech = [k for k, s in st.items() if s == "LỆCH"]; chua = [k for k, s in st.items() if s == "CHƯA ĐỌC"]
        if lech: lv = worse(lv, "TB"); why.append("KL BoQ ≠ KL BV ở " + ", ".join(lech))
        if chua: lv = worse(lv, "TB"); why.append("chưa có KL theo BV ở " + ", ".join(chua))
        muc[ma], ly_do[ma] = lv, why
        if ma not in iss_of:  # issue nháp chỉ cho mã chưa có issue nào
            lk = tr.get("Dòng BoQ gốc") or []
            if lech:
                ds = [f"{k}: BoQ {kl.get((ma, k), 0):g} ≠ BV {bvkl[(ma, k)]:g}" for k in lech]
                moi.append({"ID": new_id(), "Loại": "BoQ↔BV", "Mã": ma, "Mức": "TB", "Nội dung": "KL lệch — " + "; ".join(ds), "Nguồn": lk, "Trạng thái": "MỞ"})
            if kt:
                moi.append({"ID": new_id(), "Loại": "Thiếu tờ / KHÔNG THẤY", "Mã": ma, "Mức": "CAO" if "Kích thước BV" in kt else "TB",
                            "Nội dung": "Trường quyết định giá chưa đọc được: " + ", ".join(kt), "Nguồn": lk, "Trạng thái": "MỞ"})
            if gl:
                moi.append({"ID": new_id(), "Loại": "Claude↔Gemini", "Mã": ma, "Mức": "TB", "Nội dung": "Gemini lệch: " + gl[:200], "Nguồn": lk, "Trạng thái": "MỞ"})
            iss_of[ma] = [m_["ID"] for m_ in moi if m_["Mã"] == ma]
    t_rows, m_rows = [], []
    for ma in tdict:
        r = {"Mã": ma, "Độ tin cậy": muc[ma]}
        if iss_of.get(ma): r["Issue"] = ", ".join(sorted(set(iss_of[ma]), key=nat_key))
        if ma in ngoai:
            v = next((str(x.get("Giá trị")) for x in qs.get(ma, []) if x.get("Trường") == "Phạm vi nhà máy"), "")
            r["Giả định tính giá"] = "Không thuộc scope nhà máy (" + (v.split("–", 1)[1].strip() if "–" in v else v.strip()) + ")"
        t_rows.append(r)
    for r in MAP:
        k = (str(r["Mã"]), str(r["Khu"])); ma = k[0]
        if ma in ngoai: m_rows.append({"Mã": ma, "Khu": k[1], "Độ tin cậy": "CAO"}); continue
        s = map_state[k]
        m_rows.append({"Mã": ma, "Khu": k[1], "Trạng thái": s, "Độ tin cậy": worse(muc[ma], {"KHỚP": "CAO", "LỆCH": "TB", "CHƯA ĐỌC": "THAP"}[s]),
                       **({"Issue": ", ".join(sorted(set(iss_of[ma]), key=nat_key))} if s == "LỆCH" and iss_of.get(ma) else {})})
    write_tool(P, ra, "BOQ_TONGHOP", "doichieu", t_rows); write_tool(P, ra, "BOQ_MAP", "doichieu", m_rows); write_tool(P, ra, "ISSUES", "doichieu", moi)
    d = ra or P["data"]; os.makedirs(d, exist_ok=True)
    json.dump({"muc": muc, "ly_do": ly_do}, open(os.path.join(d, "_doichieu_tool.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    from collections import Counter
    print(f"Độ tin cậy nháp ({len(ngoai)} mã ngoài scope = CAO): {dict(Counter(v for k, v in muc.items() if k not in ngoai))}; KL từng khu: {dict(Counter(map_state.values()))}; issue nháp {len(moi)}")
    for b in kt_chuoi_sai(QS): print("  CHUỖI SAI", b)
    print("LLM còn phải làm: Mã hoàn thiện (BV↔ML), Giả định tính giá, Hành động/RFI cho issue, RFI_LOG, ASK_USER; lý do từng mã ở _doichieu_tool.json")


# ---- B6 do: tool ghi bao nhiêu ô
def cmd_do(goc, dot):
    P = paths(goc, dot)
    sheets, cols, merged, errs = load_merged(P)
    from collections import Counter
    tot = Counter()
    for sh, o in OWN.items():
        if isinstance(o, dict):
            for b in o.values(): tot[b] += 1
        else:
            for b, r in zip(o, merged[sh]): tot[b] += len([v for v in r.values() if v not in (None, "")])
    z = sum(tot.values()); a = sum(v for b, v in tot.items() if b.endswith("-tool"))
    print(f"Ô có dữ liệu (không tính khóa): {z}; tool ghi và còn nguyên: {a} ({a * 100 // max(z, 1)}%)")
    for b, v in sorted(tot.items()): print(f"  {b}: {v}")


# ---------------------------------------------------------------- main
if __name__ == "__main__":
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit(0)
    c, r = a[0], a[1:]
    if c == "kiemke": (cmd_ds if "--ds" in r[2:] else cmd_kiemke)(r[0], r[1])
    elif c == "tim": cmd_tim(r[0], r[1], r[2], r[3:])
    elif c == "ungvien": cmd_ungvien(r[0], r[1], r[2:])
    elif c == "chu": cmd_chu(r[0], r[1], r[2:])
    elif c == "so": cmd_so(r[0], r[1], r[2:])
    elif c == "anh": cmd_anh(r[0], r[1], r[2], r[3:])
    elif c == "xem": cmd_xem(r[0], r[1:])
    elif c == "mail": cmd_mail(r[0])
    elif c == "gemini": cmd_gemini(r)
    elif c == "build": sys.exit(1 if cmd_build(r[0], r[1])["E"] else 0)
    elif c == "check": cmd_check(r[0], r[1], r[2:])
    elif c == "cot": cmd_cot(r)
    elif c == "boq-tach": cmd_boq_tach(r[0], r[1], r[2:])
    elif c == "doichieu": cmd_doichieu(r[0], r[1], r[2:])
    elif c == "scope": cmd_scope(r[0], r[1], r[2:])
    elif c == "loc": cmd_loc(r[0], r[1], r[2:])
    elif c == "do": cmd_do(r[0], r[1])
    else: print(__doc__); sys.exit(2)
