"""Mô tả costing (cột AJ 'GHI CHÚ COSTING', MATRIX FITFUR-LOOSE) cho từng dòng BoQ đợt 2a — viết theo BoQ của CHÍNH dòng đó
(mã hoàn thiện, phần không bao gồm theo vị trí; KHÔNG ghi chú nội bộ RFI/giả định — user sửa 25/10), không chép từ block dùng chung (memory boq-desc-per-row).
Ra: 02_takeoff\\MO_TA_2a.xlsx (anh duyệt cột cmt) + 02_takeoff\\MO_TA_2a_lenh.txt (lệnh L.desc cho Claude trong Excel)."""
import json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout.reconfigure(encoding="utf-8")
DIR = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL\02_takeoff"
NL = "\n"
NEN = "Tấm nền MDF E1 chống ẩm - dày 17mm Sơn lót" + NL + "Khung xương gỗ sơn lót chống ẩm"
VEN = ("Tấm nền MDF E1 chống ẩm - dày 9mm 1M Sơn lót" + NL + "Tấm hoàn thiện MDF E1 chống ẩm - dày 9mm 1M Veneer WD-01 Sơn PU hoàn thiện"
       + NL + "Khung xương gỗ sơn lót chống ẩm")
KBG = lambda *x: NL + "Không bao gồm:" + "".join(NL + "+ " + i for i in x)

# (dòng BoQ, mã, mô tả, căn cứ)
D = [
    (14, "F007", "Gỗ Ash sơn PU hoàn thiện 18x80mm (màu theo WD-01)", "BoQ WD-01, cao 80; GR-TY-702 #01/#02; chung block F012"),
    (15, "F008", "Gỗ Ash sơn PU hoàn thiện 18x80mm (màu theo WD-03)",
     "BoQ WD-03 = gỗ kỹ thuật sồi (Finishes p22); user: tính như len thường; chung block F012"),
    (13, "F006", VEN,
     "BoQ WD-01; GR-TY-702 #01 thay thế; I01"),
    (30, "F022", VEN, "BoQ WD-01; PS-TY-702 #01; chung block với F006"),
    (31, "F023", "4 thanh gỗ Ash 40x140x2100mm sơn PU hoàn thiện (màu theo WD-01)", "BoQ 4nr 140x40x2100; PS-TY-702 #08"),
    (32, "F024", "Nan gỗ Ash sơn trắng hoàn thiện (cấu tạo như F009)" + NL,
     "PS-TY-701 #02 chỉ có lưới gió MEP; mô tả BoQ trùng F009"),
    (37, "F028", "Kích thước W1000xH1000 mm (tính trên 1m2 mặt phẳng trần):" + NL
     + "Khung trần gỗ: 4 thanh viền gỗ Ash 50x70mm chạy xung quanh, sơn PU hoàn thiện (màu theo WD-01)" + KBG("Tấm trần"),
     "LL-TY-701 #09; BoQ ĐVT m2 → 1 m2 mặt phẳng + 4 viền gỗ quanh chu vi (user sửa block 2026-10-25: header 1000x1000, nan 50x70 E=4 D=1000)"),
    (95, "F086", NEN + KBG("Giấy dán tường vinyl WC-01"), "giữ nguyên mô tả có sẵn của F086"),
    (97, "F088", NEN + KBG("Kim loại PVD MT-02"), "BoQ MT-02 ốp cửa thang; LL-TY-702 #03/#08"),
    (98, "F089", NEN + KBG("Gương MR-01", "Nẹp, khung kim loại MT-02"), "BoQ MR-01; LL-TY-704 #06/#10"),
    (99, "F090", NEN + KBG("Giấy dán tường WC-08", "Nẹp kim loại MT-02"), "BoQ WC-08 + MT-02; LL-TY-701 #02/#03, LL-TY-704 #13"),
    (100, "F091", NEN + KBG("Gương MR-03", "Nẹp kim loại MT-02"), "BoQ MR-03; LL-TY-703 #08 (gỗ WD-01 cabin không có dòng BoQ)"),
    (103, "F094", NEN + KBG("Giấy dán tường vinyl WC-01"), "BoQ WC-01; JS-65-702 #02"),
    (104, "F095", NEN + KBG("Giấy dán tường WC-04", "Nẹp kim loại MT-02"),
     "BoQ WC-04 + MT-02, cao 1300; PS-TY-706/707/708"),
    (105, "F096", NEN + KBG("Giấy dán tường WC-04"), "BoQ WC-04; PS-TY-702 #02"),
    (106, "F097", NEN + KBG("Giấy dán tường WC-06"), "BoQ WC-06; PS-TY-702 #02"),
    (107, "F098", NEN + KBG("Giấy dán tường WC-07", "Nẹp kim loại MT-02"), "BoQ WC-07 + MT-02 cao 1300 (BV 1260); JS-65-706/707"),
    (108, "F099", NEN + KBG("Giấy dán tường vinyl WC-03"), "BoQ WC-03; CO-TY-701 #01/#02"),
]
# F012 (dòng 20) giữ mô tả có sẵn — không ghi
assert len({d[0] for d in D}) == len(D)

out = os.path.join(DIR, "MO_TA_2a.xlsx")
if os.path.exists(out) and any(r[4] for r in openpyxl.load_workbook(out)["MO_TA"].iter_rows(min_row=2, values_only=True)):
    sys.exit("MO_TA_2a.xlsx đã có cmt của anh — không ghi đè, sửa theo cmt trước")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "MO_TA"
ws.append(["Dòng BoQ", "Mã", "Mô tả costing đề xuất (cột AJ)", "Căn cứ", "cmt"])
for j, w in enumerate([9, 8, 70, 50, 30], 1):
    c = ws.cell(1, j); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    ws.column_dimensions[c.column_letter].width = w
for d in D:
    ws.append(list(d))
    for j in range(1, 6):
        ws.cell(ws.max_row, j).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
wb.save(out)

cmd = "return await L.desc(" + json.dumps({"items": [{"row": r, "text": t} for r, _, t, _ in D if r != 95], "mark": True},
                                         ensure_ascii=False) + ")"
open(os.path.join(DIR, "MO_TA_2a_lenh.txt"), "w", encoding="utf-8").write(cmd)
print(len(D), "dòng →", out)
