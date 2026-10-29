"""Trích ngược Chiet tinh đã dựng → AI_INPUT v2 (tblAiInput + tblAiModule), để:
  1) có dữ liệu thật chạy thử check_ai_input.py;
  2) kiểm khứ hồi: AI_INPUT đã ghi vào Excel ↔ trích lại từ Excel phải trùng nhau.

Mỗi block Chiet tinh → tìm mẫu gốc (chuỗi nhãn A giống nhất) → ô khác mẫu:
  NHAP / LINK  → tblAiModule (ô nhập, ghi được bằng L.fromAiInput);
  DE_SO / SUA_CONG_THUC / XOA_CONG_THUC / dòng thêm-bớt → sheet NGOAI_EDITS (không ghi được bằng ô nhập, cần duyệt).
Công thức trỏ vào hàng trong block được viết lại thành {h+n} (n = dòng tính từ header của MẪU).

Dùng: py -I extract_ai_input.py <file CPSX> <AI_INPUT ra.xlsx> [RULES_MASTER.xlsx] [RunId]
"""
import os, sys, re, datetime, difflib, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cpsx_blocks import Book, LIB, CT, is_f, match_source, diff_block
from module_catalog import rule_ids, find_id
import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

AI_COLS = ["BlockId", "ThaoTac", "Ma", "DongBoQ", "QuanHe", "Ten", "Loai", "Mau", "MauTen", "W", "D", "H", "SL", "GiaTheo", "VatLieu", "HoanThien",
           "ThongSo", "Nguon", "DoTinCay", "Trang", "GhiChu", "RunId"]
MOD_COLS = ["BlockId", "Module", "Lan", "Off", "Cot", "GiaTri", "Loai", "Nguon", "Trang", "DoTinCay", "GhiChu"]
NGOAI_COLS = ["BlockId", "Ma", "HangExcel", "Module", "Cot", "Loai", "GiaTri", "Mau_GiaTri"]
REF = re.compile(r"(?<![!A-Za-z_'])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(!])")


def relink(f, blk, d2s):
    """Hàng trong block (không có tên sheet) → {h+n} theo dòng của mẫu; trả None nếu trỏ vào dòng thêm."""
    bad = []

    def sub(m):
        r = int(m[4])
        if blk["hdr"] <= r <= blk["end"]:
            n = d2s.get(r - blk["hdr"])
            if n is None:
                bad.append(r)
                return m[0]
            return f"{m[1]}{m[2]}{m[3]}{{h+{n}}}" if n else f"{m[1]}{m[2]}{m[3]}{{h}}"
        return m[0]
    out = REF.sub(sub, f)
    return (None, bad) if bad else (out, [])


def boq_rows(book, blk, by_hdr):
    """Dòng BoQ của block = dòng link ở A (mã) ∪ các dòng AG trong công thức E (SL)
    ∪ dòng BoQ của block khác mà E trỏ tới (block thành phần dùng chung, vd E = E1082 + E1179)."""
    f = str(book.cell(CT, blk["hdr"], "E") or "")
    rows = [blk["boq"][1]] if blk["boq"] else []
    rows += [int(x) for x in re.findall(r"!\$?AG\$?(\d+)", f) if int(x) not in rows]
    for x in re.findall(r"(?<![!A-Z$])\$?E\$?(\d+)(?!\d)", f):
        o = by_hdr.get(int(x))
        if o is not None and o is not blk:
            rows += [r for r in boq_rows(book, o, by_hdr) if r not in rows]
    return rows


def main():
    src, dst = sys.argv[1], sys.argv[2]
    ids = rule_ids(sys.argv[3]) if len(sys.argv) > 3 else {}
    run = sys.argv[4] if len(sys.argv) > 4 else "REV-" + datetime.datetime.now().strftime("%Y%m%d-%H%M")
    book = Book(src, LIB + [CT])
    # Sheet Kitchen chỉ có hiệu lực khi user cho phép dùng Matrix (R0076): mặc định KHÔNG đưa vào ứng viên đoán mẫu.
    # Bật bằng biến môi trường CPSX_ALLOW_KITCHEN=1 (hoặc tham số thứ 5 'kitchen'). Nếu không, block Template giống Kitchen
    # (tủ, ngăn kéo, ghép hộp, len chân) bị đoán nhầm sang Kitchen#19 vì biên điểm chỉ ~0.01–0.03 (SOFITEL RE00).
    allow_k = os.environ.get("CPSX_ALLOW_KITCHEN") == "1" or "kitchen" in [a.lower() for a in sys.argv[5:]]
    lib = [b for s in LIB if s != "Kitchen" or allow_k for b in book.blocks(s)]
    ai, mod, ngoai, seen = [], [], [], {}
    bq_sheet = None
    cts = book.blocks(CT)
    by_hdr = {b["hdr"]: b for b in cts}
    for blk in cts:
        code = blk["name"]
        seen[code] = seen.get(code, 0) + 1
        bid = f"{code}-{seen[code]}"
        ratio, s = match_source(book, blk, lib)
        hit = find_id(ids, s["sheet"], s["name"]) if s else None
        edits, added, removed = diff_block(book, blk, s)
        rows = boq_rows(book, blk, by_hdr)
        bq_sheet = bq_sheet or (blk["boq"][0] if blk["boq"] else None)
        # cặp dòng đích → mẫu (để viết lại tham chiếu trong block)
        d2s = {}
        la, lb = book.labels(blk), book.labels(s)
        for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, la, lb, autojunk=False).get_opcodes():
            if tag == "equal" or (tag == "replace" and i2 - i1 == j2 - j1):
                d2s.update(zip(range(i1, i2), range(j1, j2)))
        V = lambda c: book.cell(CT, blk["hdr"], c, "V")
        ai.append({"BlockId": bid, "ThaoTac": "MOI", "Ma": code, "DongBoQ": ",".join(map(str, rows)),
                   "Ten": None, "Loai": None, "Mau": hit["id"] if hit else None, "MauTen": f"{s['sheet']}!{s['name']}",
                   "W": V("B"), "D": V("C"), "H": V("D"), "SL": V("E"), "GiaTheo": V("M") if V("M") in ("md", "m2") else None, "_hdr": blk["hdr"],
                   "VatLieu": book.cell(CT, blk["hdr"] + 1, "M", "V"), "HoanThien": book.cell(CT, blk["hdr"] + 1, "N", "V"),
                   "ThongSo": book.cell(CT, blk["hdr"], "K", "V"), "Nguon": "CPSX", "DoTinCay": "CAO", "Trang": os.path.basename(src),
                   "GhiChu": f"khớp mẫu {ratio:.2f}; +{len(added)}/-{len(removed)} dòng", "RunId": run})
        for e in edits:
            if e["Loai"] in ("NHAP", "LINK"):
                v = e["GiaTri"]
                if is_f(v):
                    v2, bad = relink(v, blk, d2s)
                    if v2 is None:
                        ngoai.append({"BlockId": bid, "Ma": code, "HangExcel": blk["hdr"] + e["off_dich"], "Module": e["Module"],
                                      "Cot": e["Cot"], "Loai": "LINK_DONG_THEM", "GiaTri": v, "Mau_GiaTri": e["Mau_GiaTri"]})
                        continue
                    v = v2
                mod.append({"BlockId": bid, "Module": e["Module"], "Lan": e["Lan"], "Off": e["off"], "Cot": e["Cot"],
                            "GiaTri": v, "Loai": e["Loai"], "Nguon": "CPSX", "Trang": f"{CT}!{e['Cot']}{blk['hdr'] + e['off_dich']}",
                            "DoTinCay": "CAO", "GhiChu": None})
            else:
                ngoai.append({"BlockId": bid, "Ma": code, "HangExcel": blk["hdr"] + e["off_dich"], "Module": e["Module"],
                              "Cot": e["Cot"], "Loai": e["Loai"], "GiaTri": e["GiaTri"], "Mau_GiaTri": e["Mau_GiaTri"]})
        for r in added:
            ngoai.append({"BlockId": bid, "Ma": code, "HangExcel": r, "Module": book.cell(CT, r, "A", "V"), "Cot": "",
                          "Loai": "DONG_THEM", "GiaTri": None, "Mau_GiaTri": None})
        for r in removed:
            ngoai.append({"BlockId": bid, "Ma": code, "HangExcel": None, "Module": book.cell(s["sheet"], r, "A", "V"), "Cot": "",
                          "Loai": "DONG_BOT", "GiaTri": None, "Mau_GiaTri": f"{s['sheet']}!{r}"})
    # Tên/Loại từ BoQ; QuanHe: đơn giá BoQ (AH) cộng H của nhiều block → các block là THANH_PHAN (mỗi block mang đủ KL),
    # ngược lại nhiều block cùng 1 dòng là TACH (chia SL)
    if bq_sheet:
        bb = Book(src, [bq_sheet])
        bf = openpyxl.load_workbook(src, read_only=True)[bq_sheet]
        ah = {i: str(r[0] or "") for i, r in enumerate(bf.iter_rows(min_col=34, max_col=34, values_only=True), start=1)}
        for a in ai:
            r = int(a["DongBoQ"].split(",")[0]) if a["DongBoQ"] else None
            if r:
                a["Ten"] = bb.cell(bq_sheet, r, "C", "V")
                a["Loai"] = bb.cell(bq_sheet, r, "B", "V")
                a["QuanHe"] = "THANH_PHAN" if len(re.findall(r"Chiet tinh'?!", ah.get(r, ""))) > 1 else "TACH"
                # GiaTheo: AH trỏ N của block = giá theo md/m2 (E đã quy đổi); trỏ H = giá theo đơn vị E (= ĐVT BoQ)
                if not re.search(rf"Chiet tinh'?!\$?N\$?{a['_hdr']}(?!\d)", ah.get(r, "")):
                    a["GiaTheo"] = None
    write(dst, [("AI_INPUT", "tblAiInput", AI_COLS, ai), ("AI_MODULE", "tblAiModule", MOD_COLS, mod),
                ("NGOAI_EDITS", "tblNgoaiEdits", NGOAI_COLS, ngoai)])
    sys.stdout.reconfigure(encoding="utf-8")
    print(f"{len(ai)} block, {len(mod)} ô nhập, {len(ngoai)} ngoài ô nhập:", dict(collections.Counter(n["Loai"] for n in ngoai)))
    print("RunId", run, "→", dst)


def write(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, tname, cols, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        # công thức ghi dạng chữ để Power Query đọc nguyên văn, không tính trong file này
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"
        ref = f"A1:{get_column_letter(len(cols))}{max(2, len(rows) + 1)}"
        t = Table(displayName=tname, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "A2"
    wb.save(path)


if __name__ == "__main__":
    main()
