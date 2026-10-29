"""AI_INPUT đợt 2a SOFITEL (ván nền, len, nan, veneer) — nguồn quyết định: QUAN_SAT_SOFITEL.xlsx sheet TOM_TAT + user 2026-10-24/25.
Bản gộp (2026-10-25, user): mục cùng cấu tạo = 1 block, E header = Σ AG các dòng, đơn giá mọi dòng link H block;
cùng cấu tạo với block có sẵn → SUA block đó (gộp vào), không tạo block mới. Cần lib v2.7.
Ghi ra 02_takeoff\\AI_INPUT_batch2a_gop.xlsx (bản 18 block riêng lẻ cũ: AI_INPUT_batch2a.xlsx)."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from extract_ai_input import AI_COLS, MOD_COLS, write

sys.stdout.reconfigure(encoding="utf-8")
CPSX = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL\261020_SOFITEL_CPSX_RE00.xlsm"
OUT = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL\02_takeoff\AI_INPUT_batch2a_gop.xlsx"
RUN = "BATCH2A-GOP-20261025"
QS = "QUAN_SAT_SOFITEL.xlsx (2a-20261024)"
BQ = "'MATRIX FITFUR-LOOSE'"

# nguồn: hdr Chiet tinh RE00
SRC = {"F012": 441, "F017": 735, "F011": 303, "F009": 192, "F086": 1983}

B = []
def blk(bid, mas, src, w, d, h, edits=(), op="MOI", quan="TACH", mau="", nguon="SCAN", tin="CAO", trang="", ghi=""):
    """mas: các mã BoQ của block, mã đầu = mã block link A (SUA: mã block có sẵn)."""
    B.append(dict(bid=bid, mas=mas.split(), src=src, w=w, d=d, h=h, edits=list(edits), op=op, quan=quan, mau=mau, nguon=nguon, tin=tin,
                  trang=trang, ghi=ghi))

LEN = 'Template "Len chân, thanh lam gỗ"'
blk("F012-GOP", "F012 F007 F008", "F012", 20, 80, 1000, op="SUA", mau=LEN, trang="ID-GR-TY-702 #01/#02",
    ghi="Gộp F007 (WD-01) + F008 (WD-03, user: tính như len gỗ thường) vào block len F012 có sẵn; RFI-05 vật liệu WD-03 vẫn mở")

VEN = 'Template "Vách, trần phẳng"'
blk("F006-1", "F006 F022", "F017", 1000, 50, 1000, [("Tấm hoàn thiện", 1, 3, "N", "WD-01", "NHAP")], mau=VEN, tin="TB",
    trang="ID-GR-TY-702 #01; ID-PS-TY-702 #01",
    ghi="Vách veneer WD-01 (F006 foyer: tờ MK-TY-702 không có, RFI-01; F022 lối vào phòng ngủ PS); chép F017 (WD-04) đổi veneer → WD-01")

TRONG = 'Template "TRỐNG"'
NAN = "Nan chỉ ngang chân (40) - gỗ ASH"
blk("F023-1", "F023", "F011", 40, 40, 2100,
    [("VÁCH, TRẦN", 1, 2, "E", 4, "NHAP"), (NAN, 1, 3, "C", 140, "NHAP")], mau=TRONG, trang="ID-PS-TY-702 #08",
    ghi="4 thanh 40x140x2100 / Nr (cao 2100 theo BoQ, không thấy trên tờ)")
blk("F024-1", "F024", "F009", 1000, 25, 300, mau=TRONG, tin="TB", trang="ID-PS-TY-701 #02 (không thấy nan gỗ)",
    ghi="Chép F009 (mô tả BoQ trùng), đoạn dài 1000 = 1 m để đơn giá H theo m. Chi tiết nan chưa thấy (RFI-04)")

# F028: khung trần gỗ, đơn vị 1000x1000 = CHỈ 4 thanh viền gỗ chu vi, không có tấm nền (user 2026-10-25)
blk("F028-1", "F028", "F011", 50, 40, 4000, [(NAN, 1, 3, "C", 70, "NHAP")], mau=TRONG,
    trang="ID-LL-TY-701 #09", ghi="Khung trần gỗ: chỉ 4 thanh viền 50x70 x 1000 = 1 thanh dài 4000 cho mỗi m2, không tấm nền")

blk("F086-GOP", "F086 F088 F089 F090 F091 F094 F095 F096 F097 F098 F099", "F086", 1000, 50, 1000, op="SUA", mau=VEN,
    trang="LL-TY-701/702/703/704; JS-65-702/706/707; PS-TY-702/706/707/708; CO-TY-701",
    ghi="Gộp 10 mục ván nền MDF 17 + khung vào block F086 có sẵn; giấy/gương/kim loại ngoài phạm vi SOFITEL")

# BoQ RE00: mã → (dòng, mô tả, loại), KL
wb = openpyxl.load_workbook(CPSX, read_only=True)
boq = {str(r[0]).strip(): (i, r[2], r[1]) for i, r in enumerate(
    wb["MATRIX FITFUR-LOOSE"].iter_rows(min_row=8, max_col=7, values_only=True), start=8) if r[0]}
hdr = {i: r for i, r in enumerate(wb["Chiet tinh"].iter_rows(min_row=1, max_row=2100, max_col=5, values_only=True), start=1) if i in SRC.values()}
wb.close()
kl = {str(r[0]).strip(): r[32] for r in openpyxl.load_workbook(CPSX, read_only=True, data_only=True)["MATRIX FITFUR-LOOSE"].iter_rows(
    min_row=8, max_col=33, values_only=True) if r[0]}

ai, mod = [], []
for b in B:
    rows = [boq[m][0] for m in b["mas"]]
    if b["op"] == "SUA":  # block có sẵn phải đang link A + E về dòng của mã đầu
        h = hdr[SRC[b["src"]]]
        assert f"!A{rows[0]}" in str(h[0]).replace("$", "") and f"!AG{rows[0]}" in str(h[4]).replace("$", ""), (b["bid"], h)
    r0, ten, loai = boq[b["mas"][0]]
    ai.append({"BlockId": b["bid"], "ThaoTac": b["op"], "Ma": " ".join(b["mas"]), "DongBoQ": ",".join(map(str, rows)), "QuanHe": b["quan"],
               "Ten": ten, "Loai": loai, "Mau": b["mau"], "MauTen": f"Chiet tinh!#{SRC[b['src']]}", "W": b["w"], "D": b["d"], "H": b["h"],
               "SL": sum(kl[m] for m in b["mas"]), "GiaTheo": None, "Nguon": b["nguon"], "DoTinCay": b["tin"], "Trang": b["trang"],
               "GhiChu": f"{'sửa' if b['op'] == 'SUA' else 'chép'} {b['src']} (hdr {SRC[b['src']]}). " + b["ghi"], "RunId": RUN})
    _, sw, sd, sh, _ = hdr[SRC[b["src"]]]
    edits = [("hdr", 1, 0, c, v, "NHAP") for c, v, s in (("B", b["w"], sw), ("C", b["d"], sd), ("D", b["h"], sh)) if v != s]
    e = "=" + "+".join(f"{BQ}!AG{r}" for r in rows)
    edits.append(("hdr", 1, 0, "E", e, "LINK"))
    for m, lan, off, cot, val, typ in edits + b["edits"]:
        mod.append({"BlockId": b["bid"], "Module": m, "Lan": lan, "Off": off, "Cot": cot, "GiaTri": val, "Loai": typ,
                    "Nguon": b["nguon"], "Trang": QS, "DoTinCay": b["tin"], "GhiChu": None})

write(OUT, [("AI_INPUT", "tblAiInput", AI_COLS, ai), ("AI_MODULE", "tblAiModule", MOD_COLS, mod)])
print(len(ai), "block,", len(mod), "ô nhập →", OUT)
for a in ai:
    print(f"  {a['BlockId']:9} {a['ThaoTac']} dòng {a['DongBoQ']}  SL={a['SL']:.3f}")
