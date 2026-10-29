"""Kiểm khứ hồi sau L.fromAiInput: đọc lại file CPSX đã lưu, so với AI_INPUT.

Mỗi BlockId:
  - tìm block theo sheet ẩn _AIRUN (BlockId, ThaoTac, Hang, RunId do L.fromAiInput ghi);
  - MOI: so block mới với block nguồn (MauTen) — khác biệt phải đúng bằng các dòng tblAiModule (NHAP/LINK),
    ngoài ra không được có ô khác; link BoQ: A trỏ dòng DongBoQ đầu, đơn giá BoQ (AH) trỏ H của block mới;
  - SUA: từng ô trong tblAiModule phải có đúng giá trị/công thức đã khai;
  - không có lỗi (#REF!, #N/A…) trong block; RunId khớp.
Công thức trong AI_INPUT dùng {h}, {h+n}, {boq}, {B} được thay như applyEdits rồi so (bỏ khoảng trắng, '=+').

Dùng: py -I roundtrip_check.py <AI_INPUT.xlsx> <file CPSX đã lưu>
"""
import os, sys, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cpsx_blocks import Book, CT, LIB, norm, txt, diff_block
import openpyxl

ERR = re.compile(r"^#(N/A|NAME\?|REF!|VALUE!|DIV/0!|NUM!|NULL!|SPILL!|CALC!)")


def table(wb, name):
    for ws in wb.worksheets:
        if name in ws.tables:
            rows = list(ws[ws.tables[name].ref])
            head = [c.value for c in rows[0]]
            return [dict(zip(head, [c.value for c in r])) for r in rows[1:] if any(c.value is not None for c in r)]
    return []


def canon(v):
    s = str(txt(v) if v is not None else "").replace(" ", "")
    return s.replace("=+", "=")


def expand(v, hdr, boq_row, boq_sheet):
    if not (isinstance(v, str) and v.startswith("=")):
        return v
    q = boq_sheet if re.match(r"^[A-Za-z_][A-Za-z0-9_.]*$", boq_sheet) else "'" + boq_sheet.replace("'", "''") + "'"
    v = re.sub(r"\{h([+-]\d+)?\}", lambda m: str(hdr + int(m[1] or 0)), v)
    return v.replace("{boq}", str(boq_row)).replace("{B}", q)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ai_path, cpsx = sys.argv[1], sys.argv[2]
    wa = openpyxl.load_workbook(ai_path)
    AI, MOD = table(wa, "tblAiInput"), table(wa, "tblAiModule")
    book = Book(cpsx, LIB + [CT])
    cts = book.blocks(CT)
    # block của từng BlockId: sheet ẩn _AIRUN (L.fromAiInput ghi; bản ghi sau cùng thắng)
    by_hdr = {b["hdr"]: b for b in cts}
    notes = {}
    wr = openpyxl.load_workbook(cpsx, read_only=True)
    if "_AIRUN" in wr.sheetnames:
        for r in wr["_AIRUN"].iter_rows(min_row=2, values_only=True):
            if r[0] and r[2] in by_hdr:
                notes[r[0]] = (by_hdr[r[2]], str(r[1]), str(r[3] or ""))
    # thiếu bản ghi (lib cũ ghi comment không được) → tìm: SUA = block MauTen; MOI = block khác nguồn có A trỏ dòng BoQ đầu
    for a in AI:
        if a["BlockId"] in notes:
            continue
        op = str(a.get("ThaoTac") or "MOI").upper()
        nm = str(a["MauTen"]).partition("!")[2]
        first = str(a.get("DongBoQ") or "").split(",")[0].strip()
        if op == "SUA":
            hit = [b for b in cts if (nm[1:].isdigit() and b["hdr"] == int(nm[1:])) or norm(b["name"]) == norm(nm)]
        else:
            hit = [b for b in cts if b["boq"] and str(b["boq"][1]) == first]
        if len(hit) == 1:
            notes[a["BlockId"]] = (hit[0], op, a.get("RunId") or "")
            print(f"(không có _AIRUN cho {a['BlockId']}: tìm theo {'MauTen' if op == 'SUA' else 'link BoQ dòng ' + first} → hàng {hit[0]['hdr']})")
    boq_sheet = cts[0]["boq"][0] if cts and cts[0]["boq"] else None
    wv = openpyxl.load_workbook(cpsx, read_only=True, data_only=True)
    wf = openpyxl.load_workbook(cpsx, read_only=True)
    BQF, BQV = wf[boq_sheet], wv[boq_sheet]
    loi, ok = [], 0
    for a in AI:
        bid, op = a["BlockId"], str(a.get("ThaoTac") or "MOI").upper()
        if bid not in notes:
            loi.append(f"{bid}: không thấy block có ghi chú AI_INPUT {bid}")
            continue
        blk, op2, run = notes[bid]
        if op2 != op or (a.get("RunId") and run != a["RunId"]):
            loi.append(f"{bid}: ghi chú {op2} RunId={run} ≠ AI_INPUT {op} RunId={a.get('RunId')}")
        rows = [int(x) for x in str(a.get("DongBoQ") or "").split(",") if x.strip().isdigit()]
        mods = [m for m in MOD if m["BlockId"] == bid]
        want = {(int(m["Off"]), str(m["Cot"]).upper()): expand(m["GiaTri"], blk["hdr"], rows[0] if rows else 0, boq_sheet) for m in mods}
        # 1) ô đã khai phải có đúng giá trị
        for (off, col), v in want.items():
            got = txt(book.cell(CT, blk["hdr"] + off, col))
            if canon(got) != canon(v):
                loi.append(f"{bid}: {col}{blk['hdr'] + off} = {got!r}, AI_INPUT = {v!r}")
            else:
                ok += 1
        # 2) MOI: ngoài các ô đã khai, block mới phải giống nguồn (dịch dòng BoQ nguồn → dòng mới không tính là khác)
        if op == "MOI":
            sn, _, nm = str(a["MauTen"]).partition("!")
            if sn == CT:
                src = [b for b in cts if b is not blk and ((nm[1:].isdigit() and b["hdr"] == int(nm[1:])) or norm(b["name"]) == norm(nm))]
            else:
                src = [b for b in book.blocks(sn) if norm(b["name"]) == norm(nm)]
            if len(src) != 1:
                loi.append(f"{bid}: không xác định được block nguồn {a['MauTen']} ({len(src)})")
            else:
                edits, added, removed = diff_block(book, blk, src[0])
                extra = [e for e in edits if (e["off"], e["Cot"]) not in want]
                # SUA trên chính nguồn trong cùng lần chạy (vd D02 E) làm nguồn khác đi: bỏ qua ô đó
                sua_src = {(int(m["Off"]), str(m["Cot"]).upper()) for x in AI if str(x.get("ThaoTac")).upper() == "SUA"
                           and x["MauTen"] == a["MauTen"] for m in MOD if m["BlockId"] == x["BlockId"]}
                extra = [e for e in extra if (e["off"], e["Cot"]) not in sua_src]
                for e in extra:
                    loi.append(f"{bid}: ô {e['Cot']}{blk['hdr'] + e['off_dich']} khác nguồn mà AI_INPUT không khai ({e['Loai']}: {str(e['GiaTri'])[:60]})")
                if added or removed:
                    loi.append(f"{bid}: số dòng khác nguồn (+{len(added)}/-{len(removed)})")
        # link BoQ: A trỏ dòng đầu (cả SUA: block có sẵn vẫn giữ dòng của nó); AH mọi dòng khai trỏ H block (gộp nhiều mục cùng cấu tạo)
        if rows:
            a_f = str(book.cell(CT, blk["hdr"], "A") or "")
            if not re.search(rf"!\$?A\$?{rows[0]}(?!\d)", a_f):
                loi.append(f"{bid}: A{blk['hdr']} = {a_f} không trỏ BoQ dòng {rows[0]}")
            for r in rows:
                ah = str(BQF[f"AH{r}"].value or "")
                if not re.search(rf"Chiet tinh'?!\$?H\$?{blk['hdr']}(?!\d)", ah):
                    loi.append(f"{bid}: BoQ AH{r} = {ah} không trỏ H{blk['hdr']}")
        # 3) lỗi công thức trong block (giá trị đã tính khi lưu)
        errs = [f"{chr(64 + j)}{r}" for r in range(blk["hdr"], blk["end"] + 1) for j in range(1, 19)
                if isinstance(book.cell(CT, r, j, "V"), str) and ERR.match(book.cell(CT, r, j, "V"))]
        if errs:
            loi.append(f"{bid}: lỗi công thức {errs[:8]}")
        h, e = book.cell(CT, blk["hdr"], "H", "V"), book.cell(CT, blk["hdr"], "E", "V")
        bq = [(r, BQV[f"AG{r}"].value, BQV[f"AH{r}"].value) for r in rows]
        print(f"{bid:7} {op} block {blk['hdr']}-{blk['end']}  E={e}  H={h:,.0f}  BoQ (dòng, KL, đơn giá)={bq}")
    print(f"\n{ok} ô khớp AI_INPUT, {len(loi)} lỗi")
    for x in loi:
        print("  LỖI", x)
    return 1 if loi else 0


if __name__ == "__main__":
    sys.exit(main())
