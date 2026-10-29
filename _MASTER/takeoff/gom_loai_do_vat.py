"""Gom 84 mục chưa tính (SOFITEL RE00) theo LOẠI ĐỒ VẬT. Cột Mẫu để trống, user điền 1 lần/loại.
Loại gán theo tên BoQ + ĐVT (chưa đọc bản vẽ). Chỉ điền sẵn Mẫu ở loại user đã xác nhận (NHAN-DANG-MAU §0)."""
import sys, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, r"D:\COSTING EXCEL\_MASTER\takeoff")
from cpsx_blocks import Book

sys.stdout.reconfigure(encoding="utf-8")
RE00 = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL\261020_SOFITEL_CPSX_RE00.xlsm"
BOSUNG = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL\02_takeoff\BO_SUNG_RE00_v2.xlsx"
OUT = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL\02_takeoff\LOAI_DO_VAT_SOFITEL.xlsx"

# (loại, [mã chưa tính], [mã RE00 đã làm cùng loại], mẫu đã xác nhận, ghi chú)
LOAI = [
    ("Len chân gỗ/sơn", "F007 F008", "F012 F013 F020 F021 F027", 'Template "Len chân, thanh lam gỗ"', "user xác nhận"),
    ("Chỉ gỗ / trim / thanh nan rời", "F023 F026.1", "F011 F011.1 F025 F026", 'Template "TRỐNG"', "user xác nhận; F026.1 thực ra đã có AH→349 ở RE00, không làm"),
    ("Vách nan che điều hòa", "F024", "F009", 'Template "TRỐNG"', "user xác nhận"),
    ("Ốp veneer gỗ lên tường (m2)", "F006 F022 F028", "F010 F014 F017", "", "F028 phải mở LL-TY-701 trước; F006 dẫn tờ MK-TY-702 không có (RFI)"),
    ("Vách sơn / vách CNC", "F018 F019", "F052 F084", "", ""),
    ("Ốp trang trí cột", "F093", "", "", "chưa có block cùng loại"),
    ("Ván nền dán giấy (chỉ tính ván + khung)", "F090 F094 F095 F096 F097 F098 F099", "F015 F085 F086", "", "giấy WC ngoài phạm vi SOFITEL"),
    ("Ván nền ốp gương / kim loại", "F088 F089 F091", "F082", "", "gương MR, kim loại PVD ngoài phạm vi; F088 làm lại đợt 1"),
    ("Gương ốp có cánh mở", "F092", "F082", "", "Nr, có cánh hai bên"),
    ("Vách TV / TV panel có hốc", "F053 F054 F055 F056 F057", "F029 F030", "", "F053/F054 làm lại đợt 1; F055–F057 giấy WC-05 + WD-04"),
    ("Vách góc (corner screen)", "F058", "F016 F031 F032 F033 F034", "", ""),
    ("Đầu giường", "F064", "F035 F037", "", ""),
    ("Đầu giường + vách góc", "F062 F063", "F035 F031", "", "ghép đầu giường + vách góc (T1)"),
    ("Vách gỗ đứng / fins portal", "F059 F060 F061", "", "", "chưa có block cùng loại"),
    ("Tủ áo", "F065 F066 F067", "F038 F039 F040", 'Template "Tủ áo"', "user xác nhận; user bảo bỏ qua F065–F067"),
    ("Hốc ủi đồ (iron niche)", "F071", "F038 F039", "", "tủ áo F038/F039 có hốc ủi"),
    ("Tủ minibar cao (1900)", "F042 F072 F073", "", "", "chưa có block cùng loại ở RE00"),
    ("Tủ minibar thấp (900)", "F074", "F041", "", "REV01 có block F074, user chốt giữ cách bóc"),
    ("Vanity", "F075 F076 F077 F078", "F043", "", "đá ST ngoài phạm vi"),
    ("Bedside liền tường (FITFUR)", "F047 F079 F080", "F044 F045 F046 F048", "", ""),
    ("Tủ phòng khách + kệ", "F068 F069", "", "", "chưa có block cùng loại"),
    ("Hốc tường nhà tắm", "F070", "", "", "gương sọc MR ngoài phạm vi"),
    ("Khung bồn tắm / ghế tắm", "F050 F051 F081", "", "", "đá ngoài phạm vi: còn khung gỗ?"),
    ("Bàn rời (bàn ăn, side, lamp table)", "L005 L026 L027 L032", "L001 L002 L015 L021", "", ""),
    ("Ghế bành / lounge / armchair", "L012 L035 L037 L039 L041", "L010 L016", "", "L012/L041 ngoài trời"),
    ("Đôn / stool / ottoman", "L034 L038", "", "", "L034 từng chép từ ghế L016 → #N/A, đã xoá"),
    ("Băng ghế", "L040 L043", "", "", "L043 kèm bàn phụ"),
    ("Gương rời", "L042", "F082", "", ""),
    ("Cửa gỗ đi (swing/1 cánh/lùa gỗ)", "F005 D08 D09", "D01 D03 D07 D10 D11 D12", "Sheet Door", "user xác nhận cửa→Door; F005 ngoài trời dày 40"),
    ("Cửa lùa kính khung gỗ", "D15 D16 D17 D18", "D02", "Sheet Door", "user xác nhận cửa→Door"),
    ("Cửa kỹ thuật (riser)", "D19", "D04 D05 D06", "Sheet Door", "user xác nhận cửa→Door; mẫu cụ thể chưa rõ"),
    ("Cửa có sẵn làm lại mặt", "D20 D21 D22 D23 D24 D25 D26", "", "", "bỏ giấy/nhãn 3M thì còn gì tính? (câu hỏi mở)"),
]

# BoQ RE00: mã → (dòng, mô tả, ĐVT, hàng block AH)
wb = openpyxl.load_workbook(RE00, read_only=True)
boq = {}
for i, r in enumerate(wb["MATRIX FITFUR-LOOSE"].iter_rows(min_row=8, max_col=34, values_only=True), start=8):
    if r[0]:
        m = re.search(r"!\$?[A-Z]+\$?(\d+)", str(r[33] or ""))
        boq[str(r[0]).strip()] = (i, str(r[2] or "").strip(), str(r[6] or "").strip(), int(m[1]) if m else None)
wb.close()

bs = {}
for r in list(openpyxl.load_workbook(BOSUNG, read_only=True)["BO_SUNG"].iter_rows(values_only=True))[1:]:
    if r[0]:
        bs[str(r[0]).strip()] = r[11]

# kiểm: mọi mã trong BO_SUNG có đúng 1 loại; mã đã làm đều có AH
seen = {}
for ten, todo, done, _, _ in LOAI:
    for c in todo.split():
        assert c not in seen, c
        seen[c] = ten
    for c in done.split():
        assert boq[c][3], f"{c} chưa có AH"
thieu = set(bs) - set(seen)
thua = set(seen) - set(bs)
assert not thieu and not thua, (thieu, thua)
chua_ah_ngoai_ds = [c for c, v in boq.items() if not v[3] and c not in bs]

# danh sách tên block mẫu cho dropdown (Template + Door, không Kitchen)
book = Book(RE00, ["Template", "Door"])
maus = [f'{s} "{b["name"]}"' for s in ("Template", "Door") for b in book.blocks(s)]
maus = list(dict.fromkeys(maus))

out = openpyxl.Workbook()
H = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="305496")
Y = PatternFill("solid", fgColor="FFF2CC")
G = PatternFill("solid", fgColor="E2EFDA")
WR = Alignment(wrap_text=True, vertical="top")


def head(ws, cols, widths):
    ws.append(cols)
    for j, w in enumerate(widths, 1):
        c = ws.cell(1, j)
        c.font, c.fill, c.alignment = H, HF, WR
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


ws = out.active
ws.title = "LOAI"
head(ws, ["STT", "Loại đồ vật", "Số mục", "Mã chưa tính", "ĐVT", "Block RE00 cùng loại (mã @ hàng)",
          "Mẫu (anh điền)", "Trạng thái mẫu", "Ghi chú", "cmt"],
     [5, 34, 7, 30, 8, 36, 40, 14, 48, 30])
for k, (ten, todo, done, mau, note) in enumerate(LOAI, 1):
    codes = todo.split()
    dvt = ", ".join(sorted({boq[c][2] for c in codes}))
    ref = "; ".join(f"{c}@{boq[c][3]}" for c in done.split())
    ws.append([k, ten, len(codes), " ".join(codes), dvt, ref, mau,
               "ĐÃ XÁC NHẬN" if mau else "CHỜ ĐIỀN", note, None])
    for j in range(1, 11):
        ws.cell(k + 1, j).alignment = WR
    ws.cell(k + 1, 7).fill = G if mau else Y
n = len(LOAI) + 1

ws2 = out.create_sheet("MUC")
head(ws2, ["Mã", "Dòng BoQ (RE00)", "Hạng mục", "ĐVT", "Loại đồ vật", "Mẫu (tra từ LOAI)", "Bộ tờ phải đọc",
           "Nguồn gán loại"], [8, 9, 70, 6, 34, 40, 40, 30])
for i, c in enumerate(sorted(seen, key=lambda x: (x[0] != "F", x[0] != "L", x)), 2):
    ws2.append([c, boq[c][0], boq[c][1], boq[c][2], seen[c], None, bs.get(c), "tên BoQ + ĐVT, chưa đọc bản vẽ"])
    ws2.cell(i, 6).value = f'=IFERROR(INDEX(LOAI!$G$2:$G${n},MATCH(E{i},LOAI!$B$2:$B${n},0)),"")'
    ws2.cell(i, 3).alignment = WR

ws3 = out.create_sheet("_MAU")
for m in maus:
    ws3.append([m])
ws3.sheet_state = "hidden"
dv = DataValidation(type="list", formula1=f"=_MAU!$A$1:$A${len(maus)}", allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv)
dv.add(f"G2:G{n}")

ws4 = out.create_sheet("GHI_CHU")
for line in [
    "Mục đích: anh điền cột G (Mẫu) 1 lần cho mỗi loại; sheet MUC tự tra mẫu cho từng mã.",
    "Loại gán theo tên BoQ + ĐVT, CHƯA đọc bản vẽ. Sai loại → ghi ở cột cmt.",
    "Cột G có dropdown tên block Template/Door của RE00 (không Kitchen); gõ tay vẫn được.",
    "Ô xanh = quy tắc anh đã xác nhận (NHAN-DANG-MAU §0); ô vàng = chờ điền. Không gợi ý mẫu khi không có căn cứ.",
    "Block RE00 cùng loại = mã đã có AH trỏ về block (để chép, ThaoTac MOI từ Chiet tinh).",
    f"Mã RE00 chưa có AH nhưng KHÔNG có trong BO_SUNG_v2: {', '.join(chua_ah_ngoai_ds)} — cần anh xem.",
]:
    ws4.append([line])
ws4.column_dimensions["A"].width = 120

out.save(OUT)
print(len(LOAI), "loại;", len(seen), "mục;", len(maus), "mẫu dropdown")
print("chưa AH ngoài DS:", chua_ah_ngoai_ds)
print("chưa có block cùng loại:", [t for t, _, d, _, _ in LOAI if not d])
