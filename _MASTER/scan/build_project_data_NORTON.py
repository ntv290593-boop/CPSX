import json,csv,re,collections
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
D='/tmp/claude-0/norton/'
IN='..\\00_input\\261023_D01\\'
pages=json.load(open(D+'pages.json'));_pp=json.load(open(D+'sched_pp.json'));sched=[dict(r,file=f,no=next((x['no'] for x in pages if x['file']==f and x['page']==r['page']),'')) for f,v in _pp.items() for r in v]
spec=json.load(open(D+'spec0814.json'));boq=[b for b in json.load(open(D+'boq.json')) if b['type']]
imap=list(csv.DictReader(open('/mnt/project-files/NORTON/_intake_map_D01.csv',encoding='utf-8-sig')))
imap.append({'OldRelPath':'NORTON_ Báo giá hạng mục cửa gỗ.msg','NewRelPath':'04_TaiLieuKhac\\NORTON_ Báo giá hạng mục cửa gỗ.msg','Ext':'.msg','SizeBytes':''})
pdfrel={x['NewRelPath'].split('\\')[-1][:-4]:x['NewRelPath'] for x in imap if x['Ext']=='.pdf'}
dwg={}
for x in imap:
  if x['Ext'].lower()=='.dwg':
    m=re.search(r'Sheet - ([A-Z]{2}-[\dA-Z]+-[\dA-Z]+)',x['NewRelPath'])
    if m: dwg[(x['NewRelPath'].split('\\')[1],m.group(1).replace('-','.'))]=x['NewRelPath']
dwgfold={'SB_TENDER_PLOT B_GENERAL, BASEMENT_2026 10 08':'DWG_GEN-BASEMENT','SB_TENDER_PLOT B_BLK 01_2026 10 08':'DWG_BLK01'}
GEN='SB_TENDER_PLOT B_GENERAL, BASEMENT_2026 10 08'
B1,B2,B3='SB_TENDER_PLOT B_BLK 01_2026 10 08','SB_TENDER_PLOT B_BLK 02_2026 10 08','SB_TENDER_PLOT B_BLK 03_2026 10 08'
SPEC='02_Spec\\DPA_Section_08_14_00_Wood_Doors.pdf'
MLX='02_Spec\\Skybridge_Overall Material List_2029 09 29__2.xlsx'
BOQF='01_BoQ\\BOQ-Cửa gỗ.xlsx'
WAR='02_Spec\\Yeu cau bao hanh.pdf'
H=Font(bold=True,color='FFFFFF');HF=PatternFill('solid',fgColor='1F4E78');LK=Font(color='0563C1',underline='single')
wb=Workbook();wb.remove(wb.active)
def sheet(name,cols,widths=None):
  ws=wb.create_sheet(name);ws.append(cols)
  for c in ws[1]: c.font=H;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center')
  ws.freeze_panes='A2'
  for i,w in enumerate(widths or [],1): ws.column_dimensions[get_column_letter(i)].width=w
  return ws
# ---- link registry: (label, relpath from input, page, subaddress)
def short(f): return f.replace('SB_TENDER_PLOT B_','').replace('_2026 10 08','')
SHEETS=json.load(open(D+'sheetmap.json'))
def P(f,p,lab=None):
  n=next((x['no'] for x in pages if x['file']==f and x['page']==p),'')
  lab=lab or f"{short(f)} p{p}"+(f" {n.split('.',3)[-1].strip()}" if n else '')
  sh=SHEETS.get(f,{}).get(str(p))
  return (lab,'03_BanVe\\'+sh,'','') if sh else (lab,pdfrel[f],p,'')
def PG(f,p): return (f"↗ File gốc {short(f)} p{p}",pdfrel[f],p,'')
def SP(p,lab): return (f"08 14 00 p{p} {lab}",SPEC,p,'')
def MLD(row,lab): return (f"ML DOORS r{row} {lab}",MLX,'',f"DOORS!B{row}")
def MLI(ref): n=int(ref.split('.')[1]); return (f"ML {ref}",MLX,'',f"IRONMONGERY (STEEL & WOOD)!B{n+4}")
def BQ(s,r): return (f"BoQ {s}!{r}",BOQF,'',f"{s}!A{r}")
def FL(rel,lab=None,p=''): return (lab or rel.split('\\')[-1],rel,p,'')
REG=[]
def setlinks(ws,r,c,items,text=None):
  items=[i for i in items if i]
  if not items: return
  if not any(i[0].startswith('↗') for i in items):
    sh=next((i for i in items if i[1].startswith('03_BanVe\\_sheets')),None)
    if sh:
      src=next(k for k,v in SHEETS.items() if sh[1][len('03_BanVe\\'):] in v.values())
      pn=int(next(k for k,v in SHEETS[src].items() if v==sh[1][len('03_BanVe\\'):]))
      items=items+[PG(src,pn)]
  cell=ws.cell(r,c)
  nreal=sum(not i[0].startswith('↗') for i in items)
  cell.value=text or (items[0][0] if nreal==1 else f"{items[0][0]}  ▾{nreal}")
  cell.hyperlink=Hyperlink(ref=cell.coordinate,location=f"'{ws.title}'!{cell.coordinate}",display=str(cell.value));cell.font=LK
  for it in items: REG.append((ws.title,cell.coordinate)+it)
SCHED={'BLK01':(B1,81),'BLK02':(B2,93),'BLK03':(B3,97),'BASEMENT':(GEN,101)}
MLDOOR={'DH114':(4,'8.1 EI30'),'ID-DH112':(5,'8.2 phòng ngủ'),'DH117':(6,'WPB'),'DH118':(6,'WPB'),'ID-DH111':(6,'WPB')}
DET=[P(GEN,54,'GEN p54 CD.00.68 chi tiết cửa gỗ'),P(GEN,55,'GEN p55 CD.00.68A')]
SPECS=[SP(6,'§3.02.A cánh ≥45mm, nẹp 24mm'),SP(6,'§3.02.B cấu tạo'),SP(7,'§3.02.C loại cửa'),SP(5,'§2.02 vật liệu')]
# ---------- PROJECT
ws=sheet('PROJECT',['Trường','Giá trị','Nguồn','→ OverView'],[28,70,40,14])
rows=[('Tên dự án','TỔ HỢP THƯƠNG MẠI SB – Chung cư khu B (NORTON) / Skybridge Plot B',[P(GEN,1,'Bìa GEN p1')],''),
 ('CĐT','Skybridge Company Limited (Gamuda Land)',[FL(WAR,'Bảo hành p1',1)],''),
 ('Tư vấn KT','DPA (design) / ICU (local architect)',[P(GEN,54,'Title block GEN p54')],''),
 ('Gói thầu','Cửa gỗ + phụ kiện cửa (Doors & Ironmongery)',[BQ('Sum',1)],''),
 ('Người yêu cầu','Phạm Quốc Bảo (Tender) – 22/10/2026',[FL('04_TaiLieuKhac\\NORTON_ Báo giá hạng mục cửa gỗ.msg','Email .msg')],''),
 ('Hạn nộp giá','26/10/2026',[FL('04_TaiLieuKhac\\NORTON_ Báo giá hạng mục cửa gỗ.msg','Email .msg')],''),
 ('Đợt hồ sơ','261023_D01 (198 file + 1 .msg)',[],''),
 ('Bảo hành','Cửa gỗ: 5 năm sau TOC (mục 6); phụ kiện/kim khí: 5 năm (mục 10); liên đới CĐT–Tổng thầu–Nhà SX, ký dạng Deed',[FL(WAR,'Bảo hành p1 mục 6',1),FL(WAR,'Bảo hành p2 mục 10',2)],''),
 ('Thương hiệu cửa','An Cường / Sunwood / Woodsland',[MLD(4,'8.1'),MLD(5,'8.2'),MLD(6,'WPB')],''),
 ('Hoàn thiện','Laminate An Cường (8.1, WPB); Melamine Yên Lâm/An Cường (8.2)',[MLD(4,'8.1'),MLD(5,'8.2')],'C (HT)'),
 ('Phụ kiện','Assa Abloy / Häfele / Dormakaba; khóa điện tử Häfele/Adel/Philips 912.20.693 (PC rate)',[MLI('9.01'),MLI('9.02')],'D45'),
 ('Chống cháy','DH114 EI30 (lõi rockwool 40kg/m3, cánh 46mm, khung 38mm)',[P(B1,81),MLD(4,'8.1')],'C39–C51'),
 ('Chống ẩm','DH117/118/ID-DH111: WPB 6mm chịu nước',[MLD(6,'WPB'),BQ('B4.6',23)],'C40')]
for i,(a,b,lk,o) in enumerate(rows,2):
  ws.append([a,b,'',o]);setlinks(ws,i,3,lk)
ws.cell(8,3,'_intake_map_D01.csv').hyperlink='_intake_map_D01.csv';ws.cell(8,3).font=LK
# ---------- FILES
ws=sheet('FILES',['ID','Loại','Tên file','Link','Đường dẫn gốc (trong zip)','Ext','MB','Trạng thái scan'],[7,14,60,8,60,7,8,30])
cat={'01_BoQ':'BoQ','02_Spec':'Spec','03_BanVe':'Bản vẽ','04_TaiLieuKhac':'Tài liệu khác'}
for i,x in enumerate(imap,1):
  p=x['NewRelPath'];fn=p.split('\\')[-1];mb=round(int(x['SizeBytes'])/1e6,1) if x['SizeBytes'] else ''
  st='DWG – link theo trang ở PAGES' if x['Ext'].lower()=='.dwg' else {'.pdf':'Đã trích text','.xlsx':'Đã đọc'}.get(x['Ext'].lower(),'Đã đọc')
  if 'TENDER_PLOT B_2026' in fn: st='CHƯA ĐỌC (523MB, cần tách trên máy)'
  if 'bao hanh' in fn: st='Đọc bằng ảnh (font lỗi)'
  ws.append([f'F{i:03d}',cat[p.split('\\')[0]],fn,'',x['OldRelPath'],x['Ext'],mb,st]);setlinks(ws,i+1,4,[FL(p,'Mở')])
ws.auto_filter.ref=ws.dimensions
# ---------- PAGES
ws=sheet('PAGES',['File','Trang','Số BV','Tên BV','Ver','Mã cửa xuất hiện (SL)','Mở PDF','DWG'],[30,6,28,50,6,45,10,8])
for r,x in enumerate(pages,2):
  f=x['file']
  ws.append([short(f),x['page'],x['no'],(x['name'] or '').replace('Ě','Đ'),x['ver'],', '.join(f'{k}({v})' for k,v in sorted(x['doors'].items())),'',''])
  setlinks(ws,r,7,[P(f,x['page'],f"p{x['page']}")])
  if x['no'] and f in dwgfold:
    k=(dwgfold[f],'.'.join(x['no'].split('.')[-3:]))
    if k in dwg: setlinks(ws,r,8,[FL(dwg[k],'DWG')])
ws.auto_filter.ref=ws.dimensions
# ---------- BOQ_CDT
ws=sheet('BOQ_CDT',['Sheet','Dòng','Bill','Khu','Loại','Item','Mã cửa','Mô tả','ĐV','KL CĐT','Mô tả kỹ thuật','Link'],[8,6,8,16,8,6,10,60,5,8,80,14])
for i,b in enumerate(boq,2):
  ws.append([b['sheet'],b['row'],b['bill'],b['zone'],b.get('kind') or '',b['item'],b['type'],b['desc'],b['unit'],b['qty'],b['spec'] or '',''])
  setlinks(ws,i,12,[BQ(b['sheet'],b['row'])])
ws.auto_filter.ref=ws.dimensions
# ---------- common
S=collections.defaultdict(dict);DUP=[]
for s in sched:
  blk={'BLK 01':'BLK01','BLK 02':'BLK02','BLK 03':'BLK03'}.get(next((k for k in ['BLK 01','BLK 02','BLK 03'] if k in s['file']),''),'BASEMENT')
  if blk in S[s['type']]:
    DUP.append((s['type'],blk,S[s['type']][blk]['count'],s['count'],s['file'],s['page']))
    S[s['type']][blk]=dict(S[s['type']][blk],count=S[s['type']][blk]['count']+s['count'],conf=0.5)
  else: S[s['type']][blk]=s
blkzone=lambda z:'BASEMENT' if z=='BASEMENT' else z.split()[-1]
leaf=[b for b in boq if b['desc'].startswith('Single leaf')]
tot=collections.defaultdict(lambda:collections.defaultdict(int))
for b in leaf:
  if isinstance(b['qty'],(int,float)): tot[b['type']][blkzone(b['zone'])]+=b['qty']
iss={'DH114':'I01,I02,I03,I06,I11','DH117':'I01,I02,I03','DH118':'I01,I02,I03','ID-DH111':'I01,I02,I03','ID-DH112':'I01,I02,I03,I04,I05,I07'}
def schedlinks(t): return [P(S[t][b]['file'],S[t][b]['page'],f"{short(S[t][b]['file'])} p{S[t][b]['page']} bảng {S[t][b]['table']} dòng {S[t][b]['row']}") for b in ['BASEMENT','BLK01','BLK02','BLK03'] if b in S[t]]
def bqleaf(t): return [BQ(b['sheet'],b['row']) for b in leaf if b['type']==t]
# ---------- BOQ_MAP
ws=sheet('BOQ_MAP',['Mã cửa','Khối','KL BoQ (podium+tower)','KL bảng thống kê BV','Lệch','Kích thước (BV)','EI','Trang thống kê','Chi tiết cửa','Spec','Material List','Issue','Trạng thái','Độ tin cậy'],[10,10,12,12,8,14,7,30,30,30,22,18,10,9])
r=1
for t in ['DH114','DH117','DH118','ID-DH111','ID-DH112']:
  for blk in ['BASEMENT','BLK01','BLK02','BLK03']:
    q=tot[t].get(blk);s=S[t].get(blk)
    if q is None and not s: continue
    r+=1;sc=s['count'] if s else None
    ws.append([t,blk,q,sc,(sc-q) if (q is not None and sc is not None) else 'n/a',f"{s['W']}x{s['H']}" if s else '',s['EI'] if s else '','','','','',iss[t],
      'KHỚP' if (q==sc) else ('THIẾU' if (q is None or sc is None) else 'LỆCH'),(s['conf'] if s else 0.5)])
    if s: setlinks(ws,r,8,[P(s['file'],s['page'],f"{short(s['file'])} p{s['page']} bảng {s['table']} dòng {s['row']}")])
    setlinks(ws,r,9,DET);setlinks(ws,r,10,SPECS);setlinks(ws,r,11,[MLD(*MLDOOR[t])])
ws.auto_filter.ref=ws.dimensions
# ---------- BOQ_TONGHOP
ZC=['BASEMENT','PODIUM BLK01','TOWER BLK01','PODIUM BLK02','TOWER BLK02','PODIUM BLK03','TOWER BLK03']
G=collections.OrderedDict()
for b in boq:
  k=(b['type'],re.sub(r'\s+',' ',b['desc']).strip())
  g=G.setdefault(k,{'b':b,'z':collections.defaultdict(int),'rows':[],'na':[]})
  g['rows'].append((b['sheet'],b['row']))
  if isinstance(b['qty'],(int,float)): g['z'][b['zone']]+=b['qty']
  else: g['na'].append(b['zone'])
pos={'DH114':'Cửa chính căn hộ','DH117':'WC tầng hầm','DH118':'Podium','ID-DH111':'WC căn hộ','ID-DH112':'Phòng ngủ căn hộ'}
cols=['STT','Nhóm','Mã cửa','Vị trí','Mô tả (BoQ)','Mô tả kỹ thuật (BoQ)','Kích thước BV','EI','Ref ML','ĐV']+ZC+['Tổng BoQ','KL theo BV (cánh)','Lệch','Dòng BoQ gốc','Thống kê BV','Chi tiết cửa','Spec','Material List','Issue','Giả định tính giá','Độ tin cậy']
ws=sheet('BOQ_TONGHOP',cols,[5,9,10,16,48,70,11,6,8,5]+[9]*len(ZC)+[9,10,7,22,26,30,30,24,18,40,9])
bvtot={t:sum(s['count'] for s in S[t].values()) for t in S}
assume={'DH114':'Khóa điện tử PC rate; phụ kiện bổ sung ML chào option','ID-DH112':'Laminate theo BoQ','DH117':'Theo BoQ','DH118':'Theo BoQ','ID-DH111':'Theo BoQ'}
for i,(k,g) in enumerate(G.items(),1):
  b=g['b'];t=b['type'];isleaf=b['desc'].startswith('Single leaf')
  tq=sum(g['z'].values());bv=bvtot.get(t) if isleaf else None
  sc=next(iter(S[t].values()),None)
  m=None if isleaf else re.search(r'Ref\.?\s*(9\.\d\d)',b['desc'])
  row=[i,'Cửa' if isleaf else 'Phụ kiện',t,pos.get(t,''),k[1],b['spec'] or '',f"{sc['W']}x{sc['H']}" if (sc and isleaf) else '',(sc['EI'] or '') if (sc and isleaf) else '',(m.group(1) if m else (MLDOOR[t][1].split()[0] if isleaf else '')),b['unit']]
  row+=[g['z'][z] if z in g['z'] else ('N/A' if z in g['na'] else '') for z in ZC]
  row+=[tq,bv if bv is not None else '',(bv-tq) if bv is not None else '','','','','','',iss.get(t,''),assume.get(t,''),(min([x['conf'] for x in S[t].values()] or [0.5]) if isleaf else 0.9)]
  ws.append(row);r=i+1;c0=10+len(ZC)
  setlinks(ws,r,c0+4,[BQ(*x) for x in g['rows']])
  if isleaf:
    setlinks(ws,r,c0+5,schedlinks(t));setlinks(ws,r,c0+6,DET);setlinks(ws,r,c0+7,SPECS);setlinks(ws,r,c0+8,[MLD(*MLDOOR[t])])
  else:
    setlinks(ws,r,c0+7,[SP(2,'§1.02.B dẫn 08 70 00'),SP(6,'§3.02.A.7 gia công lắp PK')])
    if m: setlinks(ws,r,c0+8,[MLI(m.group(1))])
  for c in (5,6): ws.cell(r,c).alignment=Alignment(wrap_text=True,vertical='top')
ws.auto_filter.ref=ws.dimensions;ws.freeze_panes='F2'
# ---------- MATERIALS
ws=sheet('MATERIALS',['Mã cửa','Vị trí','Cánh','Lõi','Mặt','Khung','Hoàn thiện','Phụ kiện BoQ','Nguồn'],[10,22,10,16,18,30,26,60,34])
for i,row in enumerate([
 ('DH114','Cửa chính căn hộ','46mm EI30','Rockwool 40kg/m3','Tấm CC 6mm + MDF CA 3mm','38mm','Laminate An Cường','Khóa điện tử 4 chức năng (PC rate)'),
 ('DH117','WC tầng hầm','39mm','Honeycomb 25mm','WPB 6mm x2','45mm (17 gỗ ghép + 6 WPB), chặn 17 WPB, nẹp L 14','Laminate, nẹp PVC','Bản lề, tay co âm, lock case, tay gạt, WC cylinder, kick plate, chặn cửa'),
 ('DH118','Podium BLK01–03','39mm','Honeycomb 25mm','WPB 6mm x2','45mm','Laminate, nẹp PVC','Như DH117'),
 ('ID-DH111','WC căn hộ','39mm','Honeycomb 25mm','WPB 6mm x2','45mm','Laminate, nẹp PVC','Bản lề, lock case đen, cylinder+thumbturn, tay gạt, chặn cửa tường'),
 ('ID-DH112','Phòng ngủ căn hộ','37mm','Honeycomb 25mm','MDF CA 6mm x2','45mm dày 29mm','Laminate (BoQ) / Melamine (ML 8.2)','Như ID-DH111')],2):
  ws.append(list(row)+['']);t=row[0]
  setlinks(ws,i,9,bqleaf(t)[:1]+schedlinks(t)[:1]+[MLD(*MLDOOR[t])]+DET[:1])
# ---------- ISSUES
I=[('I01','Spec↔BoQ/BV','Tất cả trừ DH114','Spec §3.02.A.3: cánh ≥45mm "unless otherwise shown in the Drawings"; BV/BoQ/ML ghi 39mm, 37mm → BV ưu tiên, xin xác nhận','TB','RFI-01',[SP(6,'§3.02.A.3'),BQ('B8.9',29),P(GEN,54)]),
('I02','Spec↔BoQ','Tất cả cửa gỗ','Spec §3.02.A.5: nẹp cạnh gỗ cứng đặc ≥24mm; BoQ + ML: nẹp PVC (ML 8.2: PVC ≥1mm)','CAO','RFI-01',[SP(6,'§3.02.A.5'),BQ('B8.9',29),MLD(5,'8.2')]),
('I03','Spec↔BoQ','Tất cả cửa gỗ','Spec §3.02.B/C: khung gỗ cứng + plywood 3 lớp ≥6mm, veneer (§2.02.B); BoQ/ML: honeycomb + MDF/WPB + laminate','CAO','RFI-01',[SP(6,'§3.02.B'),SP(7,'§3.02.C'),SP(5,'§2.02.B'),BQ('B8.9',29)]),
('I04','ML↔BoQ','ID-DH112','Hoàn thiện: ML 8.2 = Melamine Yên Lâm/An Cường; BoQ = Laminate','TB','RFI-02',[MLD(5,'8.2'),BQ('B8.9',33)]),
('I05','ML nội bộ','ID-DH112','ML 8.2: khóa & phụ kiện cửa phòng ngủ phải đạt EI30, trong khi cửa không chống cháy','TB','RFI-02',[MLD(5,'8.2'),P(B1,81)]),
('I06','ML↔BoQ','DH114','BoQ chỉ có khóa điện tử 9.01; ML yêu cầu thêm cho cửa chính căn hộ: bản lề nặng 9.02, chặn sàn 9.08, mắt thần 9.11, thanh chắn bụi 9.12, kick plate 9.13','CAO','RFI-03',[MLI('9.02'),MLI('9.08'),MLI('9.11'),MLI('9.12'),MLI('9.13'),BQ('B8.10',29)]),
('I07','ML↔BoQ','ID-DH112','BoQ: chặn cửa gắn tường 9.09 (ML ghi cho cửa phòng tắm); ML: hít cửa nam châm 9.10 cho cửa phòng ngủ','THẤP','RFI-03',[MLI('9.09'),MLI('9.10'),BQ('B8.10',57)]),
('I08','BV↔BoQ','DH331','DH331 850x2200 có trong thống kê/mặt bằng (BLK01 175 / BLK02 189 / BLK03 149) nhưng không có chi tiết, không có trong BoQ','CAO','RFI-04',[P(B1,81),P(B2,93),P(B3,97),P(B1,84),P(B2,96)]),
('I09','BoQ nội bộ','Tất cả','BoQ dẫn "Bill Nr. 4.11/5.11" cho phụ kiện, thực tế ở Bill 4.7/5.8…','THẤP','—',[BQ('B4.6',23),BQ('B4.7',29)]),
('I10','Thiếu hồ sơ','Phụ kiện','Section 08 70 00 Ironmongery được dẫn chiếu nhưng không có trong hồ sơ','TB','RFI-05',[SP(2,'§1.02.B')]),
('I11','Giả định','DH114','Khóa điện tử là PC rate → ghi rõ giá tạm tính/loại trừ','TB','—',[BQ('B8.10',29),MLI('9.01')]),
('I12','Hồ sơ chưa đọc','—','SB_TENDER_PLOT B_2026 10 08.pdf (523MB, Vol2) chưa quét','THẤP','—',[FL('03_BanVe\\SB_TENDER_PLOT B_2026 10 08.pdf')]),
('I13','Đối chiếu KL','Tất cả','KL BoQ (podium+tower) = bảng thống kê CD.66.00 từng khối → KHỚP','OK','—',[P(B1,81),P(B2,93),P(B3,97),P(GEN,101)])]
ST={'I01':('LỆCH','Gửi RFI-01; báo giá theo BV/BoQ, ghi giả định lệch spec',0.9),'I02':('LỆCH','Gửi RFI-01; tính nẹp PVC, chào option nẹp gỗ cứng',1.0),'I03':('LỆCH','Gửi RFI-01; tính theo BoQ/ML',1.0),
'I04':('LỆCH','Gửi RFI-02; tính Laminate, chào giảm giá nếu Melamine',1.0),'I05':('LỆCH','Gửi RFI-02; phụ kiện thường',0.9),'I06':('THIẾU','Gửi RFI-03; chào option phụ kiện bổ sung',1.0),
'I07':('LỆCH','Gửi RFI-03; tính theo BoQ',0.9),'I08':('THIẾU','Gửi RFI-04; loại trừ trong báo giá',1.0),'I09':('INFO','Không cần xử lý',1.0),'I10':('THIẾU','Gửi RFI-05',1.0),
'I11':('INFO','Ghi giả định PC rate trong báo giá',1.0),'I12':('THIẾU','Tách Vol2 nếu cần',1.0),'I13':('KHỚP','—',1.0)}
for k,(t,blk,a,b,f,pg) in enumerate(DUP,14):
  I.append((f'I{k}','BV nội bộ',t,f'Bảng thống kê {blk} có 2 dòng {t} ({a} và {b}) → cộng {a+b}; ngoài gói cửa gỗ, chỉ ghi nhận','INFO','—',[P(f,pg)]))
  ST[f'I{k}']=('INFO','Không thuộc gói; ghi nhận',0.5)
ws=sheet('ISSUES',['ID','Trạng thái','Loại','Mã cửa','Nội dung','Hành động','Mức','RFI','Độ tin cậy','Nguồn (bấm chọn)'],[6,9,14,16,80,40,7,8,9,34])
order={'LỆCH':0,'THIẾU':1,'INFO':2,'KHỚP':3}
for i,x in enumerate(sorted(I,key=lambda x:(order[ST[x[0]][0]],x[0])),2):
  st,act,cf=ST[x[0]]
  ws.append([x[0],st,x[1],x[2],x[3],act,x[4],x[5],cf,''])
  setlinks(ws,i,10,x[6])
  for c in (5,6): ws.cell(i,c).alignment=Alignment(wrap_text=True,vertical='top')
ws.auto_filter.ref=ws.dimensions
# ---------- RFI_LOG
R=[('RFI-01','Cấu tạo cửa gỗ: Spec 08 14 00 (cánh ≥45mm, nẹp gỗ cứng 24mm, khung gỗ cứng + plywood/veneer) khác BoQ/ML (39/37mm honeycomb, MDF/WPB, laminate, nẹp PVC). Xin xác nhận áp dụng mô tả BoQ/ML.','I01–I03'),
('RFI-02','ID-DH112: hoàn thiện Laminate (BoQ) hay Melamine (ML 8.2)? Phụ kiện có cần EI30 không khi cửa không chống cháy?','I04–I05'),
('RFI-03','Phụ kiện DH114 (bản lề nặng 9.02, chặn sàn 9.08, mắt thần 9.11, thanh chắn bụi 9.12, kick plate 9.13) và hít cửa 9.10 cho ID-DH112 theo ML không có trong BoQ – có thuộc gói?','I06–I07'),
('RFI-04','Cửa DH331 (850x2200, ~513 bộ) có trên thống kê/mặt bằng nhưng không có chi tiết/BoQ – có thuộc gói cửa gỗ?','I08'),
('RFI-05','Xin bổ sung Section 08 70 00 Ironmongery.','I10')]
ws=sheet('RFI_LOG',['RFI','Câu hỏi','Issue','Ngày gửi','Trả lời','Trạng thái'],[8,100,10,10,40,10])
for x in R: ws.append(list(x)+['','','Nháp'])
for c in ws['B'][1:]: c.alignment=Alignment(wrap_text=True,vertical='top')
# ---------- SPEC (tự tách 08 14 00 theo PART/§)
ws=sheet('SPEC',['Mục','Tiêu đề','Nội dung','Nhãn','Trang','Mở'],[10,26,100,18,6,16])
for i,x in enumerate(spec,2):
  ws.append([x['ref'],x['title'],x['text'],', '.join(x['tags']),x['page'],''])
  setlinks(ws,i,6,[SP(x['page'],'§'+x['ref'])])
  ws.cell(i,3).alignment=Alignment(wrap_text=True,vertical='top')
ws.auto_filter.ref=ws.dimensions
# ---------- NOTES (tự lấy mục submittal/qa từ SPEC + ghi chú tay)
ws=sheet('NOTES',['#','Ghi chú','Loại','Nguồn','Độ tin cậy'],[4,100,12,30,9])
nts=[(f"§{x['ref']} {x['title']} {x['text']}".strip(),'/'.join(x['tags']),[SP(x['page'],'§'+x['ref'])],0.8) for x in spec if set(x['tags'])&{'submittal','qa'} and x['ref'].count('.')>=2]
nts+=[('Bảo hành 5 năm sau TOC (cửa gỗ mục 6, phụ kiện mục 10), liên đới Nhà SX, ký Deed','warranty',[FL(WAR,'Bảo hành p1',1),FL(WAR,'Bảo hành p2',2)],0.9),('Mặt cắt gỗ/MDF tại vị trí phụ kiện phải quét lớp bảo vệ; chân khung chống thấm đến H=300','material',[MLD(4,'8.1'),MLD(6,'WPB')],1.0)]
for i,(n,t,lk,cf) in enumerate(nts,2):
  ws.append([i-1,n,t,'',cf]);setlinks(ws,i,4,lk);ws.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top')
# ---------- ASK_USER
ws=sheet('ASK_USER',['#','Câu hỏi nội bộ','Trả lời'],[4,100,30])
for i,q in enumerate(['Báo giá theo BoQ (39/37mm honeycomb) + ghi chú lệch spec, hay chào thêm phương án theo spec?','DH331 tính hay loại trừ?','Phụ kiện bổ sung theo ML (I06) chào option riêng?','Có cần tách đọc file Vol2 523MB không?'],1): ws.append([i,q,''])
# ---------- _LINKS
ws=sheet('_LINKS',['Key','Sheet','Ô','Nhãn','File (tương đối)','Trang','Ô đích','Mở (dự phòng)','Độ tin cậy'],[22,14,7,40,60,6,30,10,9])
for j,(s,c,lab,rel,pg,sub) in enumerate(REG,2):
  ws.append([f"{s}!{c}",s,c,lab,IN+rel,pg,sub,'',1.0 if (pg or sub) else 0.8])
  x=ws.cell(j,8,'Mở');x.hyperlink=IN+rel+(('#'+sub) if sub else '');x.font=LK
# ---------- DASHBOARD
ws=wb.create_sheet('DASHBOARD',0)
ws['A1']='NORTON – Skybridge Plot B · Cửa gỗ + phụ kiện';ws['A1'].font=Font(bold=True,size=14)
info=[('Hạn nộp giá','26/10/2026'),('Đợt','261023_D01'),('File',len(imap)),('Trang BV đã quét',len(pages)),('Bảng thống kê (pdfplumber)',f'{len(sched)} dòng'),('Dòng BoQ',len(boq)),('Hạng mục BoQ (đã gộp)',len(G)),('Issue',f"{len(I)} (LỆCH {sum(v[0]=='LỆCH' for v in ST.values())} · THIẾU {sum(v[0]=='THIẾU' for v in ST.values())} · INFO {sum(v[0]=='INFO' for v in ST.values())})"),('Mục spec 08 14 00',len(spec)),('RFI nháp',len(R)),('KL BoQ ↔ BV','Khớp 100% (5 mã cửa)'),('Link','Cần add-in SRC_LINKS: bấm → mở PDF đúng trang (Edge) / xlsx đúng ô; ô có ▾ → chọn link')]
for i,(a,b) in enumerate(info,3): ws.cell(i,1,a).font=Font(bold=True);ws.cell(i,2,b)
for i,s in enumerate(wb.sheetnames[1:],3):
  c=ws.cell(i,4,s);c.hyperlink=Hyperlink(ref=c.coordinate,location=f"'{s}'!A1",display=s);c.font=LK
ws.column_dimensions['A'].width=22;ws.column_dimensions['B'].width=40;ws.column_dimensions['D'].width=14
wb.move_sheet('BOQ_TONGHOP',-(wb.sheetnames.index('BOQ_TONGHOP')-1))
wb.save(D+'out/PROJECT_DATA.xlsx')
print(wb.sheetnames,len(REG))
