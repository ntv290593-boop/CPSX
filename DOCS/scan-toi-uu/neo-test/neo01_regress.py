# -*- coding: utf-8 -*-
"""Kiểm hồi quy `pd.py neo` (N3.1/N3.2) — gọi thẳng hàm trong pd.py, không chép logic. Chỉ đọc dự án, ghi vào thư mục tạm.
A. Bench cache Blanca (_cache\\Blanca\\_data): kết quả gốc + đột biến link tờ (trang ±1 cùng file) -> tỉ lệ bắt.
B. SOFITEL\\01_scan v8 (user cho phép đo 29/10): D1d qua neo_run — nhãn cột P (Bản vẽ đối chứng) -> kỳ vọng 199/199 CÓ NEO;
   đột biến trang ±1 -> tỉ lệ bắt.
Chạy: py -I -X utf8 neo01_regress.py   (vài phút: B đọc text PDF bằng pymupdf)"""
import sys, os, re, json, copy, random, shutil, tempfile, collections
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "SKILL", "01-scan", "tool"))
import pd
random.seed(5)


def shift_links(merged, pages_of, kieu="to"):
    """Bản sao merged: mỗi link tờ (PDF, có trang) đổi sang trang ±1 trong cùng file. -> (merged mới, số link đã đổi)."""
    m, n = copy.deepcopy(merged), 0
    for sh, rows in m.items():
        for r in rows:
            for cot, v in r.items():
                ok = (sh, cot) in pd.NEO_TO or (sh == "QUAN_SAT" and cot == "Tờ · Chi tiết" and r.get("Nguồn") in ("HÌNH", "TEXT"))
                if not ok or not isinstance(v, (list, dict)): continue
                ls = v.get("links", []) if isinstance(v, dict) else v
                for l in ls:
                    if isinstance(l, list) and len(l) >= 3 and str(l[2]).isdigit() and l[1] in pages_of:
                        q = int(l[2]) + random.choice((-1, 1))
                        if 1 <= q <= pages_of[l[1]]: l[2] = q; n += 1
    return m, n


def run(P, merged, ten):
    res = pd.neo_run(P, merged)
    print(f"[{ten}]", pd.neo_tong(res))
    return res


# ---------- A. Blanca bench cache
tmp = tempfile.mkdtemp(); d = os.path.join(tmp, "01_scan", "_data", "X")
shutil.copytree(os.path.join(HERE, "..", "_cache", "Blanca", "_data"), d)
P = pd.paths(tmp, "X")
_, _, M, _ = pd.load_merged(P)
pages_of = {f[:-5]: len(json.load(open(os.path.join(d, "_text", f), encoding="utf-8"))) for f in os.listdir(os.path.join(d, "_text"))}
r0 = run(P, M, "A gốc")
M2, n = shift_links(M, pages_of)
r1 = run(P, M2, f"A đột biến trang ±1 ({n} link)")
print(f"   -> bắt {r1['tong'].get('LỆCH', 0) - r0['tong'].get('LỆCH', 0)} / {r0['tong'].get('CÓ NEO', 0)} link CÓ NEO gốc "
      "(link spec nằm trong tổng nhưng không bị đổi)")
lo = collections.Counter(re.sub(r"\d", "#", l["nhan"])[:20] for l in r1["lech"])
shutil.rmtree(tmp, ignore_errors=True)

# ---------- B. SOFITEL v8, cột P
import openpyxl, pymupdf
ROOT = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL"; SCAN = ROOT + r"\01_scan"
wb = openpyxl.load_workbook(SCAN + r"\PROJECT_DATA_SOFITEL_v8.xlsx", read_only=True)
links = collections.defaultdict(list)
for key, sh, cell, lab, f, pg, _ in wb["_LINKS"].iter_rows(min_row=2, values_only=True):
    if key and key.startswith("BOQ_TONGHOP!P") and f and pg and str(f).lower().endswith(".pdf"):
        links[key].append((lab, os.path.normpath(os.path.join(SCAN, f)), int(pg)))
so_bv = {}  # (tên file không đuôi, trang) -> Số BV
for r in wb["PAGES"].iter_rows(min_row=2, values_only=True):
    if r[0] and r[1] and r[2]: so_bv[(str(r[0]).rsplit(".", 1)[0], int(r[1]))] = str(r[2]).strip()
tmp = tempfile.mkdtemp(); P = pd.paths(tmp, "X"); os.makedirs(os.path.join(P["data"], "_text"))
P["inp"] = ROOT
fid, FILES, PAGES = {}, [], []
for v in links.values():
    for _, f, _ in v:
        if f in fid or not os.path.isfile(f): continue
        k = fid[f] = f"F{len(fid) + 1:03d}"
        doc = pymupdf.open(f)
        json.dump([p.get_text() for p in doc], open(os.path.join(P["data"], "_text", k + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
        FILES.append({"ID": k, "Đường dẫn": os.path.relpath(f, ROOT)})
        stem = os.path.basename(f).rsplit(".", 1)[0]
        for i in range(1, doc.page_count + 1):
            PAGES.append({"File ID": k, "Trang": i, "Số BV": so_bv.get((stem, i), "")})
T = [{"Mã": key, "Bản vẽ đối chứng": [[lab, fid[f], pg] for lab, f, pg in v if f in fid]} for key, v in links.items()]
M = {"FILES": FILES, "PAGES": PAGES, "BOQ_TONGHOP": T}
print(f"\nB: {sum(len(t['Bản vẽ đối chứng']) for t in T)} link cột P, {len(FILES)} file PDF, {sum(1 for p in PAGES if p['Số BV'])} trang có Số BV")
r0 = run(P, M, "B gốc (kỳ vọng D1d 199/199)")
for l in r0["lech"][:8]: print("   ", pd.neo_dong(l))
pages_of = {x["ID"]: sum(1 for p in PAGES if p["File ID"] == x["ID"]) for x in FILES}
M2, n = shift_links(M, pages_of)
r1 = run(P, M2, f"B đột biến trang ±1 ({n} link)")
print(f"   -> bắt {r1['tong'].get('LỆCH', 0)} / {n}")
for l in r1["lech"][:3]: print("   ", pd.neo_dong(l))
shutil.rmtree(tmp, ignore_errors=True)
