# -*- coding: utf-8 -*-
"""Đo N3.3 (mã VL + kích thước đọc trên bản vẽ) của `pd.py neo` — gọi thẳng hàm pd.py, không chép logic. Chỉ đọc, ghi vào thư mục tạm.
A. Bench cache Blanca (_cache\\Blanca\\_data): bản vẽ có mã VL (GL1, MT2…) và kích thước nằm trong lớp text.
   Gốc · đột biến 1 chữ số (kích thước) / đổi số trong mã (mã VL) · tráo giá trị giữa 2 mã khác nhau · dời trang ±1.
   Chỉ tính trên các ô CÓ NEO ở bản gốc.
B. SOFITEL\\01_scan v8 (user cho phép đo 29–30/10): bản vẽ GA Design vẽ mã/kích thước bằng nét -> kỳ vọng KHÔNG KIỂM ĐƯỢC toàn bộ,
   0 LỆCH. Mã: cột H 'Mã trên BV (đọc hình)' theo từng tờ. Kích thước: chuỗi R×S×C trong 'Mô tả (BoQ)' -> link cột P.
Chạy: py -I -X utf8 neo01_n33.py"""
import sys, os, re, json, copy, random, shutil, tempfile, collections, glob
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "SKILL", "01-scan", "tool"))
import pd
random.seed(5)

KIEU = ("ma", "kt")


def run(P, M, ten):
    r = pd.neo_run(P, M)
    print(f"[{ten}]", "; ".join(f"{k}: {r['theo_kieu'].get(k, {})}" for k in KIEU))
    return r


def lech_keys(r):
    return {(l["sh"], l["dong"], l["cot"]) for l in r["lech"] if l["cot"] in ("Mã hoàn thiện trên BV", "Kích thước BV")}


def neo_ok(P, M):
    """Các ô N3.3 ra CÓ NEO ở bản gốc: chạy từng ô một lần (ô = (sheet, dòng, cột giá trị))."""
    items = [x for x in pd.neo_items(M, {r["ID"]: os.path.join(P["inp"], r["Đường dẫn"]) for r in M["FILES"]}, P["inp"]) if x["kieu"] in KIEU]
    r = pd.neo_run(P, M); bad = lech_keys(r)
    ok = []
    for x in items:
        if (x["sh"], x["dong"], x["cot"]) in bad or not x["ma"]: continue
        if one(P, M, x, M[x["sh"]][x["dong"] - 2]).get("CÓ NEO"): ok.append(x)  # CÓ NEO thật: chạy lại riêng ô này
    return items, ok


def cot_gt(x): return "Giá trị" if x["sh"] == "QUAN_SAT" else "Kích thước BV"
def cot_link(x): return {"QUAN_SAT": "Tờ · Chi tiết", "BOQ_TONGHOP": "Bản vẽ đối chứng", "BOQ_MAP": "Tờ chi tiết"}[x["sh"]]


def one(P, M, x, row):
    m1 = {k: v for k, v in M.items() if k in ("FILES", "PAGES")}
    m1[x["sh"]] = [row]
    # giữ nguồn họ mã như bản gốc (neo_ho đọc cả QUAN_SAT/BOQ_TONGHOP/SPEC)
    m1["SPEC"] = [{"Mã liên quan": " ".join(f"{h}1 {h}2" for h in x["ho"])}]
    return pd.neo_run(P, m1)["theo_kieu"].get(x["kieu"], {})


def dot_bien(v, kieu, ho):
    if kieu == "kt":
        ns = pd.neo_kt_gt(v); n = random.choice(ns); s = n.split(".")[0]; i = random.randrange(len(s))
        dg = random.choice([c for c in "0123456789" if c != s[i] and not (i == 0 and c == "0")])
        return re.sub(r"(?<!\d)" + re.escape(n) + r"(?!\d)", s[:i] + dg + s[i + 1:], v, count=1)
    ms = list(pd.neo_mavl(v, ho).values()); m = random.choice(ms)
    a, b, c = pd.NEO_MAVL.fullmatch(m).groups()
    return v.replace(m, f"{a}{'-' if '-' in m else ''}{int(b) + random.choice((1, 2, 3, 7))}{c}", 1)


def bai(P, M, items, ok, pages_of, ten):
    print(f"\n== {ten}: {len(items)} ô N3.3, {len(ok)} ô CÓ NEO dùng cho đột biến")
    for k in KIEU:
        xs = [x for x in ok if x["kieu"] == k]
        if not xs: continue
        # 1. đột biến chữ số
        b = n = 0
        for x in xs:
            for _ in range(3):
                row = copy.deepcopy(M[x["sh"]][x["dong"] - 2]); row[cot_gt(x)] = dot_bien(row[cot_gt(x)], k, set(x["ho"]))
                t = one(P, M, x, row); n += 1; b += bool(t.get("LỆCH"))
        # 2. tráo giá trị với ô khác mã
        b2 = n2 = 0
        for x in xs:
            oth = [y for y in xs if y["ten"].split(" · ")[0] != x["ten"].split(" · ")[0]
                   and not {tuple(l[:2]) for l in y["nhom"]} & {tuple(l[:2]) for l in x["nhom"]}]
            if not oth: continue
            y = random.choice(oth); row = copy.deepcopy(M[x["sh"]][x["dong"] - 2])
            row[cot_gt(x)] = M[y["sh"]][y["dong"] - 2][cot_gt(y)]
            t = one(P, M, x, row); n2 += 1; b2 += bool(t.get("LỆCH"))
        # 3. dời trang ±1 (mọi link của ô)
        b3 = n3 = 0
        for x in xs:
            row = copy.deepcopy(M[x["sh"]][x["dong"] - 2]); v = row[cot_link(x)]
            ls = v.get("links", []) if isinstance(v, dict) else v; doi = 0
            for l in ls:
                if isinstance(l, list) and len(l) >= 3 and str(l[2]).isdigit() and l[1] in pages_of:
                    q = int(l[2]) + random.choice((-1, 1))
                    if 1 <= q <= pages_of[l[1]]: l[2] = q; doi += 1
            if not doi: continue
            t = one(P, M, x, row); n3 += 1; b3 += bool(t.get("LỆCH"))
        print(f"  {k}: đột biến bắt {b}/{n} ({100 * b // max(n, 1)}%) · tráo giá trị bắt {b2}/{n2} · dời trang ±1 bắt {b3}/{n3}")


# ---------- A. Blanca
tmp = tempfile.mkdtemp(); d = os.path.join(tmp, "01_scan", "_data", "X")
shutil.copytree(os.path.join(HERE, "..", "_cache", "Blanca", "_data"), d)
P = pd.paths(tmp, "X")
_, _, M, _ = pd.load_merged(P)
print("họ mã dự án (Blanca):", sorted(pd.neo_ho(M)))
pages_of = {f[:-5]: len(json.load(open(os.path.join(d, "_text", f), encoding="utf-8"))) for f in os.listdir(os.path.join(d, "_text"))}
r0 = run(P, M, "A gốc")
print("  KKĐ:", {k: v for k, v in r0["khong_kiem"].items() if "mã" in k or "kích thước" in k or "R×S×C" in k})
for l in r0["lech"]:
    if l["cot"] in ("Mã hoàn thiện trên BV", "Kích thước BV"): print("  ", pd.neo_dong(l))
items, ok = neo_ok(P, M)
bai(P, M, items, ok, pages_of, "A Blanca")
shutil.rmtree(tmp, ignore_errors=True)

# ---------- B. SOFITEL v8 (kỳ vọng KKĐ, 0 LỆCH)
import openpyxl, pymupdf
ROOT = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL"; SCAN = ROOT + r"\01_scan"
wb = openpyxl.load_workbook(SCAN + r"\PROJECT_DATA_SOFITEL_v8.xlsx", read_only=True)
links = collections.defaultdict(list)
for key, sh, cell, lab, f, pg, *_ in wb["_LINKS"].iter_rows(min_row=2, values_only=True):
    if key and key.startswith("BOQ_TONGHOP!P") and f and pg and str(f).lower().endswith(".pdf"):
        links[int(key.split("!P")[1])].append((lab, os.path.normpath(os.path.join(SCAN, f)), int(pg)))
so_bv = {}
for r in wb["PAGES"].iter_rows(min_row=2, values_only=True):
    if r[0] and r[1] and r[2]: so_bv[(str(r[0]).rsplit(".", 1)[0], int(r[1]))] = str(r[2]).strip()
tmp = tempfile.mkdtemp(); P = pd.paths(tmp, "X"); os.makedirs(os.path.join(P["data"], "_text")); P["inp"] = ROOT
fid, FILES, PAGES = {}, [], []


def F(f):
    if f not in fid:
        k = fid[f] = f"F{len(fid) + 1:03d}"; doc = pymupdf.open(f)
        json.dump([p.get_text() for p in doc], open(os.path.join(P["data"], "_text", k + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
        FILES.append({"ID": k, "Đường dẫn": os.path.relpath(f, ROOT)})
        stem = os.path.basename(f).rsplit(".", 1)[0]
        for i in range(1, doc.page_count + 1): PAGES.append({"File ID": k, "Trang": i, "Số BV": so_bv.get((stem, i), "")})
    return fid[f]


to_pg = {}  # số tờ -> link
for v in links.values():
    for lab, f, pg in v:
        if os.path.isfile(f):
            s = so_bv.get((os.path.basename(f).rsplit(".", 1)[0], pg))
            if s: to_pg.setdefault(s, [lab, F(f), pg])
SH = re.compile(r"ID-[A-Z0-9]{1,4}-[A-Z0-9]{1,4}-\d{3}")
QS, TH = [], []
for i, r in enumerate(wb["BOQ_TONGHOP"].iter_rows(min_row=2, values_only=True), 2):
    ma, h, mota = r[2], r[7], r[4]
    for seg in str(h or "").split(";"):
        m = SH.search(seg)
        if m and m.group(0) in to_pg and seg[m.end():].strip():
            QS.append({"Mã": ma, "Trường": "Mã hoàn thiện trên BV", "Nguồn": "HÌNH", "Trạng thái": "CÓ",
                       "Giá trị": seg[m.end():].strip(" :-"), "Tờ · Chi tiết": [to_pg[m.group(0)]]})
    ls = [[lab, F(f), pg] for lab, f, pg in links.get(i, []) if os.path.isfile(f)]
    if ls and pd.neo_kt_gt(mota): TH.append({"Mã": ma, "Kích thước BV": mota, "Bản vẽ đối chứng": ls})
M = {"FILES": FILES, "PAGES": PAGES, "QUAN_SAT": QS, "BOQ_TONGHOP": TH}
print(f"\nB SOFITEL: {len(QS)} ô mã (cột H theo tờ), {len(TH)} ô kích thước (Mô tả BoQ → cột P), {len(FILES)} PDF; họ mã: {sorted(pd.neo_ho(M))}")
r0 = run(P, M, "B gốc (kỳ vọng KKĐ, 0 LỆCH)")
print("  KKĐ:", r0["khong_kiem"])
for l in r0["lech"]:
    if l["cot"] in ("Mã hoàn thiện trên BV", "Kích thước BV"): print("  ", pd.neo_dong(l))
shutil.rmtree(tmp, ignore_errors=True)
