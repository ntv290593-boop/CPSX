# -*- coding: utf-8 -*-
"""Plan NEO 3b-N2 (01-scan), chỉ đọc. User cho phép dùng SOFITEL\\01_scan (đợt costing cũ) để ĐO thuật toán (29/10).
D1: giá trị LLM đọc hình trong PROJECT_DATA_SOFITEL_v8 (BOQ_TONGHOP H 'Mã trên BV', P 'Bản vẽ chi tiết', Q 'Chi tiết')
    có neo trên lớp text trang được link không; + trùng ngẫu nhiên + đột biến.
D3: số kích thước DWG (đáp án _dwg_tangA) có trong lớp text trang PDF tách sẵn (_sheets) không.
Chạy: py -I -X utf8 neo01_d1d3.py"""
import os, re, json, glob, random, collections, unicodedata
import openpyxl, pymupdf

ROOT = r"D:\COSTING EXCEL\_ACTIVE\SOFITEL"
SCAN = ROOT + r"\01_scan"
random.seed(11)

def fold(s):
    s = unicodedata.normalize("NFKC", str(s or ""))
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").upper()

CODE = re.compile(r"(?<![A-Z0-9])([A-Z]{1,4})\s*[-._]?\s*0*(\d{1,3})(?![0-9])")
def codes(s): return {a + b for a, b in CODE.findall(fold(s))}
def alnum(s): return re.sub(r"[^A-Z0-9]", "", fold(s))

_doc, _txt = {}, {}
def page_text(path, pg):
    k = (path, pg)
    if k not in _txt:
        if path not in _doc: _doc[path] = pymupdf.open(path)
        d = _doc[path]
        _txt[k] = d[pg - 1].get_text() if 1 <= pg <= d.page_count else None
    return _txt[k]

wb = openpyxl.load_workbook(SCAN + r"\PROJECT_DATA_SOFITEL_v8.xlsx", read_only=True)
links = collections.defaultdict(list)
for key, sh, cell, lab, f, pg, _ in wb["_LINKS"].iter_rows(min_row=2, values_only=True):
    if f and pg: links[key].append((lab, os.path.normpath(os.path.join(SCAN, f)), int(pg)))
# Số BV -> (file, trang) từ PAGES + FILES
files = {r[2].rsplit(".", 1)[0]: r[4] for r in wb["FILES"].iter_rows(min_row=2, values_only=True) if r[2]}
sheet_pg = {}
for r in wb["PAGES"].iter_rows(min_row=2, values_only=True):
    fn, pg, so = r[0], r[1], r[2]
    if fn and pg and so and fn in files:
        p = os.path.join(ROOT, files[fn], fn + ".pdf") if not files[fn].lower().endswith(".pdf") else os.path.join(ROOT, files[fn])
        if not os.path.isfile(p):  # FILES 'Đường dẫn' có thể là thư mục
            cand = glob.glob(os.path.join(ROOT, files[fn].split("\\")[0], "**", fn + ".pdf"), recursive=True)
            p = cand[0] if cand else p
        sheet_pg.setdefault(str(so).strip(), (p, int(pg)))
print(f"PAGES: {len(sheet_pg)} số tờ có file/trang; file mở được: {sum(os.path.isfile(v[0]) for v in sheet_pg.values())}")

T = list(wb["BOQ_TONGHOP"].iter_rows(min_row=2, values_only=True))
SH = re.compile(r"ID-[A-Z0-9]{1,4}-[A-Z0-9]{1,4}-\d{3}")

def kkd(t): return t is None or len(t.strip()) < 20

# ---------- D1a: mã đọc hình (cột H) theo từng tờ
res = collections.Counter(); items = []; lech_ex = []
for r in T:
    ma, h = r[2], r[7]
    if not h: continue
    for seg in re.split(r";", str(h)):
        m = SH.search(seg)
        if not m: continue
        sh = m.group(0); body = seg[m.end():]
        cs = {c for c in codes(body) if not c.startswith("ID")}
        if not cs: continue
        if sh not in sheet_pg: res["không có trang cho tờ"] += len(cs); continue
        p, pg = sheet_pg[sh]
        t = page_text(p, pg) if os.path.isfile(p) else None
        if kkd(t): res["KHÔNG KIỂM ĐƯỢC"] += len(cs); continue
        pc = codes(t)
        for c in cs:
            ok = c in pc
            res["CÓ NEO" if ok else "LỆCH"] += 1
            items.append((ma, sh, c, p, pg))
            if not ok and len(lech_ex) < 12: lech_ex.append((ma, sh, c))
print("\n=== D1a mã đọc hình (H) trên trang của tờ ===", dict(res))
for e in lech_ex: print("  LỆCH", e)

# ---------- D1b: trùng ngẫu nhiên — mã của tờ A kiểm trên trang tờ B khác
pages = sorted({(p, pg) for *_, p, pg in items})
hit = n = 0
for ma, sh, c, p, pg in items:
    for _ in range(3):
        q = random.choice(pages)
        if q == (p, pg): continue
        n += 1; hit += c in codes(page_text(*q))
print(f"=== D1b trùng ngẫu nhiên: mã đúng tờ A lại có trên tờ B: {hit}/{n} ({100 * hit / max(n, 1):.0f}%)")
fam = collections.Counter(re.match(r"[A-Z]+", c).group(0) for *_, c, _p, _g in [(i[0], i[1], i[2], i[3], i[4]) for i in items])
print("   họ mã:", dict(fam.most_common(10)))

# ---------- D1c: đột biến số trong mã (WD-01 -> WD-0x), có bắt không
cau = tot = 0
for ma, sh, c, p, pg in items:
    pre, num = re.match(r"([A-Z]+)(\d+)", c).groups()
    for k in random.sample([x for x in range(1, 13) if str(x) != num], 3):
        tot += 1; cau += (pre + str(k)) not in codes(page_text(p, pg))
print(f"=== D1c đột biến mã (đổi số): bắt {cau}/{tot} ({100 * cau / max(tot, 1):.0f}%)")

# ---------- D1d: số tờ trong nhãn link cột P có trên trang link (khung tên)
c1 = collections.Counter(); ex = []
for i, r in enumerate(T, start=2):
    for lab, p, pg in links.get(f"BOQ_TONGHOP!P{i}", []):
        m = SH.search(str(lab))
        if not m: continue
        t = page_text(p, pg) if os.path.isfile(p) else None
        if kkd(t): c1["KHÔNG KIỂM ĐƯỢC"] += 1; continue
        ok = alnum(m.group(0)) in alnum(t)
        c1["CÓ NEO" if ok else "LỆCH"] += 1
        if not ok and len(ex) < 8: ex.append((r[2], lab, os.path.basename(p), pg))
print("=== D1d số tờ của nhãn P trên trang link ===", dict(c1))
for e in ex: print("  LỆCH", e)
# trùng ngẫu nhiên số tờ: số tờ A trên trang B (khung tên thường có cả danh mục tờ liên quan?)
lp = [(SH.search(str(l)).group(0), p, g) for v in links.values() for l, p, g in v if SH.search(str(l)) and os.path.isfile(p)]
lp = [x for x in lp if not kkd(page_text(x[1], x[2]))]
h2 = n2 = 0
for s, p, g in random.sample(lp, min(300, len(lp))):
    q = random.choice(lp)
    if q[0] == s: continue
    n2 += 1; h2 += alnum(s) in alnum(page_text(q[1], q[2]))
print(f"   trùng ngẫu nhiên số tờ A trên trang tờ B: {h2}/{n2}")

# ---------- D1e: tên chi tiết (cột Q 'ID-.. #12: #12 DETAIL - ...') — từ của tên có trên trang
c3 = collections.Counter(); ex = []
STOPW = {"DETAIL", "PLAN", "SECTION", "ELEVATION", "TYPICAL", "AND", "THE", "WITH"}
for r in T:
    q = r[16]
    if not q: continue
    for seg in str(q).split(";;"):
        m = SH.search(seg); mt = re.search(r"#(\d{1,2})\s*([A-Z][^;#]*)$", seg.strip())
        if not m or not mt or m.group(0) not in sheet_pg: continue
        words = {w for w in re.findall(r"[A-Z]{3,}", fold(mt.group(2))) if w not in STOPW}
        if not words: continue
        t = page_text(*sheet_pg[m.group(0)])
        if kkd(t): c3["KHÔNG KIỂM ĐƯỢC"] += 1; continue
        tw = set(re.findall(r"[A-Z]{3,}", fold(t)))
        f = len(words & tw) / len(words)
        c3["CÓ NEO" if f >= 0.8 else "LỆCH"] += 1
        if f < 0.8 and len(ex) < 6: ex.append((r[2], m.group(0), mt.group(2)[:50], round(f, 2)))
print("=== D1e tên chi tiết (Q) trên trang tờ (≥80% từ) ===", dict(c3))
for e in ex: print("  LỆCH", e)

# ---------- D3: kích thước DWG có trong lớp text trang PDF tách sẵn
dw = {}
for f in glob.glob(SCAN + r"\_dwg_tangA\json\*.json"):
    j = json.load(open(f, encoding="utf-8"))
    vals = [v for d in j.get("details", []) for v in d.get("vals", [])]
    if j.get("id") and vals: dw[j["id"]] = vals
NUM = re.compile(r"(?<![\d.])\d{2,5}(?![\d])")
st = collections.Counter(); per = []
for f in glob.glob(SCAN + r"\_sheets\*\*.pdf"):
    sid = re.sub(r"^p\d+_", "", os.path.basename(f))[:-4]
    if sid not in dw: continue
    t = page_text(f, 1) or ""
    nums = set(NUM.findall(t))
    want = {str(int(round(v))) for v in dw[sid] if v >= 10}
    if not want: continue
    hitn = len(want & nums)
    per.append((sid, len(want), hitn, len(t.strip())))
    st["tờ"] += 1; st["số DWG (≥10)"] += len(want); st["có trong text"] += hitn
    st["tờ có ≥50% số DWG trong text"] += hitn >= 0.5 * len(want)
    st["tờ text < 200 ký tự"] += len(t.strip()) < 200
print("\n=== D3 kích thước DWG trong lớp text PDF ===", dict(st))
per.sort(key=lambda x: -x[2] / x[1])
print("  tốt nhất:", per[:5]); print("  kém nhất:", per[-3:])
